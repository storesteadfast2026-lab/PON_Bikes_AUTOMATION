import os
import json
import shutil
import tempfile
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from django.db import IntegrityError, transaction
from django.urls import reverse
from django.utils import timezone
from django.contrib.staticfiles import finders

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor

from .forms import ContainerForm, ProductDefinitionForm
from .models import AppSetting, Container, Customer, GeneratedExport, Group1Family, ImportBatch, ImportProfile, NormalizedLine, ProductCatalog, ProductCatalogEntry, ProductDefinition, ShortNameDictionaryRule, SourceFile
from .services import (
    EXPORT_DEFAULTS,
    _is_excel_5_95,
    build_name_dictionary,
    build_new_product_workbook,
    build_product_check_workbook,
    build_client_receiving_workbook,
    build_client_report_rows,
    build_upstockserial_csv,
    compare_products_to_catalog,
    analyse_workbook_mapping,
    normalize_container_identifier,
    parse_product_catalog,
    parse_product_moves_csv,
    product_check_summary,
    reconcile_product_movements,
    resolve_group1_family,
    suggest_mapping,
    suggest_translogic_name,
)
from .translogic_transfer import copy_file_verified, discover_numbered_import_set


class RealFileWorkflowTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.temp_media = tempfile.TemporaryDirectory()
        cls.override = override_settings(MEDIA_ROOT=cls.temp_media.name)
        cls.override.enable()

    @classmethod
    def tearDownClass(cls):
        cls.override.disable()
        cls.temp_media.cleanup()
        super().tearDownClass()

    def setUp(self):
        self.user = get_user_model().objects.create_user(username="tester", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="TCNU4108637",
            customer=self.customer,
            container_date=date(2026, 9, 15),
            year=2026,
            job_order="TEST-JO-001",
            created_by=self.user,
        )
        AppSetting.objects.create(key="TL_EXPORT_FORMAT", value="xlsx")
        self.samples = Path(__file__).resolve().parents[1] / "sample_data"

    def _upload_configure_confirm(self, filename, kind):
        sample_path = self.samples / filename
        with sample_path.open("rb") as handle:
            response = self.client.post(
                reverse("receiving:upload_source", args=[self.container.pk]),
                {"kind": kind, "file": SimpleUploadedFile(filename, handle.read())},
            )
        self.assertEqual(response.status_code, 302)
        source = SourceFile.objects.get(container=self.container, kind=kind)
        mapping = suggest_mapping(source.file.path, kind)
        response = self.client.post(reverse("receiving:configure_import", args=[source.pk]), mapping)
        self.assertEqual(response.status_code, 302)
        batch = ImportBatch.objects.get(source_file=source, status="PREVIEW")
        response = self.client.post(reverse("receiving:confirm_batch", args=[batch.pk]))
        self.assertEqual(response.status_code, 302)
        batch.refresh_from_db()
        self.assertEqual(batch.status, "CONFIRMED")
        return batch

    def test_received_scan_stream_mapping_is_preserved(self):
        sample_path = self.samples / "TCNU4108637 PON CON40 15-09-26.xlsx"
        mapping = suggest_mapping(sample_path, "RECEIVED")
        self.assertEqual(mapping["sheet_name"], "Sheet1")
        self.assertEqual(mapping["start_row"], 1)
        self.assertEqual(mapping["code_column"], "A")
        self.assertTrue(mapping["location_stream"])

    def test_end_to_end_first_stage_with_real_files(self):
        expected = self._upload_configure_confirm("TCNU4108637.xlsx", "CLIENT")
        received = self._upload_configure_confirm("TCNU4108637 PON CON40 15-09-26.xlsx", "RECEIVED")
        self.assertEqual(expected.total_units, 78)
        self.assertEqual(received.total_units, 78)

        response = self.client.get(reverse("receiving:comparison", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"]["expected"], 78)
        self.assertEqual(response.context["totals"]["received"], 78)
        self.assertEqual(response.context["totals"]["exceptions"], 0)


class LoginPresentationTests(TestCase):
    def test_login_page_uses_application_styles(self):
        response = self.client.get(reverse("login"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'css/app.css')
        self.assertContains(response, 'class="login-shell"')
        self.assertContains(response, 'class="login-panel"')

    def test_main_stylesheet_is_discoverable(self):
        self.assertIsNotNone(finders.find("css/app.css"))

    def test_short_name_suggestion_preserves_model_colour_and_size(self):
        self.assertEqual(
            suggest_translogic_name("C27 SOLOIST FRAME BLACK MAGIC 48"),
            "C27 SOLOIST FR BLACK MAGIC 48",
        )
        self.assertEqual(
            suggest_translogic_name("SOLOIST FK 57.5MM OFFSET CLOUD BREAK 48"),
            "SOLOIST FK 57.5 OFF CLD BRK 48",
        )

    def test_xml_name_dictionary_prefers_historical_short_name(self):
        entries = [
            SimpleNamespace(
                long_name="C22 Aspero 5 Campagnolo Ekar Moss Plum 56",
                short_name="C22 Asp 5 Camp Ekar MossPlum56",
            ),
            SimpleNamespace(long_name="Focus Jam2 Advance Black", short_name="FC Jam2 Adv Blk"),
            SimpleNamespace(long_name="Focus Jam2 Advance Green", short_name="FC Jam2 Adv Grn"),
        ]
        dictionary = build_name_dictionary(entries)
        self.assertEqual(
            suggest_translogic_name(
                "C22 Aspero 5 Campagnolo Ekar Moss Plum 56",
                name_dictionary=dictionary,
            ),
            "C22 Asp 5 Camp Ekar MossPlum56",
        )
        self.assertIn("FOCUS", {rule["source"] for rule in dictionary["rules"]})

    def test_admin_short_name_rule_overrides_learned_json_rule(self):
        user = get_user_model().objects.create_user(username="dictionary-admin", password="test-password")
        catalog = ProductCatalog.objects.create(
            file="product_catalogs/products_pon_pbp_auto.xml",
            original_name="products_pon_pbp_auto.xml",
            sha256="a" * 64,
            row_count=1,
            name_dictionary={
                "exact": {},
                "rules": [
                    {"source": "KALKHOFF", "target": "KH", "count": 37, "confidence": 1.0, "saving": 6}
                ],
            },
            active=True,
            uploaded_by=user,
        )
        rule = ShortNameDictionaryRule.objects.create(
            catalog=catalog,
            source="KALKHOFF",
            target="KHF",
            occurrence_count=37,
            confidence=Decimal("1.0000"),
            active=True,
        )
        effective = catalog.effective_name_dictionary()
        self.assertEqual(effective["rules"][0]["target"], "KHF")

        rule.active = False
        rule.save()
        self.assertEqual(catalog.effective_name_dictionary()["rules"], [])

    def test_xml_parser_normalises_windows_nbsp_and_reads_metadata(self):
        xml = (
            b'<?xml version="1.0"?><xdoc><table><row>'
            b'<code>ABC</code><pon_sku>ABC</pon_sku><customer>PON</customer><code2>ABC</code2>'
            b'<short_name>K26 Test Grey</short_name><long_name>K26 Test\xA0Grey</long_name>'
            b'<group1>PONKHF</group1><group2>CHG02</group2>'
            b'</row></table></xdoc>'
        )
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "products_pon_pbp_auto.xml"
            path.write_bytes(xml)
            rows = parse_product_catalog(path)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["long_name"], "K26 Test Grey")
        self.assertEqual(rows[0]["group1"], "PONKHF")
        self.assertEqual(rows[0]["group2"], "CHG02")

    def test_xls_parser_reads_expected_product_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            fake_xls = folder / "products_pon_pbp_auto.xls"
            fake_xls.write_bytes(b"xls-placeholder")
            source_xlsx = folder / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Code", "PON SKU", "Customer", "Code2", "Short Name", "Long Name", "Group 1", "Group 2"])
            sheet.append(["D680511220", "D680511220", "PON", "D680511220", "F26 THRON", "F26 THRON2 Test", "PONFCS", "CHG02"])
            workbook.save(source_xlsx)

            def fake_convert(_source, destination):
                shutil.copyfile(source_xlsx, destination)

            with patch("receiving.services._convert_xls_catalog_to_xlsx", side_effect=fake_convert):
                rows = parse_product_catalog(fake_xls)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["code"], "D680511220")
        self.assertEqual(rows[0]["long_name"], "F26 THRON2 Test")
        self.assertEqual(rows[0]["group1"], "PONFCS")
        self.assertEqual(rows[0]["group2"], "CHG02")

    def test_group1_family_resolution_prefers_specific_high_priority_rule(self):
        rules = [
            {"id": 1, "family_name": "Focus", "suffix": "FCS", "match_phrases": "focus", "priority": 100, "active": True},
            {"id": 2, "family_name": "Focus E-Bikes", "suffix": "FEB", "match_phrases": "focus e bikes\nfocus e-bike", "priority": 220, "active": True},
        ]
        resolved = resolve_group1_family("Focus E Bikes F26 Jam2 6.8", "PON", rules)
        self.assertTrue(resolved["resolved"])
        self.assertEqual(resolved["group1"], "PONFEB")
        self.assertEqual(resolved["family_name"], "Focus E-Bikes")

    def test_group1_family_resolution_blocks_unknown_family(self):
        rules = [{"id": 1, "family_name": "Cervelo", "suffix": "CVL", "match_phrases": "cervelo", "priority": 100, "active": True}]
        resolved = resolve_group1_family("Unknown Brand Model 1", "PON", rules)
        self.assertFalse(resolved["resolved"])
        self.assertIn("Administration", resolved["message"])

    def test_excel_5_95_signature_validation(self):
        self.assertTrue(_is_excel_5_95(b"OLE-prefix\x09\x08\x08\x00\x00\x05BIFF-data"))
        self.assertFalse(_is_excel_5_95(b"OLE-prefix\x09\x08\x10\x00\x00\x06BIFF8-data"))

    def test_product_dimensions_are_entered_in_cm_and_stored_in_mm(self):
        form = ProductDefinitionForm(
            data={
                "long_code": "LONG-CM",
                "short_name": "CENTIMETRE TEST BIKE",
                "full_name": "Centimetre conversion test bicycle",
                "length_cm": "200.5",
                "height_cm": "118.0",
                "width_cm": "25.5",
                "weight_kg": "29.500",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        product = form.save(commit=False)
        self.assertEqual(product.length_mm, 2005)
        self.assertEqual(product.height_mm, 1180)
        self.assertEqual(product.width_mm, 255)

    def test_cubic_group2_boundary_and_full_product_names(self):
        products = [
            SimpleNamespace(
                code="EXACT-0400", long_code="", short_name="EXACT CUBIC PRODUCT",
                full_name="Exact cubic full product name", length_mm=1000, height_mm=1000,
                width_mm=400, weight_kg="10.000", group1="PONCVL",
            ),
            SimpleNamespace(
                code="ABOVE-0400", long_code="", short_name="ABOVE CUBIC PRODUCT",
                full_name="Above cubic full product name", length_mm=1000, height_mm=1000,
                width_mm=401, weight_kg="11.000", group1="PONCVL",
            ),
        ]
        workbook = load_workbook(
            build_new_product_workbook(products, dict(EXPORT_DEFAULTS), date(2026, 9, 14)),
            data_only=True,
        )
        sheet = workbook["New Product Import"]
        self.assertEqual(sheet.max_column, 30)
        self.assertEqual(sheet["K2"].value, None)
        self.assertEqual(sheet["T2"].value, 0.4)
        self.assertEqual(sheet["K3"].value, "CHG02")
        self.assertEqual(sheet["T3"].value, 0.401)
        self.assertLessEqual(len(sheet["C3"].value), 30)
        self.assertEqual(sheet["M3"].value, "Above cubic full product name")
        self.assertEqual(sheet["AD2"].value.date(), date(2026, 9, 14))
        self.assertEqual(sheet["AD3"].value.date(), date(2026, 9, 14))
        invalid_name = SimpleNamespace(
            code="NAME-TOO-LONG", long_code="", short_name="X" * 31,
            full_name="This complete name must remain available", length_mm=1000,
            height_mm=1000, width_mm=401, weight_kg="11.000", group1="PONCVL",
        )
        with self.assertRaisesRegex(ValueError, "at most 30 characters"):
            build_new_product_workbook([invalid_name], dict(EXPORT_DEFAULTS), date(2026, 9, 14))

    def test_printable_product_list_uses_first_pallet_and_sorts_by_loc(self):
        rows = [
            {
                "code": "BIKE-001", "long_code": "LONG-1", "count": 2, "location": "T1002",
                "status": "Existing", "is_new": False, "tl_code": "TL-1", "matched_by": "code",
                "catalog_customer": "PBP",
            },
            {
                "code": "bike-001", "long_code": "LONG-1", "count": 3, "location": "T1001",
                "status": "Existing", "is_new": False, "tl_code": "TL-1", "matched_by": "code",
                "catalog_customer": "PBP",
            },
            {
                "code": "BIKE-002", "long_code": "LONG-2", "count": 4, "location": "T1003",
                "status": "New", "is_new": True, "tl_code": "", "matched_by": "",
                "catalog_customer": "", "definition": SimpleNamespace(length_mm=1600, height_mm=900, width_mm=350, weight_kg=18.5),
            },
        ]
        workbook = load_workbook(build_product_check_workbook("TEST-CONTAINER", rows), data_only=False)
        check = workbook["PON Product Check"]
        self.assertEqual(check.max_row, 4)
        self.assertEqual(check["A2"].value, "BIKE-001")
        self.assertEqual(check["C2"].value, 5)
        self.assertEqual(check["D2"].value, "T1001")
        self.assertEqual(check["A3"].value, "BIKE-002")
        self.assertEqual(check["D3"].value, "T1003")
        self.assertEqual(check["E2"].value, "Existing")
        self.assertEqual(check["E3"].value, "New")
        self.assertEqual(
            [check[f"{column}1"].value for column in ("F", "G", "H")],
            ["Length (cm)", "Width (cm)", "Height (cm)"],
        )
        self.assertEqual(check["F3"].value, 160.0)
        self.assertEqual(check["G3"].value, 35.0)
        self.assertEqual(check["H3"].value, 90.0)
        self.assertEqual(check["I3"].value, 18.5)
        self.assertEqual(check.column_dimensions["A"].width, 15)
        self.assertEqual(check.column_dimensions["C"].width, 7)
        self.assertEqual(check.column_dimensions["D"].width, 9)
        self.assertEqual(check["C2"].alignment.horizontal, "center")
        self.assertEqual(check["D2"].alignment.horizontal, "center")
        self.assertTrue(check.column_dimensions["J"].hidden)
        self.assertTrue(check.column_dimensions["K"].hidden)
        self.assertIn("$A$1:$I$4", str(check.print_area))
        self.assertEqual(check.page_setup.orientation, "portrait")

        barcode = workbook["BARCODE Converter"]
        self.assertEqual(barcode.max_row, 4)
        self.assertEqual(barcode["D3"].value, 5)
        self.assertIsNone(barcode.freeze_panes)
        self.assertEqual(len(barcode._images), 4)
        for cell in (barcode["B2"], barcode["D2"], barcode["E2"], barcode["F2"], barcode["G2"], barcode["B3"], barcode["D3"], barcode["G3"]):
            self.assertEqual(cell.font.sz, 14)
        for image in barcode._images:
            self.assertIsInstance(image.anchor, OneCellAnchor)
            self.assertGreater(image.anchor._from.colOff, 0)
            self.assertGreater(image.anchor._from.rowOff, 0)


class ProductCatalogWorkflowTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=self.temp_media.name)
        self.override.enable()
        self.user = get_user_model().objects.create_user(username="catalog-tester", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="FDCU0321443",
            customer=self.customer,
            container_date=date(2026, 9, 14),
            year=2026,
            created_by=self.user,
        )
        AppSetting.objects.create(key="TL_EXPORT_FORMAT", value="xlsx")
        source = SourceFile.objects.create(
            container=self.container,
            kind="RECEIVED",
            file=SimpleUploadedFile("received.xlsx", b"placeholder"),
            original_name="received.xlsx",
            sha256="a" * 64,
            status="CONFIRMED",
            uploaded_by=self.user,
        )
        self.batch = ImportBatch.objects.create(
            source_file=source,
            status="CONFIRMED",
            total_rows=2,
            total_units=7,
            created_by=self.user,
        )
        NormalizedLine.objects.create(
            batch=self.batch,
            source_sheet="Sheet1",
            source_row=2,
            code="04-14503",
            quantity=4,
            location="T4324",
        )
        NormalizedLine.objects.create(
            batch=self.batch,
            source_sheet="Sheet1",
            source_row=3,
            code="NEW-001",
            quantity=3,
            location="T4325",
        )

    def tearDown(self):
        self.override.disable()
        self.temp_media.cleanup()

    @staticmethod
    def _catalog_bytes():
        return b"""<?xml version=\"1.0\"?>
<xdoc><table>
<row><code>000414503192</code><pon_sku>04-14503</pon_sku><customer>PON</customer><code2>04-14503</code2><short_name>Bearing Kit BlurTR1 5010 1 TB2</short_name><long_name>Bearing Kit BlurTR1 5010 1 TB2</long_name><group1>PONSCP</group1><group2/></row>
<row><code>NEW-001-X</code><pon_sku>OTHER-SKU</pon_sku><customer>PBP</customer><code2>OTHER-CODE</code2><short_name>C22 Asp 5 Camp Ekar MossPlum56</short_name><long_name>C22 Aspero 5 Campagnolo Ekar Moss Plum 56</long_name><group1>PBPCVL</group1><group2>CHG02</group2></row>
</table></xdoc>"""

    def test_syncs_exact_internal_catalog_and_classifies_exact_matches(self):
        with tempfile.TemporaryDirectory() as folder:
            catalog_path = Path(folder) / "products_pon_pbp_auto.xml"
            catalog_path.write_bytes(self._catalog_bytes())
            with override_settings(PON_PRODUCT_CATALOG_PATH=str(catalog_path)):
                response = self.client.post(reverse("receiving:sync_product_catalog"), {"customer_id": self.customer.pk})
        self.assertEqual(response.status_code, 302)
        catalog = ProductCatalog.objects.get(active=True)
        self.assertEqual(catalog.row_count, 2)
        self.assertEqual(catalog.customer, self.customer)
        known = catalog.entries.get(pon_sku="04-14503")
        self.assertEqual(known.short_name, "Bearing Kit BlurTR1 5010 1 TB2")
        self.assertEqual(known.group1, "PONSCP")
        self.assertEqual(known.group2, "")

        response = self.client.get(reverse("receiving:product_check", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        rows = {row["code"]: row for row in response.context["rows"]}
        self.assertEqual(rows["04-14503"]["status"], "Existing")
        self.assertEqual(rows["04-14503"]["matched_by"], "pon_sku, code2")
        self.assertEqual(rows["NEW-001"]["status"], "New")
        self.assertEqual(response.context["summary"]["new_units"], 3)
        self.assertEqual(response.context["saved_definition_count"], 0)
        self.assertEqual(response.context["pending_definition_count"], 1)

    def test_new_product_editor_is_first_compact_and_tab_ready(self):
        upload = SimpleUploadedFile("products_pon_pbp_auto.xml", self._catalog_bytes(), content_type="application/xml")
        self.client.post(reverse("receiving:upload_product_catalog"), {"file": upload, "customer_id": self.customer.pk})

        response = self.client.get(reverse("receiving:product_check", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertLess(html.index("Product data for Translogic"), html.index("Product classification"))
        self.assertIn('class="card new-product-entry"', html)
        self.assertIn('class="new-product-grid new-product-grid-header"', html)
        self.assertIn('data-measure="length"', html)
        self.assertIn('data-measure="height"', html)
        self.assertIn('data-measure="width"', html)
        self.assertIn('data-measure="weight"', html)
        self.assertIn("Proposed from previous row", html)
        self.assertIn("Short name (max 30)", html)
        self.assertIn("Long name", html)
        self.assertIn("Group1 family mappings", html)
        self.assertIn("There is no per-row confirmation step", html)

    def test_export_contains_product_check_and_printable_barcodes(self):
        upload = SimpleUploadedFile("products_pon_pbp_auto.xml", self._catalog_bytes(), content_type="application/xml")
        response = self.client.post(reverse("receiving:upload_product_catalog"), {"file": upload, "customer_id": self.customer.pk})
        self.assertEqual(response.status_code, 302)
        AppSetting.objects.filter(key="TL_EXPORT_FORMAT").update(value="excel_5_95")

        response = self.client.get(reverse("receiving:export_product_check", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("PON_Product_Check_FDCU0321443", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["PON Product Check", "BARCODE Converter"])

        check = workbook["PON Product Check"]
        self.assertEqual([check.cell(1, column).value for column in range(6, 9)], ["Length (cm)", "Width (cm)", "Height (cm)"])
        statuses = {check.cell(row, 1).value: check.cell(row, 5).value for row in range(2, check.max_row)}
        self.assertEqual(statuses["04-14503"], "Existing")
        self.assertEqual(statuses["NEW-001"], "New")
        self.assertEqual(check.cell(check.max_row, 3).value, 7)

        barcode = workbook["BARCODE Converter"]
        self.assertEqual(barcode["B2"].value, "FDCU0321443")
        self.assertEqual(barcode["D3"].value, 4)
        self.assertEqual(barcode["G4"].value, 7)
        self.assertEqual(len(barcode._images), 4)
        export = GeneratedExport.objects.get(container=self.container, kind="PRODUCT_CHECK")
        self.assertEqual(export.file_format, "xlsx")
        self.assertTrue(export.original_name.endswith(".xlsx"))

    def test_step2_shows_product_check_download_before_new_product_dimensions(self):
        upload = SimpleUploadedFile("products_pon_pbp_auto.xml", self._catalog_bytes(), content_type="application/xml")
        self.client.post(reverse("receiving:upload_product_catalog"), {"file": upload, "customer_id": self.customer.pk})
        response = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertIn("Download Product Check + barcodes (.xlsx)", html)
        self.assertNotIn("Download New Product Import", html)
        self.assertIn('aria-label="New Product Import status"', html)
        self.assertIn("New Product Import", html)
        self.assertIn("Pending — complete the required Translogic product data first.", html)

    def test_step2_shows_new_product_import_only_when_existing_export_rules_are_satisfied(self):
        upload = SimpleUploadedFile("products_pon_pbp_auto.xml", self._catalog_bytes(), content_type="application/xml")
        self.client.post(reverse("receiving:upload_product_catalog"), {"file": upload, "customer_id": self.customer.pk})
        Group1Family.objects.create(
            family_name="Test family",
            suffix="TST",
            match_phrases="New product",
            priority=100,
            active=True,
        )
        ProductDefinition.objects.create(
            code="NEW-001",
            long_code="LONG-001",
            short_name="New product",
            full_name="New product",
            group1="PONTST",
            length_mm=2000,
            width_mm=250,
            height_mm=1180,
            weight_kg=Decimal("29.500"),
            confirmed=True,
            created_by=self.user,
            updated_by=self.user,
        )
        response = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
        self.assertContains(response, "Download Product Check + barcodes (.xlsx)")
        self.assertContains(response, "Download New Product Import")
        self.assertNotContains(response, 'aria-label="New Product Import status"')
        self.assertContains(response, reverse("receiving:export_new_products", args=[self.container.pk]))

    def test_product_check_prints_dimensions_in_cm_without_changing_stored_mm(self):
        upload = SimpleUploadedFile("products_pon_pbp_auto.xml", self._catalog_bytes(), content_type="application/xml")
        self.client.post(reverse("receiving:upload_product_catalog"), {"file": upload, "customer_id": self.customer.pk})
        definition = ProductDefinition.objects.create(
            code="NEW-001",
            long_code="LONG-001",
            short_name="New product",
            full_name="New product",
            group1="PONTEST",
            length_mm=2000,
            width_mm=250,
            height_mm=1180,
            weight_kg=Decimal("29.500"),
            confirmed=True,
            created_by=self.user,
            updated_by=self.user,
        )
        response = self.client.get(reverse("receiving:export_product_check", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        sheet = workbook["PON Product Check"]
        row = next(r for r in range(2, sheet.max_row) if sheet.cell(r, 1).value == "NEW-001")
        self.assertEqual([sheet.cell(row, c).value for c in range(6, 9)], [200, 25, 118])
        definition.refresh_from_db()
        self.assertEqual((definition.length_mm, definition.width_mm, definition.height_mm), (2000, 250, 1180))

    def test_new_product_import_uses_exact_columns_and_cubic_group_rule(self):
        upload = SimpleUploadedFile("products_pon_pbp_auto.xml", self._catalog_bytes(), content_type="application/xml")
        self.client.post(reverse("receiving:upload_product_catalog"), {"file": upload, "customer_id": self.customer.pk})
        Group1Family.objects.update_or_create(family_name="Cervelo", defaults={"suffix": "CVL", "match_phrases": "cervelo\nC26\nC27", "priority": 100, "active": True})
        save_response = self.client.post(
            reverse("receiving:product_check", args=[self.container.pk]),
            {
                "product-0-long_code": "LONG-001",
                "product-0-short_name": "A SHORT TRANSLOGIC NAME",
                "product-0-full_name": "Cervelo A complete bicycle product name",
                "product-0-length_cm": "200.0",
                "product-0-height_cm": "118.0",
                "product-0-width_cm": "25.0",
                "product-0-weight_kg": "29.500",
            },
        )
        self.assertRedirects(save_response, reverse("receiving:product_check", args=[self.container.pk]))
        definition = ProductDefinition.objects.get(code="NEW-001")
        self.assertEqual(definition.length_mm, 2000)
        self.assertEqual(definition.height_mm, 1180)
        self.assertEqual(definition.width_mm, 250)
        self.assertEqual(definition.group1, "PONCVL")

        status_response = self.client.get(reverse("receiving:product_check", args=[self.container.pk]))
        self.assertEqual(status_response.context["saved_definition_count"], 1)
        self.assertEqual(status_response.context["pending_definition_count"], 0)
        self.assertContains(status_response, "Saved")
        saved_form = status_response.context["product_forms"][0]["form"]
        self.assertEqual(saved_form["length_cm"].value(), Decimal("200"))
        self.assertEqual(saved_form["height_cm"].value(), Decimal("118"))
        self.assertEqual(saved_form["width_cm"].value(), Decimal("25"))
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        sheet = workbook["New Product Import"]
        self.assertEqual(sheet.max_column, 30)
        self.assertEqual([sheet.cell(1, column).value for column in range(26, 31)], ["Inner", "Outer", "Layer", "Comment", "Date"])
        self.assertEqual(sheet["C2"].value, "A SHORT TRANSLOGIC NAME")
        self.assertLessEqual(len(sheet["C2"].value), 30)
        self.assertEqual(sheet["K2"].value, "CHG02")
        self.assertEqual(sheet["T2"].value, 0.59)
        self.assertEqual(sheet["V2"].value, 1180)
        self.assertEqual(sheet["W2"].value, 250)
        self.assertEqual(sheet["X2"].value, 2000)
        self.assertEqual(sheet["P2"].value, 1)
        self.assertEqual(sheet["R2"].value, 0)
        self.assertEqual(sheet["S2"].value, 1)
        self.assertEqual(sheet["J2"].value, "PONCVL")
        self.assertEqual(sheet["M2"].value, "Cervelo A complete bicycle product name")
        self.assertEqual(sheet["U2"].value, 29.5)
        self.assertEqual(sheet["AA2"].value, 1)
        self.assertEqual(sheet["AD2"].value.date(), date(2026, 9, 14))
        yellow_columns = {"A", "C", "D", "H", "J", "M", "P", "R", "S", "T", "U", "AA", "AD"}
        for column in yellow_columns:
            self.assertEqual(sheet[f"{column}1"].fill.fill_type, "solid")
            self.assertEqual(sheet[f"{column}1"].fill.fgColor.rgb, "FFFFFF00")
        for column in {"B", "E", "F", "G", "I", "K", "L", "N", "O", "Q", "V", "W", "X", "Y", "Z", "AB", "AC"}:
            self.assertNotEqual(sheet[f"{column}1"].fill.fill_type, "solid")
        self.assertTrue(GeneratedExport.objects.filter(container=self.container, kind="NEW_PRODUCT").exists())

    def test_new_product_export_requires_user_entered_container_date(self):
        self.container.container_date = None
        self.container.save(update_fields=["container_date"])
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertRedirects(response, reverse("receiving:container_detail", args=[self.container.pk]))
        self.assertFalse(GeneratedExport.objects.filter(container=self.container, kind="NEW_PRODUCT").exists())

        response = self.client.post(
            reverse("receiving:edit_container", args=[self.container.pk]),
            {
                "job_order": "9399134",
                "identifier": self.container.identifier,
                "customer": self.customer.pk,
                "container_date": "2026-09-12",
                "year": 2026,
                "notes": "Corrected operational date",
            },
        )
        self.assertRedirects(response, reverse("receiving:container_detail", args=[self.container.pk]))
        self.container.refresh_from_db()
        self.assertEqual(self.container.container_date, date(2026, 9, 12))

    def test_accepts_customer_specific_xml_filename(self):
        upload = SimpleUploadedFile("OTHER_PRODUCTS.xml", self._catalog_bytes(), content_type="application/xml")
        response = self.client.post(
            reverse("receiving:upload_product_catalog"),
            {"file": upload, "customer_id": self.customer.pk},
        )
        self.assertEqual(response.status_code, 302)
        catalog = ProductCatalog.objects.get()
        self.assertEqual(catalog.original_name, "OTHER_PRODUCTS.xml")
        self.assertEqual(catalog.customer, self.customer)

    def test_dashboard_hides_add_customer_and_shows_customer_xml_location(self):
        self.customer.product_catalog_source_path = r"C:\Data\PON_PRODUCTS.xml"
        self.customer.product_catalog_app_path = "/data/pon_products/PON_PRODUCTS.xml"
        self.customer.save(update_fields=["product_catalog_source_path", "product_catalog_app_path"])
        response = self.client.get(reverse("receiving:dashboard"))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertNotIn("Add customer", html)
        self.assertIn("PON_PRODUCTS.xml", html)
        catalog_data = response.context["customer_catalogs"][str(self.customer.pk)]
        self.assertEqual(catalog_data["source_path"], r"C:\Data\PON_PRODUCTS.xml")


class PbpToPonMigrationTests(TestCase):
    SKU = "0L0CAB111A48"

    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=self.temp_media.name)
        self.override.enable()
        self.user = get_user_model().objects.create_user(username="pbp-migration", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="ONEU1569467",
            customer=self.customer,
            container_date=date(2026, 9, 30),
            year=2026,
            created_by=self.user,
        )
        AppSetting.objects.create(key="TL_EXPORT_FORMAT", value="xlsx")
        source = SourceFile.objects.create(
            container=self.container,
            kind="RECEIVED",
            file=SimpleUploadedFile("received.xlsx", b"placeholder"),
            original_name="received.xlsx",
            sha256="b" * 64,
            status="CONFIRMED",
            uploaded_by=self.user,
        )
        self.batch = ImportBatch.objects.create(
            source_file=source,
            status="CONFIRMED",
            total_rows=1,
            total_units=1,
            created_by=self.user,
        )
        NormalizedLine.objects.create(
            batch=self.batch,
            source_sheet="Sheet1",
            source_row=2,
            code=self.SKU,
            quantity=1,
            location="T4324",
        )
        self.catalog = ProductCatalog.objects.create(
            customer=self.customer,
            file=SimpleUploadedFile("products_pon_pbp_auto.xls", b"placeholder"),
            original_name="products_pon_pbp_auto.xls",
            sha256="c" * 64,
            row_count=1,
            active=True,
            uploaded_by=self.user,
        )
        self.pbp_entry = ProductCatalogEntry.objects.create(
            catalog=self.catalog,
            source_row=10,
            code=self.SKU,
            pon_sku=self.SKU,
            customer="PBP",
            code2=self.SKU,
            short_name="C26 CALEDONIA 105 MOCHA 48",
            long_name="C26 CALEDONIA 105 MOCHA 48",
            group1="PBPCVL",
            group2="",
            raw_data={
                "code": self.SKU,
                "pon_sku": self.SKU,
                "customer": "PBP",
                "code2": self.SKU,
                "short_name": "C26 CALEDONIA 105 MOCHA 48",
                "long_name": "C26 CALEDONIA 105 MOCHA 48",
                "group1": "PBPCVL",
                "group2": "",
                "cubic": "0.500",
                "weight": "13.250",
                "height": "900",
                "width": "300",
                "length": "1800",
            },
        )
        Group1Family.objects.update_or_create(
            family_name="Cervelo",
            defaults={
                "suffix": "CVL",
                "match_phrases": "CALEDONIA\nC26",
                "priority": 100,
                "active": True,
            },
        )

    def tearDown(self):
        self.override.disable()
        self.temp_media.cleanup()

    def _line(self, code=None):
        return SimpleNamespace(code=code or self.SKU, location="T4324", quantity=1, long_code="")

    def test_three_way_classification_prefers_pon_then_pbp_then_new(self):
        rows = compare_products_to_catalog([self._line()], [self.pbp_entry], target_customer="PON")
        self.assertEqual(rows[0]["status"], "PBP → PON")
        self.assertFalse(rows[0]["is_new"])
        self.assertTrue(rows[0]["is_pbp_to_pon"])

        pon_entry = SimpleNamespace(
            code=self.SKU,
            pon_sku=self.SKU,
            code2=self.SKU,
            customer="PON",
            short_name="PON NAME",
            long_name="PON LONG NAME",
            group1="PONCVL",
            group2="",
            raw_data={},
        )
        rows = compare_products_to_catalog([self._line()], [self.pbp_entry, pon_entry], target_customer="PON")
        self.assertEqual(rows[0]["status"], "Existing")
        self.assertEqual(rows[0]["catalog_customer"], "PON")

        rows = compare_products_to_catalog([self._line("UNKNOWN-123")], [self.pbp_entry], target_customer="PON")
        self.assertEqual(rows[0]["status"], "New")
        self.assertTrue(rows[0]["is_new"])

    def test_valid_pbp_history_does_not_request_manual_dimensions(self):
        response = self.client.get(reverse("receiving:product_check", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        row = response.context["rows"][0]
        self.assertEqual(row["status"], "PBP → PON")
        self.assertTrue(row["pbp_to_pon_data"]["valid"])
        self.assertEqual(response.context["summary"]["current_new_products"], 0)
        self.assertEqual(response.context["summary"]["current_pbp_to_pon_products"], 1)
        self.assertEqual(response.context["summary"]["original_pbp_to_pon_products"], 1)
        self.assertEqual(response.context["product_forms"], [])
        self.assertContains(response, "PBP → PON")

    def test_product_check_summary_counts_unique_codes(self):
        rows = [
            {"code": "EXISTING-1", "location": "A", "count": 1, "classification": "EXISTING"},
            {"code": "existing-1", "location": "B", "count": 2, "classification": "EXISTING"},
            {"code": "PBP-1", "location": "A", "count": 1, "classification": "PBP_TO_PON"},
            {"code": "NEW-1", "location": "A", "count": 1, "classification": "NEW"},
        ]
        summary = product_check_summary(rows, {
            "EXISTING-1": "PBP_TO_PON",
            "PBP-1": "NEW",
            "NEW-1": "EXISTING",
        })

        self.assertEqual(summary["unique_products"], 3)
        self.assertEqual(summary["current_existing_products"], 1)
        self.assertEqual(summary["current_pbp_to_pon_products"], 1)
        self.assertEqual(summary["current_new_products"], 1)
        self.assertEqual(summary["original_existing_products"], 1)
        self.assertEqual(summary["original_pbp_to_pon_products"], 1)
        self.assertEqual(summary["original_new_products"], 1)

    def test_product_check_preserves_original_classification_when_current_changes(self):
        from receiving.views import _product_check_data

        _product_check_data(self.container)
        state = self.container.product_statuses.get(code=self.SKU)
        self.assertEqual(state.import_classification, "PBP_TO_PON")

        NormalizedLine.objects.create(
            batch=self.batch, source_sheet="Sheet1", source_row=3,
            code=self.SKU, quantity=1, location="T9999",
        )
        self._add_current_pon(self.SKU)

        response = self.client.get(reverse("receiving:product_check", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["rows"]), 2)
        summary = response.context["summary"]
        self.assertEqual(summary["unique_products"], 1)
        self.assertEqual(summary["original_pbp_to_pon_products"], 1)
        self.assertEqual(summary["original_existing_products"], 0)
        self.assertEqual(summary["current_existing_products"], 1)
        self.assertEqual(summary["current_pbp_to_pon_products"], 0)
        self.assertContains(response, "EXISTING PRODUCTS")
        self.assertContains(response, "PBP → PON PRODUCTS")
        self.assertContains(response, "NEW PRODUCTS")
        self.assertNotContains(response, "Existing product rows")
        self.assertNotContains(response, "PBP → PON rows")
        self.assertNotContains(response, "New product rows")

        state.refresh_from_db()
        self.assertEqual(state.import_classification, "PBP_TO_PON")

        export = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(export.status_code, 200)
        sheet = load_workbook(BytesIO(export.content), data_only=True)["New Product Import"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, self.SKU)
        self.assertTrue(all(cell.fill.fgColor.rgb == "FFF4B183" for cell in sheet[2]))

    def test_pbp_to_pon_is_included_in_existing_new_product_import_workflow(self):
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        sheet = workbook["New Product Import"]
        self.assertEqual(sheet["A2"].value, self.SKU)
        self.assertEqual(sheet["B2"].value, self.SKU)
        self.assertEqual(sheet["C2"].value, "C26 CALEDONIA 105 MOCHA 48")
        self.assertEqual(sheet["D2"].value, "PON")
        self.assertEqual(sheet["H2"].value, "L")
        self.assertEqual(sheet["J2"].value, "PONCVL")
        self.assertEqual(sheet["K2"].value, "CHG02")
        self.assertEqual(sheet["M2"].value, "C26 CALEDONIA 105 MOCHA 48")
        self.assertEqual(sheet["P2"].value, 1)
        self.assertEqual(sheet["R2"].value, 0)
        self.assertEqual(sheet["S2"].value, 1)
        self.assertEqual(sheet["T2"].value, 0.5)
        self.assertEqual(sheet["U2"].value, 13.25)
        self.assertEqual(sheet["V2"].value, 900)
        self.assertEqual(sheet["W2"].value, 300)
        self.assertEqual(sheet["X2"].value, 1800)
        self.assertEqual(sheet["AA2"].value, 1)
        self.assertFalse(ProductDefinition.objects.filter(code=self.SKU).exists())

    def _add_current_pon(self, code):
        return ProductCatalogEntry.objects.create(
            catalog=self.catalog, source_row=20, code=code, pon_sku=code,
            customer="PON", code2=code,
            short_name="C26 CALEDONIA 105 MOCHA 48",
            long_name="C26 CALEDONIA 105 MOCHA 48",
            group1="PONCVL",
            raw_data={
                "code": code,
                "pon_sku": code,
                "customer": "PON",
                "code2": code,
                "short_name": "C26 CALEDONIA 105 MOCHA 48",
                "long_name": "C26 CALEDONIA 105 MOCHA 48",
                "group1": "PONCVL",
                "cubic": "0.252",
                "weight": "15.63",
                "height": "750",
                "width": "250",
                "length": "1340",
            },
        )

    def test_import_retains_pbp_selection_and_values_after_pon_sync(self):
        url = reverse("receiving:export_new_products", args=[self.container.pk])
        first = self.client.get(url)
        self._add_current_pon(self.SKU)
        self.pbp_entry.delete()
        second = self.client.get(url)
        self.assertEqual(second.status_code, 200)
        before = load_workbook(BytesIO(first.content), data_only=True).active
        after = load_workbook(BytesIO(second.content), data_only=True).active
        self.assertNotEqual(list(before.values), list(after.values))
        self.assertEqual(after["T2"].value, 0.252)
        self.assertEqual(after["U2"].value, 15.63)
        self.assertEqual(after["V2"].value, 750)
        self.assertEqual(after["W2"].value, 250)
        self.assertEqual(after["X2"].value, 1340)
        for cell in after[2]:
            self.assertEqual(cell.fill.fgColor.rgb, "FFF4B183")
        from receiving.views import _container_workspace_data
        self.assertTrue(_container_workspace_data(self.container)["new_product_import_download_ready"])

    def test_valid_pon_row_wins_data_without_changing_historical_pbp_classification(self):
        self.client.get(reverse("receiving:product_check", args=[self.container.pk]))
        state = self.container.product_statuses.get(code=self.SKU)
        original_history = dict(state.import_history)
        self.assertEqual(state.import_classification, "PBP_TO_PON")

        ProductCatalogEntry.objects.create(
            catalog=self.catalog,
            source_row=20,
            code=self.SKU,
            pon_sku=self.SKU,
            customer="PON",
            code2=self.SKU,
            short_name="C26 CALEDONIA 105 MOCHA 48",
            long_name="C26 CALEDONIA 105 MOCHA 48",
            group1="PONCVL",
            group2="",
            raw_data={
                "code": self.SKU,
                "pon_sku": self.SKU,
                "customer": "PON",
                "code2": self.SKU,
                "short_name": "C26 CALEDONIA 105 MOCHA 48",
                "long_name": "C26 CALEDONIA 105 MOCHA 48",
                "group1": "PONCVL",
                "group2": "",
                "cubic": "0.252",
                "weight": "15.63",
                "height": "750",
                "width": "250",
                "length": "1340",
            },
        )

        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        sheet = load_workbook(BytesIO(response.content), data_only=True)["New Product Import"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, self.SKU)
        self.assertEqual(sheet["J2"].value, "PONCVL")
        self.assertIsNone(sheet["K2"].value)
        self.assertEqual(sheet["T2"].value, 0.252)
        self.assertEqual(sheet["U2"].value, 15.63)
        self.assertEqual(sheet["V2"].value, 750)
        self.assertEqual(sheet["W2"].value, 250)
        self.assertEqual(sheet["X2"].value, 1340)
        self.assertTrue(all(cell.fill.fgColor.rgb == "FFF4B183" for cell in sheet[2]))
        self.assertFalse(ProductDefinition.objects.filter(code=self.SKU).exists())

        state.refresh_from_db()
        self.assertEqual(state.import_classification, "PBP_TO_PON")
        self.assertEqual(state.import_history, original_history)

    def test_readiness_uses_current_pon_when_historical_pbp_data_is_incomplete(self):
        from receiving.models import ContainerProductStatus
        from receiving.views import _container_workspace_data

        state = ContainerProductStatus.objects.create(
            container=self.container,
            code=self.SKU,
            was_new=False,
            import_classification="PBP_TO_PON",
            import_history={"valid": False, "missing": ["cubic", "weight", "height", "width", "length"]},
        )
        self._add_current_pon(self.SKU)
        self.pbp_entry.delete()

        workspace = _container_workspace_data(self.container)
        self.assertEqual(workspace["pending_new_codes"], [])
        self.assertTrue(workspace["new_product_import_download_ready"])

        page = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Download New Product Import")
        self.assertNotContains(page, "still require Long Name, Short Name, Group1 and dimensions")

        export = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(export.status_code, 200)
        sheet = load_workbook(BytesIO(export.content), data_only=True)["New Product Import"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, self.SKU)
        self.assertEqual(sheet["J2"].value, "PONCVL")
        self.assertEqual(sheet["T2"].value, 0.252)
        self.assertEqual(sheet["U2"].value, 15.63)

        state.refresh_from_db()
        self.assertEqual(state.import_classification, "PBP_TO_PON")
        self.assertFalse(state.import_history["valid"])

    def test_truly_new_retains_original_import_membership(self):
        from receiving.views import _product_check_data, _original_import_rows
        code = "TRULY-NEW"
        NormalizedLine.objects.create(batch=self.batch, source_sheet="Sheet1", source_row=3,
                                      code=code, quantity=1, location="T4324")
        ProductDefinition.objects.create(code=code, long_code=code, short_name="CALEDONIA",
                                         full_name="CALEDONIA", group1="PONCVL", length_mm=1000,
                                         height_mm=500, width_mm=200, weight_kg=10, created_by=self.user, updated_by=self.user)
        _product_check_data(self.container)
        self._add_current_pon(code)
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        row = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == code)
        self.assertTrue(all(cell.fill.patternType is None for cell in row))

    def test_original_existing_never_enters_import_after_master_removal(self):
        from receiving.views import _product_check_data
        entry = self._add_current_pon("ORIGINAL-EXISTING")
        NormalizedLine.objects.create(batch=self.batch, source_sheet="Sheet1", source_row=3,
                                      code=entry.code, quantity=1, location="T4324")
        _product_check_data(self.container)
        entry.delete()
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        self.assertEqual([row[0].value for row in sheet.iter_rows(min_row=2)], [self.SKU])

    def test_legacy_false_recovers_pbp_from_master_at_classification_time(self):
        from receiving.models import ContainerProductStatus
        ContainerProductStatus.objects.create(container=self.container, code=self.SKU, was_new=False)
        self.catalog.active = False
        self.catalog.save(update_fields=["active"])
        newer = ProductCatalog.objects.create(customer=self.customer, file="new-master.xml",
                                              original_name="new-master.xml", sha256="d" * 64,
                                              active=True, uploaded_by=self.user)
        ProductCatalogEntry.objects.create(catalog=newer, source_row=1, code=self.SKU,
                                          pon_sku=self.SKU, code2=self.SKU, customer="PON",
                                          short_name="C26 CALEDONIA 105 MOCHA 48",
                                          long_name="C26 CALEDONIA 105 MOCHA 48", group1="PONCVL",
                                          raw_data={
                                              "cubic": "0.252", "weight": "15.63",
                                              "height": "750", "width": "250", "length": "1340",
                                          })
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        state = self.container.product_statuses.get(code=self.SKU)
        self.assertEqual(state.import_classification, "PBP_TO_PON")
        self.assertEqual(load_workbook(BytesIO(response.content)).active["A2"].value, self.SKU)

    def test_selection_is_isolated_per_container(self):
        from receiving.views import _product_check_data, _original_import_rows
        _product_check_data(self.container)
        self._add_current_pon(self.SKU)
        other = Container.objects.create(identifier="SECOND-CONTAINER", customer=self.customer,
                                         year=2026, container_date=date(2026, 9, 30), created_by=self.user)
        from receiving.models import ContainerProductStatus
        ContainerProductStatus.objects.create(container=other, code=self.SKU, was_new=False,
                                              import_classification="EXISTING")
        rows = [{"code": self.SKU}]
        self.assertTrue(_original_import_rows(self.container, rows)[0]["is_pbp_to_pon"])
        self.assertFalse(_original_import_rows(other, rows)[0]["is_pbp_to_pon"])

    def test_container_get_uses_persisted_state_with_only_targeted_pon_lookup(self):
        from receiving.views import _product_check_data
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        _product_check_data(self.container)
        with patch("receiving.views.compare_products_to_catalog", side_effect=AssertionError("classification in GET")), \
             patch("receiving.views.parse_product_catalog", side_effect=AssertionError("master parse in GET")), \
             patch("receiving.views._product_check_data", side_effect=AssertionError("heavy read in GET")), \
             patch("receiving.views.build_new_product_workbook", side_effect=AssertionError("export in GET")), \
             CaptureQueriesContext(connection) as queries:
            for _ in range(2):
                response = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context["workspace"]["new_product_import_download_ready"])
        catalog_queries = []
        for query in queries:
            sql = query["sql"].lower()
            if "receiving_productcatalogentry" in sql:
                catalog_queries.append(sql)
            self.assertFalse(sql.lstrip().startswith(("insert", "update", "delete")))
        self.assertEqual(len(catalog_queries), 2)
        for sql in catalog_queries:
            self.assertIn('"customer" like', sql)
            self.assertIn('"code" in', sql)
            self.assertIn(self.SKU.lower(), sql)

    def test_legacy_container_get_does_not_rebuild_history(self):
        from receiving.models import ContainerProductStatus
        state = ContainerProductStatus.objects.create(container=self.container, code=self.SKU, was_new=False)
        with patch("receiving.views.compare_products_to_catalog", side_effect=AssertionError("legacy recovery in GET")), \
             patch("receiving.views.parse_product_catalog", side_effect=AssertionError("legacy parse in GET")):
            response = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
            self.assertEqual(response.status_code, 200)
            self.assertFalse(response.context["workspace"]["products_ready"])
        state.refresh_from_db()
        self.assertEqual(state.import_classification, "")
        self.assertEqual(state.import_history, {})

    def test_explicit_check_repairs_incomplete_original_pbp_history(self):
        from receiving.models import ContainerProductStatus
        ContainerProductStatus.objects.create(container=self.container, code=self.SKU, was_new=False,
                                              import_classification="PBP_TO_PON", import_history={"valid": False})
        check = self.client.get(reverse("receiving:product_check", args=[self.container.pk]))
        self.assertEqual(check.status_code, 200)
        self.assertTrue(self.container.product_statuses.get(code=self.SKU).import_history["valid"])
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        self.assertEqual(sheet["U2"].value, 13.25)
        self.assertEqual(sheet["X2"].value, 1800)
        self.assertTrue(self.container.product_statuses.get(code=self.SKU).import_history["valid"])

    def test_distinct_master_identifiers_and_historical_values_are_preserved(self):
        self.pbp_entry.code = "TL-PBP-CODE"
        self.pbp_entry.code2 = "ALTERNATE-CODE"
        self.pbp_entry.save(update_fields=["code", "code2"])
        response = self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        history = self.container.product_statuses.get(code=self.SKU).import_history
        self.assertEqual(history["code"], "TL-PBP-CODE")
        self.assertEqual(history["pon_sku"], self.SKU)
        self.assertEqual(history["code2"], "ALTERNATE-CODE")
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        self.assertEqual(sheet["A2"].value, "TL-PBP-CODE")
        self.assertEqual(sheet["B2"].value, self.SKU)
        self.assertEqual(sheet["J2"].value, "PONCVL")
        self.assertEqual(sheet["K2"].value, "CHG02")
        self.assertTrue(all(cell.fill.fgColor.rgb == "FFF4B183" for cell in sheet[2]))

    def test_recent_sidebar_uses_two_queries_for_twenty_containers(self):
        from receiving.views import _recent_container_rows
        for index in range(19):
            Container.objects.create(identifier=f"SIDEBAR-{index}", customer=self.customer,
                                     year=2026, created_by=self.user)
        with self.assertNumQueries(2):
            containers, rows = _recent_container_rows()
        self.assertEqual(len(containers), 20)
        self.assertEqual(len(rows), 20)

    def test_customer_report_does_not_mark_migration_new_and_colours_entire_row_orange(self):
        client = [SimpleNamespace(code=self.SKU, quantity=1, description="Client description")]
        received = [SimpleNamespace(code=self.SKU, quantity=1, description="")]
        rows = build_client_report_rows(
            client,
            received,
            [self.pbp_entry],
            [],
            product_statuses={self.SKU: False},
            product_classifications={self.SKU: "PBP_TO_PON"},
            target_customer="PON",
        )
        self.assertEqual(rows[0]["comment"], "")
        self.assertFalse(rows[0]["is_new"])
        self.assertTrue(rows[0]["is_pbp_to_pon"])
        self.assertEqual(rows[0]["advised"], 1)
        self.assertEqual(rows[0]["received"], 1)
        self.assertEqual(rows[0]["variance"], 0)

        workbook = load_workbook(build_client_receiving_workbook(self.container.identifier, rows), data_only=True)
        sheet = workbook["Receiving Container Data"]
        self.assertIsNone(sheet["G3"].value)
        for column in range(2, 8):
            self.assertEqual(sheet.cell(3, column).fill.fgColor.rgb, "FFF4B183")

class Stage2ServiceTests(TestCase):
    def test_product_moves_reconciliation_and_upstockserial_shape(self):
        source = BytesIO(
            b'product,stock_in_mov\r\n'
            b'"BIKE-A","   9399134"\r\n'
            b'"BIKE-A","   9399135"\r\n'
            b'"BIKE-B","   9399136"\r\n'
        )
        movements = parse_product_moves_csv(source)
        received = [
            SimpleNamespace(id=1, source_row=2, code="BIKE-A", quantity=1, location="T4324"),
            SimpleNamespace(id=2, source_row=3, code="BIKE-A", quantity=1, location="T4324"),
            SimpleNamespace(id=3, source_row=5, code="BIKE-B", quantity=1, location="T4325"),
        ]
        result = reconcile_product_movements(received, movements)
        self.assertTrue(result["ready"])
        self.assertEqual(
            [(row["movement"], row["serial_no"], row["location"]) for row in result["rows"]],
            [("9399134", "2", "T4324"), ("9399135", "3", "T4324"), ("9399136", "5", "T4325")],
        )
        content = build_upstockserial_csv(result["rows"])
        lines = content.decode("utf-8").splitlines()
        self.assertEqual(lines[0], "Movement,serial_no,Loc")
        self.assertEqual(lines[1], "   9399134,2,T4324")
        self.assertEqual(lines[2], "   9399135,3,T4324")
        self.assertEqual(lines[3], "   9399136,5,T4325")
        self.assertEqual(len(lines), 501)

    def test_upstockserial_is_sorted_by_movement_without_changing_assignment(self):
        rows = [
            {"movement": "9399161", "serial_no": "4", "location": "T4325"},
            {"movement": "9399134", "serial_no": "2", "location": "T4324"},
            {"movement": "9399135", "serial_no": "3", "location": "T4324"},
        ]
        content = build_upstockserial_csv(rows)
        lines = content.decode("utf-8").splitlines()
        self.assertEqual(lines[1], "   9399134,2,T4324")
        self.assertEqual(lines[2], "   9399135,3,T4324")
        self.assertEqual(lines[3], "   9399161,4,T4325")
        self.assertEqual(len(lines), 501)

    def test_stage2_blocks_product_count_mismatch(self):
        movements = parse_product_moves_csv(
            BytesIO(b'product,stock_in_mov\r\n"BIKE-A","   9399134"\r\n')
        )
        received = [
            SimpleNamespace(id=1, source_row=2, code="BIKE-A", quantity=1, location="T4324"),
            SimpleNamespace(id=2, source_row=3, code="BIKE-A", quantity=1, location="T4324"),
        ]
        result = reconcile_product_movements(received, movements)
        self.assertFalse(result["ready"])
        self.assertEqual(result["count_mismatches"][0]["difference"], -1)

    def test_product_moves_rejects_duplicate_movement(self):
        with self.assertRaises(ValueError):
            parse_product_moves_csv(
                BytesIO(
                    b'product,stock_in_mov\r\n'
                    b'"BIKE-A","   9399134"\r\n'
                    b'"BIKE-B","   9399134"\r\n'
                )
            )



class ClientReceivingReportServiceTests(TestCase):
    def test_report_rows_keep_client_order_calculate_variance_and_mark_new(self):
        client = [
            SimpleNamespace(code="EXIST-1", quantity=2, description="Client existing name"),
            SimpleNamespace(code="NEW-1", quantity=3, description="Client new bike long name"),
        ]
        received = [
            SimpleNamespace(code="EXIST-1", quantity=1, description=""),
            SimpleNamespace(code="NEW-1", quantity=3, description=""),
            SimpleNamespace(code="EXTRA-1", quantity=1, description="Unexpected bike"),
        ]
        catalog = [
            SimpleNamespace(code="EXIST-1", pon_sku="", code2="", long_name="XML Existing Long Name"),
            SimpleNamespace(code="EXTRA-1", pon_sku="", code2="", long_name="XML Extra Long Name"),
        ]
        definitions = [SimpleNamespace(code="NEW-1", full_name="Saved New Bike Long Name")]
        rows = build_client_report_rows(client, received, catalog, definitions, product_statuses={"NEW-1": True})
        self.assertEqual([row["code"] for row in rows], ["EXIST-1", "NEW-1", "EXTRA-1"])
        self.assertEqual(rows[0]["item_name"], "XML Existing Long Name")
        self.assertEqual(rows[0]["variance"], -1)
        self.assertEqual(rows[0]["comment"], "")
        self.assertEqual(rows[1]["item_name"], "Saved New Bike Long Name")
        self.assertEqual(rows[1]["variance"], 0)
        self.assertEqual(rows[1]["comment"], "New")
        self.assertEqual(rows[2]["advised"], 0)
        self.assertEqual(rows[2]["received"], 1)
        self.assertEqual(rows[2]["comment"], "")

    def test_report_workbook_matches_requested_columns_and_totals(self):
        rows = [
            {"code": "BIKE-1", "item_name": "Long bike one", "advised": 2, "received": 2, "variance": 0, "comment": "", "is_new": False},
            {"code": "BIKE-2", "item_name": "Long bike two", "advised": 3, "received": 2, "variance": -1, "comment": "New", "is_new": True},
        ]
        workbook = load_workbook(build_client_receiving_workbook("FDCU0321443", rows), data_only=True)
        sheet = workbook["Receiving Container Data"]
        self.assertEqual(sheet["B1"].value, "Container:")
        self.assertEqual(sheet["C1"].value, "FDCU0321443")
        self.assertEqual([sheet.cell(2, col).value for col in range(2, 8)], ["Item Code", "Item Name", "Advised", "Received", "Var", "Comment"])
        self.assertEqual([sheet.cell(3, col).value for col in range(2, 8)], ["BIKE-1", "Long bike one", 2, 2, 0, None])
        self.assertEqual([sheet.cell(4, col).value for col in range(2, 8)], ["BIKE-2", "Long bike two", 3, 2, -1, "New"])
        self.assertEqual(sheet["C5"].value, "Total")
        self.assertEqual(sheet["D5"].value, 5)
        self.assertEqual(sheet["E5"].value, 4)
        self.assertEqual(sheet["F5"].value, -1)
        for col in range(2, 8):
            self.assertEqual(sheet.cell(2, col).fill.fgColor.rgb, "FFFFFF00")


class GuidedWorkflowPresentationTests(TestCase):
    """Regression checks for the single-page guided container workflow."""

    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=self.temp_media.name)
        self.override.enable()
        self.user = get_user_model().objects.create_user(username="workflow-user", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="FDCU0321443",
            customer=self.customer,
            container_date=date(2026, 9, 14),
            year=2026,
            created_by=self.user,
        )

    def tearDown(self):
        self.override.disable()
        self.temp_media.cleanup()

    def test_container_page_is_the_guided_workflow(self):
        response = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Receive &amp; Scan")
        self.assertContains(response, "Check Products")
        self.assertContains(response, "Translogic")
        self.assertContains(response, "Working &amp; final locations")
        self.assertContains(response, "PRODUCT_MOVES.csv")
        self.assertContains(response, "UPStockSerial.csv")
        self.assertContains(response, 'data-stage-selector="1"')
        self.assertContains(response, 'data-stage-selector="2"')
        self.assertContains(response, 'data-stage-panel="1"')
        self.assertContains(response, 'data-stage-panel="2"')
        self.assertContains(response, "Generated:")
        self.assertContains(response, "Uploaded:")
        self.assertContains(response, "Expected Container ID:")
        self.assertContains(response, self.container.identifier)
        self.assertContains(response, f"*{self.container.identifier}*.xlsx")
        self.assertContains(response, 'class="container-aware-file-input"')
        self.assertContains(response, "Filename does not contain")
        self.assertContains(response, "Source file path")
        self.assertContains(response, "Translogic destination path")
        self.assertContains(response, "Copy to Translogic")
        self.assertContains(response, "Direct Translogic destination was not verified by the installer")
        self.assertNotContains(response, 'scrollIntoView')

    def test_container_detail_keeps_recent_container_sidebar_and_marks_current(self):
        other = Container.objects.create(
            identifier="ONEU1569467",
            customer=self.customer,
            container_date=date(2026, 9, 15),
            year=2026,
            created_by=self.user,
        )
        response = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Recent containers")
        self.assertContains(response, self.container.identifier)
        self.assertContains(response, other.identifier)
        self.assertContains(response, reverse("receiving:container_detail", args=[other.pk]))
        self.assertContains(response, 'aria-current="page"', count=1)
        self.assertContains(response, "current-container")
        self.assertContains(response, "container-detail-sidebar")

    def test_source_filename_container_warning_does_not_block_upload(self):
        response = self.client.post(
            reverse("receiving:upload_source", args=[self.container.pk]),
            {"kind": "CLIENT", "file": SimpleUploadedFile("manifest_without_id.xlsx", b"placeholder")},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(SourceFile.objects.filter(container=self.container, kind="CLIENT").exists())
        notices = [str(item) for item in get_messages(response.wsgi_request)]
        self.assertTrue(any(f"Filename does not contain the expected Container ID {self.container.identifier}" in item for item in notices))
        self.assertTrue(any("The upload was not blocked" in item for item in notices))

    def test_customer_report_download_name_changes_without_changing_content(self):
        content = b"unchanged-customer-report-content"
        export = GeneratedExport(
            container=self.container,
            kind="CLIENT_REPORT",
            original_name="stored-audit-name.xlsx",
            file_format="xlsx",
            row_count=1,
            sha256="b" * 64,
            generated_by=self.user,
        )
        export.file.save("stored-audit-name.xlsx", SimpleUploadedFile("stored-audit-name.xlsx", content), save=False)
        export.save()

        response = self.client.get(reverse("receiving:download_client_report", args=[export.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, content)
        self.assertEqual(
            response["Content-Disposition"],
            f'attachment; filename="{self.container.identifier} Advised.xls"',
        )

        with patch("receiving.views._ensure_current_client_report", return_value=export):
            response = self.client.get(reverse("receiving:download_current_client_report", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, content)
        self.assertEqual(
            response["Content-Disposition"],
            f'attachment; filename="{self.container.identifier} Advised.xls"',
        )

    def test_complete_is_blocked_until_outputs_are_current(self):
        response = self.client.post(reverse("receiving:mark_container_complete", args=[self.container.pk]))
        self.assertEqual(response.status_code, 302)
        self.container.refresh_from_db()
        self.assertNotEqual(self.container.status, "CLOSED")

    def test_closed_container_can_be_reopened(self):
        self.container.status = "CLOSED"
        self.container.save(update_fields=["status"])
        response = self.client.post(reverse("receiving:reopen_container", args=[self.container.pk]))
        self.assertEqual(response.status_code, 302)
        self.container.refresh_from_db()
        self.assertEqual(self.container.status, "OPEN")


class MultiContainerImportTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.temp_media = tempfile.TemporaryDirectory()
        cls.override = override_settings(MEDIA_ROOT=cls.temp_media.name)
        cls.override.enable()

    @classmethod
    def tearDownClass(cls):
        cls.override.disable()
        cls.temp_media.cleanup()
        super().tearDownClass()

    def setUp(self):
        self.user = get_user_model().objects.create_user(username="multi", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="ONEU2096348",
            customer=self.customer,
            container_date=date(2026, 9, 28),
            year=2026,
            job_order="9399134",
            created_by=self.user,
        )

    def _multi_workbook(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Purchase Lines"
        ws.append(["Buy-from Vendor Name", "No.", "Description", "Location Code", "Quantity", "Reserved Qty. (Base)"])
        ws.append(["FOCUS Bikes GmbH", "D690000001", "F26 ATLAS 8.7 Lavender XS", "3PL", 2, "ONEU 156 946 7"])
        ws.append(["FOCUS Bikes GmbH", "D690000002", "F26 ATLAS 8.7 Lavender S", "3PL", 3, "TGBU4621260"])
        ws.append(["FOCUS Bikes GmbH", "D690000003", "F26 ATLAS 8.7 Steelgrey M", "3PL", 4, "ONEU2096348"])
        ws.append(["FOCUS Bikes GmbH", "D690000004", "F26 ATLAS 8.7 Steelgrey L", "3PL", 1, "ONEU2096348"])
        extra = workbook.create_sheet("Sheet1")
        extra["C3"] = "PO000107"
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _upload_source(self):
        response = self.client.post(
            reverse("receiving:upload_source", args=[self.container.pk]),
            {"kind": "CLIENT", "file": SimpleUploadedFile("multi.xlsx", self._multi_workbook())},
        )
        self.assertEqual(response.status_code, 302)
        return SourceFile.objects.get(container=self.container, kind="CLIENT")

    def test_single_container_manifest_keeps_existing_behaviour(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Purchase Lines"
        ws.append(["No.", "Description", "Quantity", "Container"])
        ws.append(["D690000001", "F26 ATLAS Lavender XS", 2, "ONEU2096348"])
        ws.append(["D690000002", "F26 ATLAS Lavender S", 3, "ONEU2096348"])
        stream = BytesIO()
        workbook.save(stream)
        response = self.client.post(
            reverse("receiving:upload_source", args=[self.container.pk]),
            {"kind": "CLIENT", "file": SimpleUploadedFile("single.xlsx", stream.getvalue())},
        )
        self.assertEqual(response.status_code, 302)
        source = SourceFile.objects.get(container=self.container, kind="CLIENT")
        mapping = analyse_workbook_mapping(source.file.path, "CLIENT")["mapping"].copy()
        mapping.update({"save_profile": False, "profile_name": ""})
        self.client.post(reverse("receiving:configure_import", args=[source.pk]), mapping)
        batch = ImportBatch.objects.get(source_file=source, status="PREVIEW")
        self.client.post(reverse("receiving:confirm_batch", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, "CONFIRMED")
        self.assertEqual(batch.total_units, 5)
        self.assertEqual(Container.objects.count(), 1)

    def test_auto_mapping_detects_realistic_columns_and_containers(self):
        source = self._upload_source()
        analysis = analyse_workbook_mapping(source.file.path, "CLIENT")
        mapping = analysis["mapping"]
        self.assertEqual(mapping["sheet_name"], "Purchase Lines")
        self.assertEqual(mapping["start_row"], 2)
        self.assertEqual(mapping["code_column"], "B")
        self.assertEqual(mapping["description_column"], "C")
        self.assertEqual(mapping["quantity_column"], "E")
        self.assertEqual(mapping["container_column"], "F")
        self.assertEqual(mapping["long_code_column"], "")
        totals = {item["identifier"]: item["units"] for item in analysis["containers"]}
        self.assertEqual(totals, {"ONEU1569467": 2, "ONEU2096348": 5, "TGBU4621260": 3})
        self.assertEqual(normalize_container_identifier("ONEU 156 946 7"), "ONEU1569467")

    def test_confirm_splits_manifest_into_independent_pending_containers(self):
        source = self._upload_source()
        analysis = analyse_workbook_mapping(source.file.path, "CLIENT")
        mapping = analysis["mapping"].copy()
        mapping.update({"save_profile": False, "profile_name": ""})
        response = self.client.post(reverse("receiving:configure_import", args=[source.pk]), mapping)
        self.assertEqual(response.status_code, 302)
        batch = ImportBatch.objects.get(source_file=source, status="PREVIEW")
        response = self.client.post(reverse("receiving:confirm_batch", args=[batch.pk]))
        self.assertEqual(response.status_code, 302)
        batch.refresh_from_db()
        self.assertEqual(batch.total_units, 5)
        self.assertEqual(set(batch.lines.values_list("container_identifier", flat=True)), {"ONEU2096348"})

        oneu = Container.objects.get(identifier="ONEU1569467")
        tgbu = Container.objects.get(identifier="TGBU4621260")
        for created in (oneu, tgbu):
            self.assertEqual(created.status, "PENDING")
            self.assertEqual(created.job_order, "")
            self.assertTrue(created.auto_created_from_manifest)
            self.assertEqual(created.manifest_source_name, "multi.xlsx")
        self.assertEqual(
            ImportBatch.objects.get(source_file__container=oneu, source_file__kind="CLIENT", status="CONFIRMED").total_units,
            2,
        )
        self.assertEqual(
            ImportBatch.objects.get(source_file__container=tgbu, source_file__kind="CLIENT", status="CONFIRMED").total_units,
            3,
        )

    def test_existing_detected_container_is_not_duplicated(self):
        existing = Container.objects.create(
            identifier="TGBU4621260",
            customer=self.customer,
            container_date=date(2026, 9, 28),
            year=2026,
            job_order="9399999",
            created_by=self.user,
        )
        source = self._upload_source()
        mapping = analyse_workbook_mapping(source.file.path, "CLIENT")["mapping"].copy()
        mapping.update({"save_profile": False, "profile_name": ""})
        self.client.post(reverse("receiving:configure_import", args=[source.pk]), mapping)
        batch = ImportBatch.objects.get(source_file=source, status="PREVIEW")
        self.client.post(reverse("receiving:confirm_batch", args=[batch.pk]))
        self.assertEqual(Container.objects.filter(identifier="TGBU4621260").count(), 1)
        existing.refresh_from_db()
        self.assertEqual(existing.job_order, "9399999")
        self.assertEqual(existing.status, "OPEN")

    def test_auto_created_pending_container_requires_job_order_before_first_scan(self):
        source = self._upload_source()
        mapping = analyse_workbook_mapping(source.file.path, "CLIENT")["mapping"].copy()
        mapping.update({"save_profile": False, "profile_name": ""})
        self.client.post(reverse("receiving:configure_import", args=[source.pk]), mapping)
        batch = ImportBatch.objects.get(source_file=source, status="PREVIEW")
        self.client.post(reverse("receiving:confirm_batch", args=[batch.pk]))
        pending = Container.objects.get(identifier="TGBU4621260")
        response = self.client.post(
            reverse("receiving:upload_source", args=[pending.pk]),
            {"kind": "RECEIVED", "file": SimpleUploadedFile("scan.xlsx", self._multi_workbook())},
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(SourceFile.objects.filter(container=pending, kind="RECEIVED").exists())

    def test_saved_profile_is_recognised_for_same_structure(self):
        source = self._upload_source()
        first = analyse_workbook_mapping(source.file.path, "CLIENT")
        ImportProfile.objects.create(
            customer=self.customer,
            name="PON Purchase Lines",
            kind="CLIENT",
            mapping=first["mapping"],
        )
        profile = ImportProfile.objects.get(name="PON Purchase Lines")
        second = analyse_workbook_mapping(source.file.path, "CLIENT", saved_profiles=[profile])
        self.assertEqual(second["profile_match"], "PON Purchase Lines")
        self.assertEqual(second["confidence"]["code_column"], 100)
        self.assertEqual(second["confidence"]["container_column"], 100)

    def test_low_confidence_does_not_force_code_mapping(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Notes"
        ws.append(["Foo", "Bar"])
        ws.append(["hello", "world"])
        stream = BytesIO()
        workbook.save(stream)
        temp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp.write(stream.getvalue())
        temp.close()
        try:
            analysis = analyse_workbook_mapping(temp.name, "CLIENT")
            self.assertEqual(analysis["mapping"]["code_column"], "")
        finally:
            Path(temp.name).unlink(missing_ok=True)


class JobOrderUniquenessTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="job", password="test-password")
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.first = Container.objects.create(
            identifier="ONEU2096348", customer=self.customer, year=2026, job_order="9399134", created_by=self.user
        )

    def test_duplicate_job_order_is_rejected_by_form(self):
        form = ContainerForm(data={
            "job_order": "9399134",
            "identifier": "TGBU4621260",
            "customer": self.customer.pk,
            "container_date": "2026-09-28",
        })
        self.assertFalse(form.is_valid())
        self.assertIn("already assigned to container ONEU2096348", str(form.errors["job_order"]))

    def test_job_order_is_first_field(self):
        self.assertEqual(list(ContainerForm().fields)[:4], ["job_order", "identifier", "customer", "container_date"])

    def test_database_blocks_duplicate_nonblank_job_order(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Container.objects.create(
                    identifier="TGBU4621260", customer=self.customer, year=2026,
                    job_order="9399134", created_by=self.user
                )


    def test_normal_templates_do_not_offer_generate_buttons(self):
        template_dir = Path(__file__).resolve().parent / "templates" / "receiving"
        for name in ("container_detail.html", "stage2.html", "product_check.html", "comparison.html"):
            html = (template_dir / name).read_text(encoding="utf-8")
            self.assertNotIn(">Generate<", html)
            self.assertNotIn(">Generate ", html)
            self.assertNotIn(">Regenerate<", html)
            self.assertNotIn(">Regenerate ", html)
        urls = (Path(__file__).resolve().parent / "urls.py").read_text(encoding="utf-8")
        self.assertIn("copy-upstockserial", urls)
        self.assertIn("copy-to-translogic", urls)


class TranslogicFileTransferTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.media = root / "media"
        self.import_dir = root / "translogic_import"
        self.working_dir = root / "working"
        self.media.mkdir()
        self.import_dir.mkdir()
        self.working_dir.mkdir()
        self.override = override_settings(
            MEDIA_ROOT=str(self.media),
            PON_TRANSLOGIC_IMPORT_DIRECT_ENABLED=True,
            PON_TRANSLOGIC_IMPORT_DIRECT_PATH=str(self.import_dir),
            PON_UPSTOCKSERIAL_PATH=str(self.working_dir / "UPStockSerial.csv"),
            PON_UPSTOCK_WORKING_DISPLAY=r"C:\Docker-Projects\PON_Bike_Data\import\UPStockSerial.csv",
            PON_UPSTOCK_FINAL_DISPLAY=r"T:\Import\UPStockSerial.csv",
        )
        self.override.enable()
        self.user = get_user_model().objects.create_user(username="transfer", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="ONEU1569467",
            customer=self.customer,
            container_date=date(2026, 9, 29),
            year=2026,
            job_order="TRANSFER-001",
            created_by=self.user,
        )

    def tearDown(self):
        self.override.disable()
        self.temp.cleanup()

    def _export(self, kind, name, content, file_format="excel_5_95"):
        export = GeneratedExport(
            container=self.container,
            kind=kind,
            original_name=name,
            file_format=file_format,
            row_count=1,
            sha256=__import__("hashlib").sha256(content).hexdigest(),
            generated_by=self.user,
        )
        export.file.save(name, ContentFile(content), save=False)
        export.save()
        return export

    def test_numbered_import_discovery_uses_real_names_and_selects_oldest(self):
        base = 1_700_000_000
        for number in range(1, 6):
            path = self.import_dir / f"ActualTranslogicFile{number}.xls"
            path.write_bytes(f"old-{number}".encode())
            os.utime(path, (base + number * 100, base + number * 100))
        # Make file 4 the real oldest. The implementation must not assume file 1.
        os.utime(self.import_dir / "ActualTranslogicFile4.xls", (base - 500, base - 500))
        result = discover_numbered_import_set(self.import_dir)
        self.assertTrue(result["available"])
        self.assertEqual(result["pattern"], "ActualTranslogicFile{1-5}.xls")
        self.assertEqual(result["oldest"]["name"], "ActualTranslogicFile4.xls")

    def test_new_product_transfer_requires_confirmation_and_replaces_only_oldest(self):
        base = 1_700_000_000
        originals = {}
        for number in range(1, 6):
            path = self.import_dir / f"ImportReal{number}.xls"
            content = f"existing-{number}".encode()
            path.write_bytes(content)
            originals[path.name] = content
            os.utime(path, (base + number * 100, base + number * 100))
        target = self.import_dir / "ImportReal3.xls"
        os.utime(target, (base - 1000, base - 1000))
        generated = b"reviewed-new-product-import"
        self._export("NEW_PRODUCT", "NewProductImport_ONEU1569467.xls", generated)

        url = reverse("receiving:transfer_new_product_import", args=[self.container.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ImportReal{1-5}.xls")
        self.assertContains(response, "ImportReal3.xls")
        # GET is inspection only: nothing is replaced before explicit POST confirmation.
        self.assertEqual(target.read_bytes(), originals[target.name])

        expected_mtime = str(target.stat().st_mtime_ns)
        response = self.client.post(url, {
            "confirmed": "yes",
            "target_name": "ImportReal3.xls",
            "expected_destination_mtime": expected_mtime,
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(target.read_bytes(), generated)
        for number in (1, 2, 4, 5):
            path = self.import_dir / f"ImportReal{number}.xls"
            self.assertEqual(path.read_bytes(), originals[path.name])

    def test_new_product_transfer_blocks_format_mismatch_without_conversion(self):
        for number in range(1, 6):
            (self.import_dir / f"ImportReal{number}.csv").write_bytes(b"csv")
        self._export("NEW_PRODUCT", "NewProductImport_ONEU1569467.xls", b"xls-content")
        url = reverse("receiving:transfer_new_product_import", args=[self.container.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        messages = [str(message) for message in get_messages(response.wsgi_request)]
        self.assertTrue(any("Format mismatch" in message for message in messages))
        self.assertEqual((self.import_dir / "ImportReal1.csv").read_bytes(), b"csv")

    def test_upstock_transfer_get_does_not_copy_and_post_verifies_copy(self):
        source = self.working_dir / "UPStockSerial.csv"
        source.write_bytes(b"Movement,serial_no,Loc\r\n1,1,T0001\r\n")
        destination = self.import_dir / "UPStockSerial.csv"
        destination.write_bytes(b"previous")
        status = {
            "can_copy": True,
            "reason": "",
            "source_path": source,
            "source_display": r"C:\Docker-Projects\PON_Bike_Data\import\UPStockSerial.csv",
            "destination_path": destination,
            "destination_display": r"T:\Import\UPStockSerial.csv",
            "destination_exists": True,
            "destination_modified": None,
        }
        url = reverse("receiving:transfer_upstockserial", args=[self.container.pk])
        with patch("receiving.views._upstock_transfer_status", return_value=status):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(destination.read_bytes(), b"previous")
            expected_mtime = str(destination.stat().st_mtime_ns)
            response = self.client.post(url, {
                "confirmed": "yes",
                "expected_destination_mtime": expected_mtime,
            })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(destination.read_bytes(), source.read_bytes())

    def test_verified_copy_reports_replacement_and_matching_hash(self):
        source = self.working_dir / "source.bin"
        destination = self.import_dir / "destination.bin"
        source.write_bytes(b"new")
        destination.write_bytes(b"old")
        result = copy_file_verified(source, destination)
        self.assertTrue(result["replaced"])
        self.assertEqual(result["source_sha256"], result["destination_sha256"])
        self.assertEqual(destination.read_bytes(), b"new")


class ContinuitySnapshotStateTests(TestCase):
    """The continuity state exporter must be read-only and machine-readable."""

    def test_export_continuity_state_is_json_and_does_not_modify_operational_rows(self):
        user = get_user_model().objects.create_user(username="continuity-test", password="test-password")
        customer = Customer.objects.create(code="PON", name="Pon.Bike")
        Container.objects.create(
            identifier="CONTINUITY01",
            customer=customer,
            container_date=date(2026, 9, 29),
            year=2026,
            job_order="CONT-JO-001",
            created_by=user,
        )
        before = {
            "customers": Customer.objects.count(),
            "containers": Container.objects.count(),
            "sources": SourceFile.objects.count(),
            "batches": ImportBatch.objects.count(),
            "exports": GeneratedExport.objects.count(),
        }
        output = StringIO()
        call_command("export_continuity_state", stdout=output)
        payload = json.loads(output.getvalue())
        self.assertEqual(payload["application"]["project"], "PON Bikes Automation")
        self.assertEqual(payload["database"]["counts"]["containers"], before["containers"])
        self.assertIn("workflow", payload)
        self.assertIn("dictionary", payload)
        after = {
            "customers": Customer.objects.count(),
            "containers": Container.objects.count(),
            "sources": SourceFile.objects.count(),
            "batches": ImportBatch.objects.count(),
            "exports": GeneratedExport.objects.count(),
        }
        self.assertEqual(before, after)


class FirstScanScannerTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=self.temp_media.name)
        self.override.enable()
        self.user = get_user_model().objects.create_user(username="scanner", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="SCANCONT01",
            customer=self.customer,
            container_date=date(2026, 10, 1),
            year=2026,
            job_order="SCAN-JO-001",
            created_by=self.user,
        )
        source = SourceFile.objects.create(
            container=self.container,
            kind="CLIENT",
            file=SimpleUploadedFile("client.xlsx", b"client"),
            original_name="client.xlsx",
            sha256="1" * 64,
            status="CONFIRMED",
            uploaded_by=self.user,
        )
        self.client_batch = ImportBatch.objects.create(
            source_file=source,
            status="CONFIRMED",
            total_rows=2,
            total_units=5,
            confirmed_at=timezone.now(),
            created_by=self.user,
        )
        NormalizedLine.objects.create(
            batch=self.client_batch,
            source_sheet="Client",
            source_row=2,
            code="BIKE-A",
            long_code="LONG-A-0001",
            quantity=3,
            location="",
        )
        NormalizedLine.objects.create(
            batch=self.client_batch,
            source_sheet="Client",
            source_row=3,
            code="BIKE-B",
            long_code="LONG-B-0002",
            quantity=2,
            location="",
        )
        self.url = reverse("receiving:first_scan_scanner", args=[self.container.pk])

    def tearDown(self):
        self.override.disable()
        self.temp_media.cleanup()

    def _start(self, mode="AUTO"):
        from receiving.models import FirstScanSession
        response = self.client.post(self.url, {"action": "start", "mode": mode})
        self.assertEqual(response.status_code, 302)
        return FirstScanSession.objects.get(container=self.container, status__in=["ACTIVE", "PAUSED"])

    def _scan(self, value):
        return self.client.post(
            self.url,
            {"action": "scan", "scanned_value": value},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

    def test_code_scan_registers_one_bike(self):
        from receiving.models import FirstScanEvent
        session = self._start("CODE")
        self.assertEqual(self._scan("T1001").status_code, 200)
        response = self._scan("BIKE-A")
        self.assertEqual(response.status_code, 200)
        line = session.batch.lines.get()
        event = FirstScanEvent.objects.get(normalized_line=line)
        self.assertEqual(line.code, "BIKE-A")
        self.assertEqual(line.long_code, "LONG-A-0001")
        self.assertEqual(event.input_type, "CODE")

    def test_long_code_scan_resolves_product_code(self):
        from receiving.models import FirstScanEvent
        session = self._start("LONG_CODE")
        self._scan("T1001")
        response = self._scan("LONG-A-0001")
        self.assertEqual(response.status_code, 200)
        line = session.batch.lines.get()
        self.assertEqual(line.code, "BIKE-A")
        self.assertEqual(FirstScanEvent.objects.get(normalized_line=line).input_type, "LONG_CODE")

    def test_auto_mode_resolves_code_and_long_code(self):
        from receiving.models import FirstScanEvent
        session = self._start("AUTO")
        self._scan("T1001")
        self.assertEqual(self._scan("BIKE-A").status_code, 200)
        self.assertEqual(self._scan("LONG-B-0002").status_code, 200)
        self.assertEqual(list(session.batch.lines.values_list("code", flat=True)), ["BIKE-A", "BIKE-B"])
        self.assertEqual(
            list(FirstScanEvent.objects.filter(event_type="BIKE").order_by("scanned_at").values_list("input_type", flat=True)),
            ["CODE", "LONG_CODE"],
        )

    def test_pallet_scan_changes_current_pallet(self):
        session = self._start()
        first = self._scan("T1001")
        self.assertEqual(first.status_code, 200)
        session.refresh_from_db()
        self.assertEqual(session.current_pallet, "T1001")
        second = self._scan("T2002")
        self.assertEqual(second.status_code, 200)
        session.refresh_from_db()
        self.assertEqual(session.current_pallet, "T2002")

    def test_bike_is_assigned_to_current_pallet(self):
        session = self._start()
        self._scan("T4324")
        self._scan("BIKE-A")
        self.assertEqual(session.batch.lines.get().location, "T4324")

    def test_continuous_scans_require_no_confirmation(self):
        session = self._start()
        self._scan("T4324")
        first = self._scan("BIKE-A")
        second = self._scan("BIKE-A")
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(session.batch.lines.count(), 2)
        self.assertEqual(second.json()["state"]["current_total"], 2)

    def test_pause_resume_preserves_session_state_and_times(self):
        from receiving.models import FirstScanPause, FirstScanSession
        session = self._start()
        self._scan("T4324")
        self._scan("BIKE-A")
        session_id = session.pk
        pause = self.client.post(self.url, {"action": "pause"})
        self.assertEqual(pause.status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.status, "PAUSED")
        pause_record = FirstScanPause.objects.get(session=session)
        self.assertIsNotNone(pause_record.started_at)
        resume = self.client.post(self.url, {"action": "resume"})
        self.assertEqual(resume.status_code, 302)
        session.refresh_from_db()
        pause_record.refresh_from_db()
        self.assertEqual(session.pk, session_id)
        self.assertEqual(session.status, "ACTIVE")
        self.assertEqual(session.current_pallet, "T4324")
        self.assertEqual(session.batch.lines.count(), 1)
        self.assertIsNotNone(pause_record.resumed_at)
        self.assertEqual(FirstScanSession.objects.filter(container=self.container).count(), 1)

    def test_scan_and_finish_timestamps_are_stored(self):
        from receiving.models import FirstScanEvent
        session = self._start()
        self._scan("T4324")
        self._scan("BIKE-A")
        event = FirstScanEvent.objects.get(event_type="BIKE", result="SUCCESS")
        self.assertIsNotNone(session.started_at)
        self.assertIsNotNone(event.scanned_at)
        finish = self.client.post(self.url, {"action": "finish"})
        self.assertEqual(finish.status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.status, "FINISHED")
        self.assertIsNotNone(session.finished_at)

    def test_container_and_current_pallet_counters_update(self):
        self._start()
        self._scan("T1001")
        self._scan("BIKE-A")
        self._scan("BIKE-A")
        self._scan("T2002")
        response = self._scan("BIKE-B")
        state = response.json()["state"]
        self.assertEqual(state["current_total"], 3)
        self.assertEqual(state["pallet_total"], 1)
        self.assertEqual(state["expected_total"], 5)
        self.assertEqual(state["completion"], 60)

    def test_scanner_records_feed_existing_first_scan_data(self):
        from receiving.views import _stage2_data
        session = self._start()
        self._scan("T4324")
        self._scan("BIKE-A")
        received_batch, moves_import, reconciliation = _stage2_data(self.container)
        self.assertEqual(received_batch.pk, session.batch_id)
        self.assertEqual(received_batch.lines.get().code, "BIKE-A")
        self.assertIsNone(moves_import)
        self.assertIsNone(reconciliation)

    def test_existing_manual_first_scan_batch_is_reused(self):
        source = SourceFile.objects.create(
            container=self.container,
            kind="RECEIVED",
            file=SimpleUploadedFile("manual.xlsx", b"manual"),
            original_name="manual.xlsx",
            sha256="2" * 64,
            status="CONFIRMED",
            uploaded_by=self.user,
        )
        manual_batch = ImportBatch.objects.create(
            source_file=source,
            status="CONFIRMED",
            total_rows=1,
            total_units=1,
            confirmed_at=timezone.now(),
            created_by=self.user,
        )
        NormalizedLine.objects.create(
            batch=manual_batch,
            source_sheet="Manual",
            source_row=1,
            code="BIKE-A",
            long_code="LONG-A-0001",
            quantity=1,
            location="T1001",
        )
        session = self._start()
        self.assertEqual(session.batch_id, manual_batch.pk)
        self._scan("T1002")
        self._scan("BIKE-B")
        manual_batch.refresh_from_db()
        self.assertEqual(manual_batch.total_units, 2)
        self.assertEqual(manual_batch.lines.count(), 2)

    def test_unknown_barcode_is_rejected_without_changing_counts(self):
        from receiving.models import FirstScanEvent
        session = self._start()
        self._scan("T4324")
        response = self._scan("UNKNOWN-999")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["ok"])
        self.assertEqual(response.json()["level"], "error")
        session.batch.refresh_from_db()
        self.assertEqual(session.batch.total_units, 0)
        self.assertEqual(session.batch.lines.count(), 0)
        self.assertEqual(FirstScanEvent.objects.get(event_type="BIKE").result, "ERROR")

    def test_photo_links_container_pallet_and_latest_scan_without_changing_counts(self):
        from receiving.models import FirstScanPhoto
        session = self._start()
        self._scan("T4324")
        self._scan("BIKE-A")
        response = self.client.post(
            self.url,
            {
                "action": "photo",
                "photo": SimpleUploadedFile("label.jpg", b"\xff\xd8\xffevidence", content_type="image/jpeg"),
            },
        )
        self.assertEqual(response.status_code, 302)
        photo = FirstScanPhoto.objects.get()
        self.assertEqual(photo.session.container, self.container)
        self.assertEqual(photo.pallet, "T4324")
        self.assertEqual(photo.event.resolved_code, "BIKE-A")
        session.batch.refresh_from_db()
        self.assertEqual(session.batch.total_units, 1)

    def test_container_first_scan_section_has_scan_button(self):
        response = self.client.get(reverse("receiving:container_detail", args=[self.container.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.url)
        self.assertContains(response, ">Scan</a>")

    def test_scanner_screen_renders_start_and_live_views(self):
        start_page = self.client.get(self.url)
        self.assertEqual(start_page.status_code, 200)
        self.assertContains(start_page, "Start continuous scanning")
        self.assertContains(start_page, "Current scanned total")
        self._start()
        live_page = self.client.get(self.url)
        self.assertEqual(live_page.status_code, 200)
        self.assertContains(live_page, 'id="continuous-scan-form"')
        self.assertContains(live_page, "Pause Scan")
        self.assertContains(live_page, "Take Photo")
