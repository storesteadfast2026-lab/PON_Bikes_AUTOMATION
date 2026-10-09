import tempfile
from datetime import date
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from .models import Container, Customer, FirstScanEvent, ImportBatch, NormalizedLine, SourceFile
from .services import build_client_report_rows, compare_lines, parse_workbook


ANNOUNCED_LONG_CODE = "68-25281-4-849-9999"
UNADVISED_LONG_CODE = "68-25281-5-849-9999-WT"
SMALL_PART_CODE = "68-29999-1-999-9999"
SPARE_PART_CODE = "68-29998-1-999-9999"


class ManifestTwoCodeParsingTests(TestCase):
    def test_sections_repeated_long_codes_and_original_rows_are_preserved(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Packing List"
        sheet.append([])
        for column, value in zip((2, 3, 4), ("Santa Cruz - Article Number", "Name / Description", "Quantity")):
            sheet.cell(18, column, value)
        for column in range(2, 5):
            sheet.cell(19, column, "FRAMES")
        sheet.cell(20, 2, ANNOUNCED_LONG_CODE)
        sheet.cell(20, 3, "Nomad frame")
        sheet.cell(20, 4, 2)
        sheet.cell(21, 2, ANNOUNCED_LONG_CODE)
        sheet.cell(21, 3, "Nomad frame")
        sheet.cell(21, 4, 1)
        for column in range(2, 5):
            sheet.cell(22, column, "SMALL PARTS")
        sheet.cell(23, 2, SMALL_PART_CODE)
        sheet.cell(23, 3, "Rocker Link Bullit 4")
        sheet.cell(23, 4, 2)

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "TIIU4062052.xlsx"
            workbook.save(path)
            rows, errors = parse_workbook(path, {
                "sheet_name": "Packing List",
                "start_row": 20,
                "row_step": 1,
                "code_column": "B",
                "long_code_column": "B",
                "description_column": "C",
                "quantity_column": "D",
            })

        self.assertEqual(errors, [])
        self.assertEqual([row["source_row"] for row in rows], [20, 21, 23])
        self.assertEqual([row["quantity"] for row in rows], [2, 1, 2])
        self.assertEqual([row["raw_data"]["_manifest_section"] for row in rows], ["FRAMES", "FRAMES", "SMALL PARTS"])
        self.assertEqual(rows[0]["raw_data"]["B"], ANNOUNCED_LONG_CODE)


class FirstScanTwoCodesTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.settings = override_settings(MEDIA_ROOT=self.media.name)
        self.settings.enable()
        self.user = get_user_model().objects.create_user(username="two-codes", password="test-password")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="TIIU4062052",
            customer=self.customer,
            container_date=date(2026, 7, 7),
            year=2026,
            job_order="CON20",
            created_by=self.user,
        )
        source = SourceFile.objects.create(
            container=self.container,
            kind="CLIENT",
            file=SimpleUploadedFile("TIIU4062052.xlsx", b"manifest"),
            original_name="TIIU4062052.xlsx",
            sha256="a" * 64,
            status="CONFIRMED",
            uploaded_by=self.user,
        )
        self.manifest = ImportBatch.objects.create(
            source_file=source,
            status="CONFIRMED",
            mapping={"code_column": "B", "long_code_column": "B", "quantity_column": "D"},
            total_rows=4,
            total_units=11,
            confirmed_at=timezone.now(),
            created_by=self.user,
        )
        NormalizedLine.objects.create(
            batch=self.manifest,
            source_sheet="Packing List",
            source_row=76,
            code=ANNOUNCED_LONG_CODE,
            long_code=ANNOUNCED_LONG_CODE,
            description="Nmd 6 CC MX 25 LG PUR NS",
            quantity=2,
            raw_data={"B": ANNOUNCED_LONG_CODE, "D": "2", "_manifest_section": "FRAMES"},
        )
        NormalizedLine.objects.create(
            batch=self.manifest,
            source_sheet="Packing List",
            source_row=77,
            code=ANNOUNCED_LONG_CODE,
            long_code=ANNOUNCED_LONG_CODE,
            description="Nmd 6 CC MX 25 LG PUR NS",
            quantity=1,
            raw_data={"B": ANNOUNCED_LONG_CODE, "D": "1", "_manifest_section": "FRAMES"},
        )
        NormalizedLine.objects.create(
            batch=self.manifest,
            source_sheet="Packing List",
            source_row=91,
            code=SMALL_PART_CODE,
            long_code=SMALL_PART_CODE,
            description="Rocker Link Bullit 4",
            quantity=2,
            raw_data={"B": SMALL_PART_CODE, "D": "2", "_manifest_section": "SMALL PARTS"},
        )
        NormalizedLine.objects.create(
            batch=self.manifest,
            source_sheet="Packing List",
            source_row=92,
            code=SPARE_PART_CODE,
            long_code=SPARE_PART_CODE,
            description="Santa Cruz spare part",
            quantity=6,
            raw_data={"B": SPARE_PART_CODE, "D": "6", "_manifest_section": "SPARE PARTS"},
        )
        self.url = reverse("receiving:first_scan_scanner", args=[self.container.pk])

    def tearDown(self):
        self.settings.disable()
        self.media.cleanup()

    def start(self):
        response = self.client.post(self.url, {"action": "start", "mode": "AUTO", "code_mode": "TWO_CODES"})
        self.assertEqual(response.status_code, 302)
        self.container.refresh_from_db()
        self.assertEqual(self.container.first_scan_code_mode, "TWO_CODES")
        return self.container.first_scan_sessions.get(status="ACTIVE")

    def scan(self, value):
        return self.client.post(
            self.url,
            {"action": "scan", "scanned_value": value},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

    def complete_bike(self, main_code, long_code=ANNOUNCED_LONG_CODE):
        first = self.scan(main_code)
        second = self.scan(long_code)
        return first, second

    def test_main_code_is_discovered_during_scan_and_bike_counts_only_after_long_code(self):
        session = self.start()
        self.scan("T4191")
        first = self.scan("192219461014")
        session.batch.refresh_from_db()
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.json()["state"]["pending_main_code"], "192219461014")
        self.assertEqual(session.batch.total_units, 0)
        second = self.scan(ANNOUNCED_LONG_CODE)
        session.batch.refresh_from_db()
        line = session.batch.lines.get()
        self.assertEqual(second.status_code, 200)
        self.assertEqual(session.batch.total_units, 1)
        self.assertEqual((line.code, line.long_code, line.location), ("192219461014", ANNOUNCED_LONG_CODE, "T4191"))
        self.assertEqual(line.raw_data["scanner_code_mode"], "TWO_CODES")

    def test_repeated_manifest_rows_and_quantity_greater_than_one_are_aggregated(self):
        session = self.start()
        self.scan("T4191")
        for _ in range(3):
            self.complete_bike("192219461014")
        rows = compare_lines(self.manifest.lines.all(), session.batch.lines.all(), two_codes=True)
        row = next(item for item in rows if item["code"] == ANNOUNCED_LONG_CODE)
        self.assertEqual((row["expected"], row["received"], row["status"]), (3, 3, "MATCH"))
        self.assertEqual(row["source_rows"], [76, 77])
        self.assertEqual(row["sections"], ["FRAMES"])
        self.assertEqual(self.manifest.lines.count(), 4)

    def test_multiple_bikes_can_share_one_long_code(self):
        session = self.start()
        self.scan("T4191")
        self.complete_bike("192219461014")
        self.complete_bike("192219461014")
        self.assertEqual(session.batch.lines.filter(long_code=ANNOUNCED_LONG_CODE).count(), 2)

    def test_different_physical_codes_for_one_long_code_are_kept_with_review_warning(self):
        session = self.start()
        self.scan("T4191")
        self.complete_bike("192219461014")
        _, response = self.complete_bike("192219461015")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["level"], "warning")
        self.assertEqual(
            set(session.batch.lines.filter(long_code=ANNOUNCED_LONG_CODE).values_list("code", flat=True)),
            {"192219461014", "192219461015"},
        )
        row = next(item for item in compare_lines(self.manifest.lines.all(), session.batch.lines.all()) if item["code"] == ANNOUNCED_LONG_CODE)
        self.assertEqual(row["status"], "REVIEW")
        self.assertIn("multiple physical Codes", row["warning"])

    def test_unadvised_real_bike_is_registered_as_advised_zero_received_one(self):
        session = self.start()
        self.scan("T4200")
        _, response = self.complete_bike("192219543345", UNADVISED_LONG_CODE)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["level"], "warning")
        line = session.batch.lines.get()
        self.assertEqual((line.code, line.long_code), ("192219543345", UNADVISED_LONG_CODE))
        self.assertTrue(line.raw_data["unadvised"])
        report = build_client_report_rows(self.manifest.lines.all(), session.batch.lines.all())
        row = next(item for item in report if item["code"] == UNADVISED_LONG_CODE)
        self.assertEqual((row["advised"], row["received"], row["variance"]), (0, 1, -1))
        self.assertEqual(row["comment"], "Unadvised / Not Advised")
        comparison = compare_lines(self.manifest.lines.all(), session.batch.lines.all())
        self.assertEqual(next(item for item in comparison if item["code"] == UNADVISED_LONG_CODE)["status"], "UNADVISED")

    def test_small_parts_and_spare_parts_are_not_counted_or_reconciled_as_bikes(self):
        self.start()
        response = self.scan("T4191")
        state = response.json()["state"]
        self.assertEqual(state["expected_total"], 3)
        session = self.container.first_scan_sessions.get(status="ACTIVE")
        rows = compare_lines(self.manifest.lines.all(), session.batch.lines.all(), two_codes=True)
        compared_codes = {row["code"] for row in rows}
        self.assertNotIn(SMALL_PART_CODE, compared_codes)
        self.assertNotIn(SPARE_PART_CODE, compared_codes)

    def test_non_scannable_spare_part_long_code_cannot_complete_a_bike(self):
        session = self.start()
        self.scan("T4191")
        first = self.scan("192219461014")
        self.assertEqual(first.status_code, 200)
        second = self.scan(SPARE_PART_CODE)
        self.assertEqual(second.status_code, 400)
        self.assertIn("non-scannable manifest section SPARE PARTS", second.json()["message"])
        session.batch.refresh_from_db()
        self.assertEqual(session.batch.total_units, 0)
        self.assertEqual(session.batch.lines.count(), 0)
        self.assertEqual(second.json()["state"]["pending_main_code"], "192219461014")

    def test_over_received_is_kept_and_warned_instead_of_rejected(self):
        session = self.start()
        self.scan("T4191")
        for _ in range(3):
            self.complete_bike("192219461014")
        _, response = self.complete_bike("192219461014")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["level"], "warning")
        self.assertEqual(session.batch.lines.count(), 4)
        row = next(item for item in compare_lines(self.manifest.lines.all(), session.batch.lines.all()) if item["code"] == ANNOUNCED_LONG_CODE)
        self.assertEqual((row["expected"], row["received"], row["status"]), (3, 4, "OVER"))

    def test_pallet_scan_does_not_break_pending_pair(self):
        session = self.start()
        self.scan("T4191")
        self.scan("192219461014")
        self.scan("T4192")
        self.scan(ANNOUNCED_LONG_CODE)
        line = session.batch.lines.get()
        self.assertEqual(line.location, "T4191")
        session.refresh_from_db()
        self.assertEqual(session.current_pallet, "T4192")


    def test_pallet_summary_reports_unique_count_first_last_and_live_position(self):
        self.start()
        first = self.scan("T4191")
        self.assertEqual(first.json()["state"]["pallet_count"], 1)
        self.assertEqual(first.json()["state"]["pallet_position"], 1)
        self.assertIsNone(first.json()["state"]["expected_pallet_count"])
        self.assertEqual(first.json()["state"]["first_pallet"], "T4191")
        self.assertEqual(first.json()["state"]["last_pallet"], "T4191")
        second = self.scan("T4192")
        self.assertEqual(second.json()["state"]["pallet_count"], 2)
        self.assertEqual(second.json()["state"]["pallet_position"], 2)
        self.assertIsNone(second.json()["state"]["expected_pallet_count"])
        self.assertEqual(second.json()["state"]["first_pallet"], "T4191")
        self.assertEqual(second.json()["state"]["last_pallet"], "T4192")

    def test_confirmed_scanned_file_can_supply_known_total_pallet_count(self):
        source = SourceFile.objects.create(
            container=self.container,
            kind="RECEIVED",
            file=SimpleUploadedFile("scanned.xlsx", b"scanned"),
            original_name="scanned.xlsx",
            sha256="c" * 64,
            status="CONFIRMED",
            uploaded_by=self.user,
        )
        received = ImportBatch.objects.create(
            source_file=source,
            status="CONFIRMED",
            mapping={"location_stream": True},
            total_rows=3,
            total_units=3,
            confirmed_at=timezone.now(),
            created_by=self.user,
        )
        for row_number, pallet in enumerate(("T4191", "T4192", "T4193"), 1):
            NormalizedLine.objects.create(
                batch=received, source_sheet="Scanner", source_row=row_number,
                code=f"19221946101{row_number}", long_code=ANNOUNCED_LONG_CODE,
                quantity=1, location=pallet,
            )
        session = self.start()
        self.assertEqual(session.batch_id, received.pk)
        response = self.scan("T4192")
        self.assertEqual(response.json()["state"]["pallet_position"], 1)
        self.assertEqual(response.json()["state"]["expected_pallet_count"], 3)

    def test_recent_history_spans_pallets_and_keeps_two_scanned_codes(self):
        self.start()
        self.scan("T4191")
        self.complete_bike("192219461014")
        self.scan("T4192")
        _, response = self.complete_bike("192219461015")
        recent = response.json()["state"]["recent"]
        self.assertEqual([item["pallet"] for item in recent[:2]], ["T4192", "T4191"])
        self.assertEqual(recent[0]["resolved_code"], "192219461015")
        self.assertEqual(recent[0]["long_code"], ANNOUNCED_LONG_CODE)

    def test_live_scanner_page_is_minimal_and_history_starts_with_pallet(self):
        self.start()
        self.scan("T4191")
        self.complete_bike("192219461014")
        response = self.client.get(self.url)
        html = response.content.decode()
        self.assertNotIn("Last result", html)
        self.assertNotIn('id="container-progress"', html)
        self.assertNotIn('id="pallet-progress"', html)
        self.assertNotIn('class="scanner-summary-grid"', html)
        self.assertIn("Current pallet:", html)
        self.assertIn("Recent scans", html)
        self.assertIn("scanner-history-code", html)
        history_start = html.index('<tbody id="recent-scans-body">')
        history_end = html.index('</tbody>', history_start)
        history_html = html[history_start:history_end]
        self.assertLess(history_html.index("T4191"), history_html.index("192219461014"))
        self.assertLess(history_html.index("192219461014"), history_html.index(ANNOUNCED_LONG_CODE))

    def test_one_code_mode_keeps_existing_registration_logic(self):
        other = Container.objects.create(
            identifier="ONECODE001",
            customer=self.customer,
            year=2026,
            job_order="ONE-CODE",
            created_by=self.user,
        )
        source = SourceFile.objects.create(
            container=other,
            kind="CLIENT",
            file=SimpleUploadedFile("one.xlsx", b"one"),
            original_name="one.xlsx",
            sha256="b" * 64,
            status="CONFIRMED",
            uploaded_by=self.user,
        )
        batch = ImportBatch.objects.create(
            source_file=source,
            status="CONFIRMED",
            total_rows=1,
            total_units=1,
            confirmed_at=timezone.now(),
            created_by=self.user,
        )
        NormalizedLine.objects.create(batch=batch, source_sheet="Manifest", source_row=1, code="BIKE-A", quantity=1)
        url = reverse("receiving:first_scan_scanner", args=[other.pk])
        self.client.post(url, {"action": "start", "mode": "AUTO", "code_mode": "ONE_CODE"})
        self.client.post(url, {"action": "scan", "scanned_value": "T1001"}, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        response = self.client.post(url, {"action": "scan", "scanned_value": "BIKE-A"}, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 200)
        session = other.first_scan_sessions.get(status="ACTIVE")
        self.assertEqual(session.batch.lines.get().code, "BIKE-A")
        self.assertFalse(FirstScanEvent.objects.filter(session=session, message__startswith="Waiting for Long Code").exists())
