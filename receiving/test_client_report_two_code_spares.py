from io import BytesIO
from types import SimpleNamespace

from django.test import SimpleTestCase
from openpyxl import load_workbook

from receiving.services import (
    build_client_receiving_workbook,
    build_client_report_rows,
    resolve_client_report_two_codes,
)


class ClientReportTwoCodeSparePartsTests(SimpleTestCase):
    LONG_A = "58-27322-351-3-941-232702"
    LONG_B = "68-25312-4-865-9999"

    def client_line(self, code, qty, description, section="PEDAL 58"):
        return SimpleNamespace(
            code=code,
            long_code="",
            quantity=qty,
            description=description,
            raw_data={"_manifest_section": section},
        )

    def received_line(self, physical, long_code, qty=1):
        return SimpleNamespace(
            code=physical,
            long_code=long_code,
            quantity=qty,
            description="",
            raw_data={},
        )

    def test_detects_long_code_reconciliation_when_stored_mode_is_blank(self):
        client = [
            self.client_line(self.LONG_A, 2, "Bike A"),
            self.client_line(self.LONG_B, 1, "Bike B"),
        ]
        received = [
            self.received_line("192219530376", self.LONG_A),
            self.received_line("192219530377", self.LONG_A),
            self.received_line("192219530500", self.LONG_B),
        ]
        self.assertTrue(resolve_client_report_two_codes(client, received, configured_two_codes=False))

    def test_one_code_flow_is_not_changed_when_code_overlap_is_better(self):
        client = [self.client_line("BIKE-A", 1, "Bike A", section="")]
        received = [SimpleNamespace(code="BIKE-A", long_code="", quantity=1, description="", raw_data={})]
        self.assertFalse(resolve_client_report_two_codes(client, received, configured_two_codes=False))

    def test_received_quantities_reconcile_by_long_code_and_spare_part_is_separate(self):
        client = [
            self.client_line(self.LONG_A, 2, "Nmd 7 CC MX 27 MD EARTH X0 AXS-C RSV"),
            self.client_line("04-27492", 2, "Rocker Link Bullit 4", section="SMALL PARTS"),
        ]
        received = [
            self.received_line("192219530376", self.LONG_A),
            self.received_line("192219530377", self.LONG_A),
        ]
        rows = build_client_report_rows(
            client,
            received,
            catalog_entries=[],
            product_definitions=[],
            two_codes=True,
        )
        bike = rows[0]
        spare = rows[1]
        self.assertEqual(bike["code"], self.LONG_A)
        self.assertEqual(bike["advised"], 2)
        self.assertEqual(bike["received"], 2)
        self.assertEqual(bike["variance"], 0)
        self.assertFalse(bike["is_spare_part"])

        self.assertTrue(spare["is_spare_part"])
        self.assertEqual(spare["section_name"], "SMALL PARTS")
        self.assertEqual(spare["item_name"], "Rocker Link Bullit 4")
        self.assertEqual(spare["advised"], 2)
        self.assertIsNone(spare["received"])
        self.assertIsNone(spare["variance"])

    def test_workbook_has_separate_small_parts_section_with_client_name(self):
        rows = [
            {
                "code": self.LONG_A,
                "item_name": "Bike A",
                "advised": 2,
                "received": 2,
                "variance": 0,
                "comment": "",
                "is_new": False,
                "is_pbp_to_pon": False,
                "is_spare_part": False,
                "section_name": "",
            },
            {
                "code": "04-27492",
                "item_name": "Rocker Link Bullit 4",
                "advised": 2,
                "received": None,
                "variance": None,
                "comment": "Not part of bike First Scan",
                "is_new": False,
                "is_pbp_to_pon": False,
                "is_spare_part": True,
                "section_name": "SMALL PARTS",
            },
        ]
        wb = load_workbook(build_client_receiving_workbook("TIIU4062052", rows), data_only=True)
        ws = wb["Receiving Container Data"]

        self.assertEqual(ws["C4"].value, "Bike Total")
        self.assertEqual(ws["B5"].value, "SMALL PARTS")
        self.assertEqual(ws["B6"].value, "04-27492")
        self.assertEqual(ws["C6"].value, "Rocker Link Bullit 4")
        self.assertEqual(ws["D6"].value, 2)
        self.assertIsNone(ws["E6"].value)
        self.assertIsNone(ws["F6"].value)
        self.assertEqual(ws["C7"].value, "Spare Parts Total")
