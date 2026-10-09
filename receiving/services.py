import csv
import hashlib
import io
import re
import subprocess
import tempfile
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_UP
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from PIL import Image, ImageDraw
from reportlab.graphics.barcode.code128 import Code128


LOCATION_PATTERN = re.compile(r"^T\d{4}$", re.IGNORECASE)
CONTAINER_PATTERN = re.compile(r"^[A-Z]{4}\d{7}$")
PRODUCT_CODE_PATTERN = re.compile(r"^(?=.*\d)[A-Z0-9][A-Z0-9._/-]{5,19}$", re.IGNORECASE)
PON_LONG_CODE_PATTERN = re.compile(
    r"^(?:58-\d{5}-\d{3}-\d-\d{3}-\d{6}(?:-\d{2})?|68-\d{5}-\d-\d{3}-\d{4}(?:-WT)?)$",
    re.IGNORECASE,
)
PRODUCT_CATALOG_PATTERN = re.compile(r"^.+\.(?:xls|xml)$", re.IGNORECASE)
PRODUCT_CATALOG_HEADERS = ("code", "pon_sku", "customer", "code2", "short_name", "long_name", "group1", "group2")
PRODUCT_CATALOG_OPTIONAL_FIELDS = ("cubic", "weight", "height", "width", "length")
HEADER_ALIASES = {
    "code_column": {
        "sku", "code", "product", "product code", "bike code", "item", "item code",
        "item no", "item number", "no", "no.", "article", "article no", "product no",
    },
    "long_code_column": {"longcode", "long code", "long sku", "long_sku", "long item code"},
    "description_column": {"description", "product description", "item description", "full name", "name"},
    "quantity_column": {"quantity", "qty", "qty.", "units", "unit qty", "quantity base", "base quantity"},
    "container_column": {"container", "container id", "container no", "container number", "container #"},
}


NEW_PRODUCT_HEADERS = [
    "Code", "long_sku", "Name", "Customer", "Pick_loc", "Prod_Ref", "outer_bar", "Status", "Group",
    "Group1", "Group2", "man_id", "Full_Name", "Buy", "Class", "Quantity", "Unit", "Pallet", "Lift",
    "Cubic", "Weight", "Height", "Width", "Length", "sub_risk", "Inner", "Outer", "Layer", "Comment", "Date",
]

EXPORT_DEFAULTS = {
    "TL_EXPORT_FORMAT": "excel_5_95",
    "TL_CUSTOMER": "PON",
    "TL_STATUS": "L",
    "TL_GROUP2_CUBIC_THRESHOLD": "0.4",
    "TL_GROUP2_LARGE": "CHG02",
    "TL_GROUP2_DEFAULT": "",
    "TL_QUANTITY": "1",
    "TL_PALLET": "0",
    "TL_LIFT": "1",
    "TL_OUTER": "1",
}


def is_pon_long_code(value):
    return bool(PON_LONG_CODE_PATTERN.fullmatch(str(value or "").strip().upper()))


NON_SCANNABLE_MANIFEST_SECTIONS = {"SMALL PARTS", "SPARE PARTS"}


def manifest_section(line):
    if isinstance(line, dict):
        raw_data = line.get("raw_data", {}) or {}
    else:
        raw_data = getattr(line, "raw_data", {}) or {}
    return str(raw_data.get("_manifest_section", "") or "").strip()


def normalise_manifest_section(value):
    return re.sub(r"[\s_-]+", " ", str(value or "").strip().upper()).strip()


def is_non_scannable_manifest_section(value):
    return normalise_manifest_section(value) in NON_SCANNABLE_MANIFEST_SECTIONS


def is_non_scannable_manifest_line(line):
    return is_non_scannable_manifest_section(manifest_section(line))


def summarize_import_rows(rows):
    """Return physical-bike and spare-part quantities without dropping manifest rows."""
    summary = {
        "data_rows": 0,
        "bike_rows": 0,
        "bike_units": 0,
        "spare_rows": 0,
        "spare_units": 0,
    }
    for row in rows:
        quantity = int((row.get("quantity", 0) if isinstance(row, dict) else getattr(row, "quantity", 0)) or 0)
        summary["data_rows"] += 1
        if is_non_scannable_manifest_line(row):
            summary["spare_rows"] += 1
            summary["spare_units"] += quantity
        else:
            summary["bike_rows"] += 1
            summary["bike_units"] += quantity
    return summary


def _line_reconciliation_code(line, use_long_code=False):
    code = str(getattr(line, "code", "") or "").strip().upper()
    long_code = str(getattr(line, "long_code", "") or "").strip().upper()
    return (long_code or code) if use_long_code else code


def _uses_two_code_reconciliation(received_lines):
    for line in received_lines:
        raw_data = getattr(line, "raw_data", {}) or {}
        if str(raw_data.get("scanner_code_mode", "") or "").upper() == "TWO_CODES":
            return True
    return False


def first_scan_expected_units(client_lines, two_codes=False):
    scannable_lines = [line for line in client_lines if not is_non_scannable_manifest_line(line)]
    if not two_codes:
        return sum(int(getattr(line, "quantity", 0) or 0) for line in scannable_lines)
    return sum(
        int(getattr(line, "quantity", 0) or 0)
        for line in scannable_lines
        if is_pon_long_code(_line_reconciliation_code(line, use_long_code=True))
    )

SHORT_NAME_REPLACEMENTS = [
    ("FRAMESET", "FR SET"),
    ("FRAME", "FR"),
    ("BICYCLE", "BIKE"),
    ("ELECTRIC", "ELEC"),
    ("MOUNTAIN", "MTB"),
    ("OFFSET", "OFF"),
    ("CLOUD", "CLD"),
    ("BREAK", "BRK"),
    ("CARBON", "CRBN"),
    ("ALUMINIUM", "ALU"),
    ("ALUMINUM", "ALU"),
    ("LIMITED", "LTD"),
    ("EDITION", "ED"),
    ("MATTE", "MT"),
    ("GLOSS", "GLS"),
    ("BLACK", "BLK"),
    ("WHITE", "WHT"),
    ("GREEN", "GRN"),
    ("YELLOW", "YLW"),
    ("ORANGE", "ORG"),
    ("PURPLE", "PPL"),
    ("SILVER", "SLV"),
    ("MAGIC", "MGC"),
]


def available_server_files(root_path, container_identifier):
    """Return safe Excel candidates without exposing files outside the configured root."""
    root = Path(root_path)
    if not root.exists() or not root.is_dir():
        return [], False
    resolved_root = root.resolve()
    candidates = []
    for path in resolved_root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        try:
            relative = path.resolve().relative_to(resolved_root)
        except ValueError:
            continue
        name_match = container_identifier.upper() in path.name.upper()
        parent_match = container_identifier.upper() in str(relative.parent).upper()
        candidates.append(
            {
                "relative_path": relative.as_posix(),
                "name": path.name,
                "size": path.stat().st_size,
                "container_match": name_match or parent_match,
            }
        )
    candidates.sort(key=lambda item: (not item["container_match"], item["relative_path"].lower()))
    return candidates[:250], True


def resolve_server_file(root_path, relative_path):
    root = Path(root_path).resolve(strict=True)
    candidate = (root / relative_path).resolve(strict=True)
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ValueError("The selected file is outside the configured server folder.") from exc
    if not candidate.is_file() or candidate.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("The selected server file is not an accepted Excel workbook.")
    return candidate


def available_product_catalog_files(root_path):
    """Return supported Translogic product master files (.xls primary; .xml legacy)."""
    root = Path(root_path)
    if not root.exists() or not root.is_dir():
        return [], False
    resolved_root = root.resolve()
    candidates = []
    for path in resolved_root.rglob("*"):
        if not path.is_file() or not PRODUCT_CATALOG_PATTERN.fullmatch(path.name):
            continue
        try:
            relative = path.resolve().relative_to(resolved_root)
        except ValueError:
            continue
        candidates.append(
            {
                "relative_path": relative.as_posix(),
                "name": path.name,
                "size": path.stat().st_size,
                "modified_at": path.stat().st_mtime,
            }
        )
    candidates.sort(key=lambda item: (-item["modified_at"], item["relative_path"].lower()))
    return candidates[:100], True


def resolve_product_catalog_file(root_path, relative_path):
    root = Path(root_path).resolve(strict=True)
    candidate = (root / relative_path).resolve(strict=True)
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ValueError("The selected product file is outside the configured server folder.") from exc
    if not candidate.is_file() or not PRODUCT_CATALOG_PATTERN.fullmatch(candidate.name):
        raise ValueError("Select a supported product master file (.xls).")
    return candidate


def file_sha256(uploaded_file):
    digest = hashlib.sha256()
    for chunk in uploaded_file.chunks():
        digest.update(chunk)
    uploaded_file.seek(0)
    return digest.hexdigest()


def path_sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalise_catalog_xml_bytes(data):
    """Legacy XML compatibility for catalogues imported before the XLS switch."""
    return data.replace(b"\xc2\xa0", b" ").replace(b"\xa0", b" ")


def _parse_product_catalog_xml(path):
    """Read the former XML product master for backwards compatibility."""
    data = _normalise_catalog_xml_bytes(Path(path).read_bytes())
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError as exc:
        raise ValueError(f"The legacy product XML is not valid: {exc}.") from exc

    table = root.find("table") if root.tag == "xdoc" else root.find("./table")
    if table is None:
        raise ValueError("The legacy product XML does not contain the expected <table> element.")

    rows = []
    for source_row, row in enumerate(table.findall("row"), start=1):
        raw = {child.tag: _clean_text(child.text) for child in row}
        missing = [name for name in PRODUCT_CATALOG_HEADERS if name not in raw]
        if missing:
            raise ValueError(
                f"Legacy product row {source_row} is missing required field(s): {', '.join(missing)}."
            )
        rows.append(_normalise_product_catalog_row(raw, source_row))

    rows = [row for row in rows if row is not None]
    if not rows:
        raise ValueError("The legacy product XML does not contain usable product rows.")
    return rows


def _catalog_header_key(value):
    value = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    aliases = {
        "code": "code",
        "product_code": "code",
        "product": "code",
        "pon_sku": "pon_sku",
        "ponsku": "pon_sku",
        "customer": "customer",
        "code2": "code2",
        "code_2": "code2",
        "short_name": "short_name",
        "shortname": "short_name",
        "long_name": "long_name",
        "longname": "long_name",
        "group1": "group1",
        "group_1": "group1",
        "group2": "group2",
        "group_2": "group2",
        "cubic": "cubic",
        "weight": "weight",
        "kg": "weight",
        "height": "height",
        "width": "width",
        "length": "length",
        "lenght": "length",
    }
    return aliases.get(value, value)


def _excel_catalog_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\xa0", " ").strip()


def _normalise_product_catalog_row(raw, source_row):
    normalized = {
        "code": _excel_catalog_text(raw.get("code", "")).upper(),
        "pon_sku": _excel_catalog_text(raw.get("pon_sku", "")).upper(),
        "customer": _excel_catalog_text(raw.get("customer", "")).upper(),
        "code2": _excel_catalog_text(raw.get("code2", "")).upper(),
        "short_name": re.sub(r"\s+", " ", _excel_catalog_text(raw.get("short_name", ""))).strip(),
        "long_name": re.sub(r"\s+", " ", _excel_catalog_text(raw.get("long_name", ""))).strip(),
        "group1": _excel_catalog_text(raw.get("group1", "")).upper(),
        "group2": _excel_catalog_text(raw.get("group2", "")).upper(),
    }
    if not any(normalized[name] for name in ("code", "pon_sku", "code2")):
        return None
    if len(normalized["short_name"]) > 30:
        raise ValueError(
            f"Product master row {source_row} has a short_name longer than 30 characters: "
            f"{normalized['short_name']}"
        )
    normalized_raw = {_catalog_header_key(k): _excel_catalog_text(v) for k, v in raw.items()}
    return {"source_row": source_row, **normalized, "raw_data": normalized_raw}


def _convert_xls_catalog_to_xlsx(path, destination):
    """Convert legacy BIFF .xls to temporary .xlsx using Gnumeric/ssconvert."""
    command = ["ssconvert", str(path), str(destination)]
    try:
        completed = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
    except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
        raise RuntimeError(
            "Reading the .xls product master requires Gnumeric/ssconvert in the application container."
        ) from exc
    if completed.returncode or not Path(destination).is_file():
        detail = (completed.stderr or completed.stdout or "unknown conversion error").strip()
        raise RuntimeError(f"Could not read the .xls product master: {detail}")


def _parse_product_catalog_xls(path):
    """Read products_pon_pbp_auto.xls using the expected product-master columns."""
    with tempfile.TemporaryDirectory(prefix="pon_product_catalog_") as folder:
        converted = Path(folder) / "product_catalog.xlsx"
        _convert_xls_catalog_to_xlsx(Path(path), converted)
        workbook = load_workbook(converted, data_only=True, read_only=True)
        try:
            selected = None
            header_row = None
            header_map = None
            required = set(PRODUCT_CATALOG_HEADERS)
            supported = required | set(PRODUCT_CATALOG_OPTIONAL_FIELDS)
            for worksheet in workbook.worksheets:
                for row_number, values in enumerate(worksheet.iter_rows(min_row=1, max_row=25, values_only=True), start=1):
                    mapped = {}
                    for index, value in enumerate(values):
                        key = _catalog_header_key(value)
                        if key in supported and key not in mapped:
                            mapped[key] = index
                    if required.issubset(mapped):
                        selected = worksheet
                        header_row = row_number
                        header_map = mapped
                        break
                if selected is not None:
                    break

            if selected is None:
                raise ValueError(
                    "The .xls product master does not contain the required columns: "
                    + ", ".join(PRODUCT_CATALOG_HEADERS)
                )

            rows = []
            for source_row, values in enumerate(
                selected.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                raw = {name: (values[index] if index < len(values) else None) for name, index in header_map.items()}
                row = _normalise_product_catalog_row(raw, source_row)
                if row is not None:
                    rows.append(row)
            if not rows:
                raise ValueError("The .xls product master does not contain usable product rows.")
            return rows
        finally:
            workbook.close()


def parse_product_catalog(path):
    """Read the configured product master. XLS is primary; XML remains legacy-compatible."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return _parse_product_catalog_xls(path)
    if suffix == ".xml":
        return _parse_product_catalog_xml(path)
    raise ValueError("The product master must be an Excel .xls file.")

def _normalise_name_key(value):
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ").strip()).upper()


def _name_tokens(value):
    return re.findall(r"[A-Za-z0-9]+(?:[./*+\-][A-Za-z0-9]+)*|[^\s]", str(value or ""))


def build_name_dictionary(catalog_entries):
    """Learn conservative long-name -> short-name matches and token abbreviations.

    Exact historical matches are preferred. Token rules are only accepted when the XML
    shows the same shortening at least twice with >=75% agreement, which prevents one-off
    or contradictory product descriptions from becoming global rules.
    """
    exact_candidates = defaultdict(Counter)
    token_candidates = defaultdict(Counter)

    for entry in catalog_entries:
        if isinstance(entry, dict):
            long_raw = entry.get("long_name", "")
            short_raw = entry.get("short_name", "")
        else:
            long_raw = getattr(entry, "long_name", "")
            short_raw = getattr(entry, "short_name", "")
        long_name = re.sub(r"\s+", " ", str(long_raw or "").replace("\xa0", " ").strip())
        short_name = re.sub(r"\s+", " ", str(short_raw or "").replace("\xa0", " ").strip())
        if not long_name or not short_name or len(short_name) > 30:
            continue
        exact_candidates[_normalise_name_key(long_name)][short_name] += 1

        long_tokens = _name_tokens(long_name)
        short_tokens = _name_tokens(short_name)
        if len(long_tokens) != len(short_tokens):
            continue
        for source, target in zip(long_tokens, short_tokens):
            source_key = _normalise_name_key(source)
            target_key = _normalise_name_key(target)
            if (
                source_key == target_key
                or not source_key
                or not target_key
                or not source_key.isalpha()
                or not target_key.isalpha()
                or len(target) >= len(source)
                or source_key[0] != target_key[0]
            ):
                continue
            token_candidates[source_key][target] += 1

    exact = {}
    for key, candidates in exact_candidates.items():
        ranked = candidates.most_common()
        top_value, top_count = ranked[0]
        total = sum(candidates.values())
        # Preserve case-only variants, or a strongly dominant historical result.
        normalised_outputs = {_normalise_name_key(value) for value in candidates}
        if len(normalised_outputs) == 1 or top_count / total >= 0.75:
            exact[key] = top_value

    rules = []
    for source, candidates in token_candidates.items():
        target, count = candidates.most_common(1)[0]
        total = sum(candidates.values())
        if count >= 2 and count / total >= 0.75:
            rules.append(
                {
                    "source": source,
                    "target": target,
                    "count": count,
                    "confidence": count / total,
                    "saving": len(source) - len(target),
                }
            )
    rules.sort(key=lambda item: (-item["saving"], -item["count"], item["source"]))
    return {"exact": exact, "rules": rules}


def _catalog_index(catalog_entries):
    index = defaultdict(list)
    for entry in catalog_entries:
        for field in ("code", "pon_sku", "code2"):
            value = str(getattr(entry, field, "") or "").strip().upper()
            if value:
                index[value].append(entry)
    return index


def _catalog_match_for_target(code, catalog_index, target_customer=""):
    matches = catalog_index.get(str(code or "").strip().upper(), [])
    target = str(target_customer or "").strip().upper()
    if target == "PON":
        pon_match = next((item for item in matches if str(getattr(item, "customer", "") or "").strip().upper() == "PON"), None)
        if pon_match:
            return pon_match, "EXISTING"
        pbp_match = next((item for item in matches if str(getattr(item, "customer", "") or "").strip().upper() == "PBP"), None)
        if pbp_match:
            return pbp_match, "PBP_TO_PON"
        return None, "NEW"
    if target:
        target_match = next(
            (item for item in matches if str(getattr(item, "customer", "") or "").strip().upper() == target),
            None,
        )
        if target_match:
            return target_match, "EXISTING"
    return (matches[0], "EXISTING") if matches else (None, "NEW")


def classify_product_codes(codes, catalog_entries, target_customer=""):
    """Return the current three-way product classification for each code."""
    index = _catalog_index(catalog_entries)
    result = {}
    for code in codes:
        normalized = str(code or "").strip().upper()
        if not normalized:
            continue
        _, classification = _catalog_match_for_target(normalized, index, target_customer)
        result[normalized] = classification
    return result


def _positive_decimal(value):
    text = str(value or "").replace("\xa0", " ").strip().replace(",", ".")
    if not text:
        return None
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return number if number > 0 else None


def pbp_to_pon_product_data(row):
    """Extract reusable PBP Translogic values without changing stored units.

    Height/width/length remain in the millimetres supplied by Translogic. Weight
    remains kilograms and Cubic remains the historical Translogic cubic value.
    """
    if not row.get("is_pbp_to_pon"):
        return {"valid": False, "missing": []}
    raw = dict(row.get("catalog_raw_data") or {})
    cubic = _positive_decimal(raw.get("cubic"))
    weight = _positive_decimal(raw.get("weight"))
    height = _positive_decimal(raw.get("height"))
    width = _positive_decimal(raw.get("width"))
    length = _positive_decimal(raw.get("length") or raw.get("lenght"))
    short_name = str(row.get("catalog_short_name") or "").strip()
    long_name = str(row.get("catalog_long_name") or "").strip()
    missing = []
    for key, value in (("cubic", cubic), ("weight", weight), ("height", height), ("width", width), ("length", length)):
        if value is None:
            missing.append(key)
    if not short_name or len(short_name) > 30:
        missing.append("short_name")
    if not long_name:
        missing.append("long_name")
    return {
        "valid": not missing,
        "missing": missing,
        "code": str(row.get("catalog_code") or row.get("code") or "").strip().upper(),
        "pon_sku": str(row.get("catalog_pon_sku") or "").strip().upper(),
        "code2": str(row.get("catalog_code2") or "").strip().upper(),
        "short_name": short_name,
        "long_name": long_name,
        "cubic": cubic,
        "weight_kg": weight,
        "height_mm": int(height) if height is not None and height == height.to_integral_value() else height,
        "width_mm": int(width) if width is not None and width == width.to_integral_value() else width,
        "length_mm": int(length) if length is not None and length == length.to_integral_value() else length,
    }


def compare_products_to_catalog(received_lines, catalog_entries, target_customer=""):
    grouped = {}
    for line in received_lines:
        key = (line.code.strip().upper(), line.location.strip().upper())
        item = grouped.setdefault(
            key,
            {
                "code": key[0],
                "long_code": "",
                "count": 0,
                "location": key[1],
            },
        )
        item["count"] += line.quantity
        if not item["long_code"] and line.long_code:
            item["long_code"] = line.long_code.strip().upper()

    catalog_index = _catalog_index(catalog_entries)
    result = []
    for item in grouped.values():
        match, classification = _catalog_match_for_target(item["code"], catalog_index, target_customer)
        matched_by = []
        if match:
            for field, label in (("code", "code"), ("pon_sku", "pon_sku"), ("code2", "code2")):
                if str(getattr(match, field, "") or "").strip().upper() == item["code"]:
                    matched_by.append(label)
        row = {
            **item,
            "classification": classification,
            "is_new": classification == "NEW",
            "is_pbp_to_pon": classification == "PBP_TO_PON",
            "status": "PBP → PON" if classification == "PBP_TO_PON" else ("New" if classification == "NEW" else "Existing"),
            "tl_code": match.code if match else "",
            "matched_by": ", ".join(matched_by),
            "catalog_customer": match.customer if match else "",
            "catalog_code": match.code if match else "",
            "catalog_pon_sku": match.pon_sku if match else "",
            "catalog_code2": match.code2 if match else "",
            "catalog_short_name": match.short_name if match else "",
            "catalog_long_name": match.long_name if match else "",
            "catalog_group1": match.group1 if match else "",
            "catalog_group2": match.group2 if match else "",
            "catalog_raw_data": dict(getattr(match, "raw_data", {}) or {}) if match else {},
        }
        row["pbp_to_pon_data"] = pbp_to_pon_product_data(row)
        result.append(row)
    result.sort(key=lambda row: (row["location"], row["code"]))
    return result


def product_check_summary(rows, original_classifications=None):
    """Summarize unique products while keeping original and current decisions separate."""
    current_by_code = {}
    for row in rows:
        code = str(row.get("code") or "").strip().upper()
        if not code:
            continue
        classification = row.get("classification") or ""
        if code not in current_by_code or not current_by_code[code]:
            current_by_code[code] = classification

    persisted = {
        str(code or "").strip().upper(): classification or ""
        for code, classification in (original_classifications or {}).items()
    }
    original_by_code = {
        code: persisted.get(code, "")
        for code in current_by_code
    }

    def count(classifications, value):
        return sum(1 for classification in classifications.values() if classification == value)

    return {
        "unique_products": len(current_by_code),
        "total_units": sum(row["count"] for row in rows),
        "original_existing_products": count(original_by_code, "EXISTING"),
        "original_pbp_to_pon_products": count(original_by_code, "PBP_TO_PON"),
        "original_new_products": count(original_by_code, "NEW"),
        "current_existing_products": count(current_by_code, "EXISTING"),
        "current_pbp_to_pon_products": count(current_by_code, "PBP_TO_PON"),
        "current_new_products": count(current_by_code, "NEW"),
        "existing_units": sum(row["count"] for row in rows if row.get("classification") == "EXISTING"),
        "pbp_to_pon_units": sum(row["count"] for row in rows if row.get("classification") == "PBP_TO_PON"),
        "new_units": sum(row["count"] for row in rows if row.get("classification") == "NEW"),
    }

def _code128_png(value, module_width=1, height=38, quiet_modules=10):
    """Render a standards-based Code 128 pattern directly to a crisp PNG."""
    barcode = Code128(str(value))
    if barcode.validate() != str(value):
        raise ValueError(f"The value cannot be encoded as Code 128: {value}")
    barcode.encode()
    pattern = barcode.decompose()
    widths = [ord(token.lower()) - ord("a") + 1 for token in pattern]
    quiet = quiet_modules * module_width
    image_width = sum(widths) * module_width + (quiet * 2)
    image = Image.new("RGB", (image_width, height), "white")
    draw = ImageDraw.Draw(image)
    position = quiet
    for token, units in zip(pattern, widths):
        width = units * module_width
        if token.isupper():
            draw.rectangle((position, 0, position + width - 1, height - 1), fill="black")
        position += width
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=False)
    output.seek(0)
    return output


def _group_product_rows_by_code(rows):
    """Build one printable row per code, using its first pallet in LOC order."""
    grouped = {}
    for row in rows:
        code = str(row["code"]).strip().upper()
        item = grouped.setdefault(
            code,
            {
                **row,
                "code": code,
                "count": 0,
                "_locations": [],
            },
        )
        item["count"] += int(row.get("count") or 0)
        location = str(row.get("location") or "").strip().upper()
        if location and location not in item["_locations"]:
            item["_locations"].append(location)
        for field in ("long_code", "tl_code", "matched_by", "catalog_customer"):
            if not item.get(field) and row.get(field):
                item[field] = row[field]

    result = []
    for code in grouped:
        item = grouped[code]
        locations = sorted(item.pop("_locations"))
        item["location"] = locations[0] if locations else ""
        result.append(item)
    return sorted(result, key=lambda item: (item["location"], item["code"]))

def _order_product_check_rows(rows):
    """Keep historical PBP migrations together before the physical-order rows."""
    printable_rows = _group_product_rows_by_code(rows)
    pbp_rows = sorted(
        (row for row in printable_rows if row.get("is_pbp_to_pon")),
        key=lambda row: row["code"],
    )
    existing_rows = [
        row for row in printable_rows
        if not row.get("is_pbp_to_pon") and not row.get("is_new")
    ]
    truly_new_rows = [
        row for row in printable_rows
        if row.get("is_new") and not row.get("is_pbp_to_pon")
    ]
    return pbp_rows + existing_rows + truly_new_rows




def _excel_column_width_pixels(width):
    """Approximate Excel's standard character-width conversion to pixels."""
    return int((float(width) * 7) + 5)


def _center_image_in_cell(image, row_number, column_number, column_width, row_height_points):
    """Return a one-cell drawing anchor centred within the target Excel cell."""
    cell_width = _excel_column_width_pixels(column_width)
    cell_height = points_to_pixels(row_height_points)
    horizontal_offset = max(0, round((cell_width - image.width) / 2))
    vertical_offset = max(0, round((cell_height - image.height) / 2))
    return OneCellAnchor(
        _from=AnchorMarker(
            col=column_number - 1,
            row=row_number - 1,
            colOff=pixels_to_EMU(horizontal_offset),
            rowOff=pixels_to_EMU(vertical_offset),
        ),
        ext=XDRPositiveSize2D(
            cx=pixels_to_EMU(image.width),
            cy=pixels_to_EMU(image.height),
        ),
    )


# PON_PBP_GREEN_NEW_BLOCK_1008_1500
# PON_PRODUCT_CHECK_DIMENSION_ORDER_1008_1540
# PON_PBP_PERSISTED_VISUAL_ORDER_1009_0815
def build_product_check_workbook(container_identifier, rows):
    """Create the review sheet used to identify new Translogic products."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PON Product Check"
    printable_rows = _order_product_check_rows(rows)

    headers = [container_identifier, "Long_Code", "Count", "LOC", "Status", "Length (cm)", "Height (cm)", "Width (cm)", "Kg", "TL Code", "Matched By"]
    worksheet.append(headers)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="FFF200")
    new_fill = PatternFill("solid", fgColor="FCE8E6")
    existing_fill = PatternFill("solid", fgColor="E6F4EA")
    dimension_fill = PatternFill("solid", fgColor="FFF4CC")
    migration_fill = PatternFill("solid", fgColor="FFFFFF00")

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="000000")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in printable_rows:
        definition = row.get("definition")
        history = row.get("pbp_to_pon_data") or {}
        use_history = bool(row.get("is_pbp_to_pon") and history.get("valid"))
        worksheet.append(
            [
                row["code"],
                row["long_code"],
                row["count"],
                row["location"],
                row.get("status", ""),
                round(float(history["length_mm"]) / 10, 1) if use_history else (round(definition.length_mm / 10, 1) if definition else ""),
                round(float(history["height_mm"]) / 10, 1) if use_history else (round(definition.height_mm / 10, 1) if definition else ""),
                round(float(history["width_mm"]) / 10, 1) if use_history else (round(definition.width_mm / 10, 1) if definition else ""),
                float(history["weight_kg"]) if use_history else (definition.weight_kg if definition else ""),
                row["tl_code"],
                row["matched_by"],
            ]
        )
        current = worksheet.max_row
        for column, cell in enumerate(worksheet[current], start=1):
            cell.border = border
            if column in (1, 2):
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            elif column <= 9:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        status_cell = worksheet.cell(current, 5)
        if row["is_new"]:
            status_cell.fill = new_fill
            status_cell.font = Font(bold=True, color="B3261E")
        elif row.get("is_pbp_to_pon"):
            status_cell.fill = migration_fill
            status_cell.font = Font(bold=True, color="006100")
        else:
            status_cell.fill = existing_fill
            status_cell.font = Font(bold=True, color="166534")
        if (row["is_new"] or row.get("is_pbp_to_pon")) and not definition and not (row.get("pbp_to_pon_data") or {}).get("valid"):
            for column in range(6, 10):
                worksheet.cell(current, column).fill = dimension_fill
        if row.get("is_pbp_to_pon"):
            for cell in worksheet[current]:
                cell.fill = migration_fill

    total_row = worksheet.max_row + 1
    worksheet.cell(total_row, 1, "Total")
    worksheet.cell(total_row, 3, sum(row["count"] for row in printable_rows))
    for cell in worksheet[total_row]:
        cell.border = border
        cell.font = Font(bold=True)

    widths = {"A": 15, "B": 22, "C": 7, "D": 9, "E": 7, "F": 11.5, "G": 11.5, "H": 11.5, "I": 7, "J": 22, "K": 20}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.column_dimensions["J"].hidden = True
    worksheet.column_dimensions["K"].hidden = True
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:I{max(1, worksheet.max_row - 1)}"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24
    worksheet.print_title_rows = "1:1"
    worksheet.print_area = f"A1:I{total_row}"
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.2
    worksheet.page_margins.right = 0.2
    worksheet.page_margins.top = 0.4
    worksheet.page_margins.bottom = 0.4
    worksheet.page_margins.header = 0.2
    worksheet.page_margins.footer = 0.2

    barcode_sheet = workbook.create_sheet("BARCODE Converter")
    barcode_sheet.sheet_view.showGridLines = False
    barcode_sheet["B2"] = container_identifier
    barcode_sheet["D2"] = "Qty"
    barcode_sheet["E2"] = "Barcode Item"
    barcode_sheet["F2"] = "Barcode Qty"
    barcode_sheet["G2"] = "Running total"

    for cell in (barcode_sheet["B2"], barcode_sheet["D2"], barcode_sheet["E2"], barcode_sheet["F2"], barcode_sheet["G2"]):
        cell.font = Font(bold=True, size=14)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    barcode_sheet.row_dimensions[2].height = 30

    quantities_by_code = {row["code"]: row["count"] for row in printable_rows}

    barcode_streams = []
    running_total = 0
    for row_number, (code, quantity) in enumerate(sorted(quantities_by_code.items()), start=3):
        running_total += quantity
        barcode_sheet.cell(row_number, 2, code)
        barcode_sheet.cell(row_number, 4, quantity)
        barcode_sheet.cell(row_number, 7, running_total)
        barcode_sheet.row_dimensions[row_number].height = 43

        item_stream = _code128_png(code)
        quantity_stream = _code128_png(quantity)
        barcode_streams.extend([item_stream, quantity_stream])
        item_image = WorksheetImage(item_stream)
        quantity_image = WorksheetImage(quantity_stream)
        item_anchor = _center_image_in_cell(item_image, row_number, 5, 34, 43)
        quantity_anchor = _center_image_in_cell(quantity_image, row_number, 6, 18, 43)
        barcode_sheet.add_image(item_image, item_anchor)
        barcode_sheet.add_image(quantity_image, quantity_anchor)

        for column in (2, 4, 5, 6, 7):
            cell = barcode_sheet.cell(row_number, column)
            cell.font = Font(size=14)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if column != 2 else "left", vertical="center")

    barcode_sheet.column_dimensions["A"].width = 3
    barcode_sheet.column_dimensions["B"].width = 23
    barcode_sheet.column_dimensions["C"].hidden = True
    barcode_sheet.column_dimensions["D"].width = 10
    barcode_sheet.column_dimensions["E"].width = 34
    barcode_sheet.column_dimensions["F"].width = 18
    barcode_sheet.column_dimensions["G"].width = 14
    barcode_sheet.freeze_panes = None
    barcode_sheet.print_title_rows = "2:2"
    barcode_sheet.print_area = f"B2:G{max(2, barcode_sheet.max_row)}"
    barcode_sheet.page_setup.orientation = "portrait"
    barcode_sheet.page_setup.fitToWidth = 1
    barcode_sheet.page_setup.fitToHeight = 0
    barcode_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    barcode_sheet.oddHeader.center.text = "PON BARCODE CONVERTER"
    barcode_sheet.oddFooter.center.text = "Page &P of &N"

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output



def parse_product_moves_csv(source):
    """Parse the Translogic PRODUCT_MOVES.CSV export.

    The verified operational file contains exactly two columns: ``product`` and
    ``stock_in_mov``. Movement values are padded with spaces in the source; the
    database stores the significant digits and the UPStockSerial writer restores
    the fixed width required by the legacy import file.
    """
    if hasattr(source, "read"):
        data = source.read()
        try:
            source.seek(0)
        except Exception:
            pass
    else:
        data = Path(source).read_bytes()

    if isinstance(data, str):
        text = data
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")

    reader = csv.DictReader(io.StringIO(text, newline=""))
    if not reader.fieldnames:
        raise ValueError("PRODUCT_MOVES.CSV has no header row.")

    header_map = {str(name or "").strip().lower(): name for name in reader.fieldnames}
    missing = [name for name in ("product", "stock_in_mov") if name not in header_map]
    if missing:
        raise ValueError(
            "PRODUCT_MOVES.CSV must contain the columns product and stock_in_mov. "
            f"Missing: {', '.join(missing)}."
        )

    product_key = header_map["product"]
    movement_key = header_map["stock_in_mov"]
    rows = []
    seen_movements = set()
    errors = []
    for source_row, raw in enumerate(reader, start=2):
        product = str(raw.get(product_key, "") or "").strip().upper()
        movement = str(raw.get(movement_key, "") or "").strip()
        if not product and not movement:
            continue
        if not product:
            errors.append(f"row {source_row}: product is blank")
            continue
        if not movement:
            errors.append(f"row {source_row}: stock_in_mov is blank")
            continue
        if not movement.isdigit():
            errors.append(f"row {source_row}: stock_in_mov '{movement}' is not numeric")
            continue
        if len(movement) > 10:
            errors.append(f"row {source_row}: stock_in_mov '{movement}' is longer than 10 characters")
            continue
        if movement in seen_movements:
            errors.append(f"row {source_row}: duplicate stock_in_mov '{movement}'")
            continue
        seen_movements.add(movement)
        rows.append({"source_row": source_row, "product": product, "movement": movement})

    if errors:
        detail = "; ".join(errors[:10])
        if len(errors) > 10:
            detail += f"; and {len(errors) - 10} more error(s)"
        raise ValueError(detail)
    if not rows:
        raise ValueError("PRODUCT_MOVES.CSV contains no movement rows.")
    return rows


def reconcile_product_movements(received_lines, movement_rows):
    """Match each Translogic movement to one physical bike from the first scan.

    Matching is by product code, not by global row position. Within each product,
    the first-scan source-row order is retained. The source row becomes serial_no,
    exactly as in the verified legacy UPStockSerial.csv process.
    """
    units_by_product = defaultdict(list)
    issues = []
    serials = set()

    for line in sorted(received_lines, key=lambda item: (item.source_row, item.id or 0)):
        product = str(line.code or "").strip().upper()
        quantity = int(line.quantity or 0)
        location = str(line.location or "").strip().upper()
        if quantity != 1:
            issues.append(
                f"First-scan row {line.source_row} ({product}) has quantity {quantity}; "
                "Stage 2 requires one physical bike per scan row."
            )
            continue
        if not product:
            issues.append(f"First-scan row {line.source_row} has no product code.")
            continue
        if not LOCATION_PATTERN.fullmatch(location):
            issues.append(
                f"First-scan row {line.source_row} ({product}) has invalid or missing location '{location}'."
            )
            continue
        serial_no = str(line.source_row)
        if serial_no in serials:
            issues.append(f"Duplicate first-scan serial/source row {serial_no}.")
            continue
        serials.add(serial_no)
        units_by_product[product].append(
            {"product": product, "serial_no": serial_no, "location": location, "source_row": line.source_row}
        )

    received_counts = Counter({product: len(units) for product, units in units_by_product.items()})
    movement_counts = Counter(row["product"] for row in movement_rows)
    product_counts = []
    all_products = sorted(set(received_counts) | set(movement_counts))
    for product in all_products:
        received = received_counts.get(product, 0)
        movements = movement_counts.get(product, 0)
        product_counts.append(
            {
                "product": product,
                "received": received,
                "movements": movements,
                "difference": movements - received,
                "status": "MATCH" if received == movements else "MISMATCH",
            }
        )

    count_mismatches = [row for row in product_counts if row["status"] != "MATCH"]
    mapping_rows = []
    offsets = defaultdict(int)
    for movement in movement_rows:
        product = movement["product"]
        index = offsets[product]
        units = units_by_product.get(product, [])
        if index < len(units):
            unit = units[index]
            mapping_rows.append(
                {
                    "product": product,
                    "movement": movement["movement"],
                    "serial_no": unit["serial_no"],
                    "location": unit["location"],
                    "status": "MATCH",
                }
            )
            offsets[product] += 1
        else:
            mapping_rows.append(
                {
                    "product": product,
                    "movement": movement["movement"],
                    "serial_no": "",
                    "location": "",
                    "status": "UNMATCHED",
                }
            )

    ready = not issues and not count_mismatches and len(mapping_rows) == len(movement_rows)
    return {
        "ready": ready,
        "issues": issues,
        "product_counts": product_counts,
        "count_mismatches": count_mismatches,
        "rows": mapping_rows,
        "received_units": sum(received_counts.values()),
        "movement_units": len(movement_rows),
        "matched_units": sum(1 for row in mapping_rows if row["status"] == "MATCH"),
    }


# PON_UPSTOCK_NO_EMPTY_ROWS_1008.1313
def build_upstockserial_csv(rows, max_rows=500):
    """Build the three-column UPStockSerial.csv file.

    The historical Excel worksheet had capacity through row 501, but the CSV output
    must contain only real bike rows. ``max_rows`` remains a safety limit; no empty
    padding rows are appended.
    """
    if len(rows) > max_rows:
        raise ValueError(
            f"UPStockSerial.csv supports a maximum of {max_rows} bike rows in the verified template; "
            f"this container has {len(rows)}."
        )
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(["Movement", "serial_no", "Loc"])

    # The legacy Excel process writes UPStockSerial.csv in ascending Movement order.
    # PRODUCT_MOVES.CSV can arrive in product/receiving order, so sort only the
    # final output rows here.  The movement-to-serial/location association created
    # during reconciliation is preserved unchanged.
    ordered_rows = sorted(
        rows,
        key=lambda row: int(str(row.get("movement", "") or "").strip() or "0"),
    )

    for row in ordered_rows:
        movement = str(row.get("movement", "") or "").strip()
        serial_no = str(row.get("serial_no", "") or "").strip()
        location = str(row.get("location", "") or "").strip().upper()
        if not movement or not serial_no or not location:
            raise ValueError("Cannot generate UPStockSerial.csv while a matched row is incomplete.")
        writer.writerow([movement.rjust(10), serial_no, location])
    return output.getvalue().encode("utf-8")


# PON_CLIENT_REPORT_TWO_CODE_SPARES_1008.1032
def resolve_client_report_two_codes(client_lines, received_lines, configured_two_codes=False):
    """Resolve the customer-report reconciliation key without changing First Scan.

    A stored TWO_CODES mode is authoritative. When older containers have a blank
    stored mode, compare the confirmed client manifest against the confirmed
    First Scan and choose Long Code only when it has strictly better overlap
    than physical Code. This preserves the normal one-code flow.
    """
    if configured_two_codes:
        return True

    client_keys = set()
    for line in client_lines:
        key = _line_reconciliation_code(line, use_long_code=True)
        if key and is_pon_long_code(key):
            client_keys.add(key)

    if not client_keys:
        return False

    code_overlap = 0
    long_overlap = 0
    for line in received_lines:
        quantity = int(getattr(line, "quantity", 0) or 0)
        code = str(getattr(line, "code", "") or "").strip().upper()
        long_code = str(getattr(line, "long_code", "") or "").strip().upper()
        if code in client_keys:
            code_overlap += quantity
        if long_code in client_keys:
            long_overlap += quantity

    return long_overlap > 0 and long_overlap > code_overlap


def build_client_report_rows(client_lines, received_lines, catalog_entries=(), product_definitions=(), product_statuses=None, product_classifications=None, target_customer="", two_codes=None):
    """Build final customer receiving rows, separating bikes from spare parts.

    Bike reconciliation uses physical Code for normal one-code flows and Long Code
    for confirmed/detected two-code flows. Non-bike manifest rows are retained as
    spare-parts rows using the exact client description and manifest section, but
    they are not treated as missing bikes and do not receive a bike First Scan count.
    """
    client_lines = list(client_lines)
    received_lines = list(received_lines)
    if two_codes is None:
        two_codes = resolve_client_report_two_codes(
            client_lines,
            received_lines,
            configured_two_codes=_uses_two_code_reconciliation(received_lines),
        )

    advised = defaultdict(int)
    received = defaultdict(int)
    client_descriptions = {}
    received_descriptions = {}
    physical_codes = defaultdict(list)
    ordered_codes = []
    seen_codes = set()
    spare_parts = []
    spare_index = {}

    def remember(code):
        if code and code not in seen_codes:
            ordered_codes.append(code)
            seen_codes.add(code)

    for line in client_lines:
        code = _line_reconciliation_code(line, use_long_code=two_codes)
        if not code:
            continue

        if two_codes and not is_pon_long_code(code):
            section_name = manifest_section(line) or "SPARE PARTS"
            client_name = str(getattr(line, "description", "") or "").replace("\xa0", " ").strip()
            spare_key = (section_name, code, client_name)
            quantity = int(getattr(line, "quantity", 0) or 0)
            if spare_key in spare_index:
                spare_parts[spare_index[spare_key]]["advised"] += quantity
            else:
                spare_index[spare_key] = len(spare_parts)
                spare_parts.append(
                    {
                        "code": code,
                        "item_name": client_name or code,
                        "advised": quantity,
                        "received": None,
                        "variance": None,
                        "comment": "Not part of bike First Scan",
                        "is_new": False,
                        "is_pbp_to_pon": False,
                        "classification": "SPARE_PART",
                        "is_spare_part": True,
                        "section_name": section_name,
                    }
                )
            continue

        remember(code)
        advised[code] += int(getattr(line, "quantity", 0) or 0)
        description = re.sub(r"\s+", " ", str(getattr(line, "description", "") or "").replace("\xa0", " ").strip())
        if description and code not in client_descriptions:
            client_descriptions[code] = description

    for line in received_lines:
        code = _line_reconciliation_code(line, use_long_code=two_codes)
        if not code:
            continue
        if two_codes and not is_pon_long_code(code):
            continue
        remember(code)
        received[code] += int(getattr(line, "quantity", 0) or 0)
        physical_code = str(getattr(line, "code", "") or "").strip().upper()
        if physical_code and physical_code not in physical_codes[code]:
            physical_codes[code].append(physical_code)
        description = re.sub(r"\s+", " ", str(getattr(line, "description", "") or "").replace("\xa0", " ").strip())
        if description and code not in received_descriptions:
            received_descriptions[code] = description

    catalogue_rows = list(catalog_entries)
    catalogue_index = _catalog_index(catalogue_rows)

    definitions = {
        str(getattr(item, "code", "") or "").strip().upper(): item
        for item in product_definitions
        if str(getattr(item, "code", "") or "").strip()
    }

    rows = []
    for code in ordered_codes:
        lookup_code = physical_codes.get(code, [code])[0]
        catalogue, _ = _catalog_match_for_target(lookup_code, catalogue_index, target_customer)
        definition = definitions.get(lookup_code) or definitions.get(code)
        if catalogue and str(getattr(catalogue, "long_name", "") or "").strip():
            long_name = re.sub(r"\s+", " ", str(catalogue.long_name).replace("\xa0", " ").strip())
        elif definition and str(getattr(definition, "full_name", "") or "").strip():
            long_name = re.sub(r"\s+", " ", str(definition.full_name).replace("\xa0", " ").strip())
        else:
            long_name = client_descriptions.get(code) or received_descriptions.get(code) or code

        advised_qty = advised.get(code, 0)
        received_qty = received.get(code, 0)
        classification = (product_classifications or {}).get(lookup_code, "")
        historical_new = product_statuses is not None and lookup_code in product_statuses and bool(product_statuses[lookup_code])
        is_new = historical_new or classification == "NEW" or (not classification and catalogue is None)
        is_pbp_to_pon = (not is_new) and classification == "PBP_TO_PON"
        rows.append(
            {
                "code": code,
                "item_name": long_name,
                "advised": advised_qty,
                "received": received_qty,
                "variance": (advised_qty - received_qty) if two_codes else (received_qty - advised_qty),
                "comment": (
                    "Unadvised / Not Advised"
                    if two_codes and advised_qty == 0 and received_qty
                    else ("New" if is_new else "")
                ),
                "is_new": is_new,
                "is_pbp_to_pon": is_pbp_to_pon,
                "classification": "PBP_TO_PON" if is_pbp_to_pon else ("NEW" if is_new else "EXISTING"),
                "is_spare_part": False,
                "section_name": "",
            }
        )

    rows.extend(spare_parts)
    return rows


def build_client_receiving_workbook(container_identifier, rows):
    """Create the final customer receiving workbook with spare parts separated."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Receiving Container Data"
    worksheet.sheet_view.showGridLines = False

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    yellow = PatternFill("solid", fgColor="FFFFFF00")
    migration_green = PatternFill("solid", fgColor="FFC6EFCE")
    spare_header_fill = PatternFill("solid", fgColor="FFD9E1F2")

    worksheet["B1"] = "Container:"
    worksheet["C1"] = str(container_identifier or "").strip()
    worksheet["B1"].font = Font(bold=True, size=11)
    worksheet["C1"].font = Font(bold=True, size=11)

    headers = ["Item Code", "Item Name", "Advised", "Received", "Var", "Comment"]
    for column, value in enumerate(headers, start=2):
        cell = worksheet.cell(2, column, value)
        cell.font = Font(bold=True, size=10)
        cell.fill = yellow
        cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    bike_rows = [row for row in rows if not row.get("is_spare_part")]
    spare_rows = [row for row in rows if row.get("is_spare_part")]

    row_number = 3
    for row in bike_rows:
        values = [
            row.get("code", ""),
            row.get("item_name", ""),
            int(row.get("advised", 0) or 0),
            int(row.get("received", 0) or 0),
            int(row.get("variance", 0) or 0),
            row.get("comment", ""),
        ]
        for column, value in enumerate(values, start=2):
            cell = worksheet.cell(row_number, column, value)
            cell.border = border
            cell.font = Font(size=10)
            cell.alignment = Alignment(
                horizontal="center" if column in (4, 5, 6) else "left",
                vertical="center",
                wrap_text=column == 3,
            )
            if row.get("is_pbp_to_pon"):
                cell.fill = migration_green
        worksheet.row_dimensions[row_number].height = 19
        row_number += 1

    worksheet.cell(row_number, 3, "Bike Total" if spare_rows else "Total")
    worksheet.cell(row_number, 4, sum(int(row.get("advised", 0) or 0) for row in bike_rows))
    worksheet.cell(row_number, 5, sum(int(row.get("received", 0) or 0) for row in bike_rows))
    worksheet.cell(row_number, 6, sum(int(row.get("variance", 0) or 0) for row in bike_rows))
    for column in range(3, 7):
        worksheet.cell(row_number, column).font = Font(bold=True, size=10)
        worksheet.cell(row_number, column).border = border
    row_number += 1

    if spare_rows:
        current_section = None
        for row in spare_rows:
            section = str(row.get("section_name", "") or "SPARE PARTS").strip()
            if section != current_section:
                if current_section is not None:
                    row_number += 1
                worksheet.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=7)
                section_cell = worksheet.cell(row_number, 2, section)
                section_cell.font = Font(bold=True, size=10)
                section_cell.fill = spare_header_fill
                section_cell.border = border
                section_cell.alignment = Alignment(horizontal="left", vertical="center")
                for column in range(2, 8):
                    worksheet.cell(row_number, column).fill = spare_header_fill
                    worksheet.cell(row_number, column).border = border
                current_section = section
                row_number += 1

            values = [
                row.get("code", ""),
                row.get("item_name", ""),
                int(row.get("advised", 0) or 0),
                "",
                "",
                row.get("comment", "Not part of bike First Scan"),
            ]
            for column, value in enumerate(values, start=2):
                cell = worksheet.cell(row_number, column, value)
                cell.border = border
                cell.font = Font(size=10)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            worksheet.row_dimensions[row_number].height = 19
            row_number += 1

        worksheet.cell(row_number, 3, "Spare Parts Total")
        worksheet.cell(row_number, 4, sum(int(row.get("advised", 0) or 0) for row in spare_rows))
        for column in range(3, 7):
            worksheet.cell(row_number, column).font = Font(bold=True, size=10)
            worksheet.cell(row_number, column).border = border
        row_number += 1

    worksheet.column_dimensions["B"].width = 34
    worksheet.column_dimensions["C"].width = 46
    worksheet.column_dimensions["D"].width = 12
    worksheet.column_dimensions["E"].width = 12
    worksheet.column_dimensions["F"].width = 10
    worksheet.column_dimensions["G"].width = 27
    worksheet.freeze_panes = "B3"
    worksheet.print_title_rows = "1:2"
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def calculate_cubic(length_mm, height_mm, width_mm):
    value = (Decimal(length_mm) * Decimal(height_mm) * Decimal(width_mm)) / Decimal("1000000000")
    return value.quantize(Decimal("0.001"), rounding=ROUND_UP)


def _normalise_family_match(value):
    value = str(value or "").replace("\xa0", " ").upper()
    value = re.sub(r"[^A-Z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def resolve_group1_family(description, customer, family_rules):
    """Resolve Group1 conservatively from Django-admin family mappings.

    A result is accepted only when the highest-priority matching family is unique.
    Overlapping broad/specific rules are handled by priority; ties across different
    suffixes are treated as ambiguous and must be resolved in Django administration.
    """
    text = _normalise_family_match(description)
    customer_prefix = _normalise_family_match(customer).replace(" ", "")
    if not text:
        return {"resolved": False, "reason": "missing_name", "message": "LongName is empty, so Group1 cannot be assigned."}
    if not customer_prefix:
        return {"resolved": False, "reason": "missing_customer", "message": "The Translogic customer prefix is not configured."}

    matches = []
    for rule in family_rules:
        active = rule.get("active", True) if isinstance(rule, dict) else getattr(rule, "active", True)
        if not active:
            continue
        phrases_raw = rule.get("match_phrases", "") if isinstance(rule, dict) else getattr(rule, "match_phrases", "")
        phrases = [_normalise_family_match(item) for item in str(phrases_raw or "").splitlines()]
        phrases = [item for item in phrases if item]
        matching_phrases = [phrase for phrase in phrases if f" {phrase} " in f" {text} "]
        if not matching_phrases:
            continue
        family_name = rule.get("family_name", "") if isinstance(rule, dict) else getattr(rule, "family_name", "")
        suffix = rule.get("suffix", "") if isinstance(rule, dict) else getattr(rule, "suffix", "")
        priority = int(rule.get("priority", 100) if isinstance(rule, dict) else getattr(rule, "priority", 100))
        rule_id = rule.get("id") if isinstance(rule, dict) else getattr(rule, "pk", None)
        suffix = str(suffix or "").strip().upper()
        if not suffix:
            continue
        matches.append({
            "id": rule_id,
            "family_name": str(family_name or "").strip(),
            "suffix": suffix,
            "priority": priority,
            "matched_phrase": max(matching_phrases, key=len),
        })

    if not matches:
        return {
            "resolved": False,
            "reason": "no_match",
            "message": "No clear Group1 family relation was found. Add this relation in Administration → Group1 family mappings.",
        }

    highest = max(item["priority"] for item in matches)
    top = [item for item in matches if item["priority"] == highest]
    distinct_suffixes = {item["suffix"] for item in top}
    if len(distinct_suffixes) != 1:
        families = ", ".join(sorted({item["family_name"] or item["suffix"] for item in top}))
        return {
            "resolved": False,
            "reason": "ambiguous",
            "message": f"Group1 is ambiguous between: {families}. Adjust the relation or priority in Administration → Group1 family mappings.",
            "matches": top,
        }

    chosen = sorted(top, key=lambda item: (-len(item["matched_phrase"]), item["family_name"]))[0]
    chosen["resolved"] = True
    chosen["group1"] = f"{customer_prefix}{chosen['suffix']}"
    return chosen


def _replace_learned_token(value, source, target):
    pattern = rf"(?<!\w){re.escape(source)}(?!\w)"
    return re.sub(pattern, target, value, flags=re.IGNORECASE)


def suggest_translogic_name(description, max_length=30, name_dictionary=None):
    """Suggest a meaningful Translogic name using the historical XML before fallbacks."""
    original = re.sub(r"\s+", " ", str(description or "").replace("\xa0", " ").strip())
    if not original:
        return ""

    dictionary = name_dictionary or {}
    exact = dictionary.get("exact", {})
    historical = exact.get(_normalise_name_key(original))
    if historical and len(historical) <= max_length:
        return historical

    value = original
    if len(value) <= max_length:
        return value

    # Apply only conservative rules learned repeatedly from the uploaded XML.
    for rule in dictionary.get("rules", []):
        changed = _replace_learned_token(value, rule["source"], rule["target"])
        if changed != value:
            value = re.sub(r"\s+", " ", changed).strip()
            if len(value) <= max_length:
                return value

    value = re.sub(r"\b(\d+(?:\.\d+)?)MM\b", r"\1", value, flags=re.IGNORECASE)
    if len(value) <= max_length:
        return value

    # Stable business fallbacks remain after the XML-derived dictionary.
    for source, replacement in SHORT_NAME_REPLACEMENTS:
        value = re.sub(rf"\b{re.escape(source)}\b", replacement, value, flags=re.IGNORECASE)
        value = re.sub(r"\s+", " ", value).strip()
        if len(value) <= max_length:
            return value

    words = [word for word in value.split() if word.upper() not in {"THE", "AND", "WITH", "OF"}]
    value = " ".join(words)
    if len(value) <= max_length:
        return value

    # Last resort: remove internal vowels from descriptive words while preserving the
    # first and final tokens, which commonly hold the family/model and size.
    while len(" ".join(words)) > max_length:
        candidates = []
        for index, word in enumerate(words):
            if index in {0, len(words) - 1} or len(word) <= 3:
                continue
            removable = [
                position for position in range(1, len(word) - 1)
                if word[position].upper() in "AEIOU"
            ]
            if removable:
                candidates.append((len(word), index, removable[-1]))
        if not candidates:
            break
        _, word_index, character_index = max(candidates)
        word = words[word_index]
        words[word_index] = word[:character_index] + word[character_index + 1:]
    value = " ".join(words)
    if len(value) > max_length:
        raise ValueError("A meaningful 30-character product name could not be suggested automatically.")
    return value


def build_new_product_workbook(definitions, configuration, container_date, pbp_to_pon_codes=None):
    """Build the exact 30-column Translogic New Product Import worksheet."""
    if not container_date:
        raise ValueError("Enter the container date before generating New Product Import.")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "New Product Import"
    worksheet.append(NEW_PRODUCT_HEADERS)
    threshold = Decimal(str(configuration["TL_GROUP2_CUBIC_THRESHOLD"]))

    historical_pbp_codes = None if pbp_to_pon_codes is None else {
        str(code or "").strip().upper() for code in pbp_to_pon_codes
    }

    for product in definitions:
        if len(product.short_name) > 30:
            raise ValueError(f"{product.code}: column C may contain at most 30 characters.")
        if not getattr(product, "group1", ""):
            raise ValueError(f"{product.code}: Group1 has not been resolved. Add the family relation in Django Administration and review the product again.")
        cubic_override = getattr(product, "cubic_override", None)
        cubic = Decimal(str(cubic_override)) if cubic_override is not None else calculate_cubic(product.length_mm, product.height_mm, product.width_mm)
        group2 = configuration["TL_GROUP2_LARGE"] if cubic > threshold else configuration["TL_GROUP2_DEFAULT"]
        worksheet.append(
            [
                product.code,
                product.long_code,
                product.short_name,
                getattr(product, "customer_override", configuration["TL_CUSTOMER"]),
                "", "", "",
                getattr(product, "status_override", configuration["TL_STATUS"]),
                "",
                product.group1,
                group2,
                "",
                product.full_name,
                "", "",
                int(getattr(product, "quantity_override", configuration["TL_QUANTITY"])),
                "",
                int(getattr(product, "pallet_override", configuration["TL_PALLET"])),
                int(getattr(product, "lift_override", configuration["TL_LIFT"])),
                float(cubic),
                float(product.weight_kg),
                product.height_mm,
                product.width_mm,
                product.length_mm,
                "", "",
                int(getattr(product, "outer_override", configuration["TL_OUTER"])),
                "", "",
                container_date,
            ]
        )

        product_code = str(getattr(product, "code", "") or "").strip().upper()
        is_historical_pbp = (
            getattr(product, "is_pbp_to_pon", False)
            if historical_pbp_codes is None
            else product_code in historical_pbp_codes
        )
        if is_historical_pbp:
            for cell in worksheet[worksheet.max_row]:
                cell.fill = PatternFill("solid", fgColor="FFFFFF00")

    yellow_fill = PatternFill("solid", fgColor="FFFFFF00")
    yellow_columns = {"A", "C", "D", "H", "J", "M", "P", "R", "S", "T", "U", "AA", "AD"}
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        if cell.column_letter in yellow_columns:
            cell.fill = yellow_fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:AD{max(1, worksheet.max_row)}"
    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 32
    worksheet.column_dimensions["M"].width = 45
    worksheet.column_dimensions["AC"].width = 24
    worksheet.column_dimensions["AD"].width = 12
    for column in ("T", "U", "V", "W", "X"):
        worksheet.column_dimensions[column].width = 12
    for row in worksheet.iter_rows(min_row=2, min_col=20, max_col=21):
        row[0].number_format = "0.000"
        row[1].number_format = "0.000"
    for cell in worksheet["AD"][1:]:
        cell.number_format = "dd/mm/yyyy"
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def convert_excel_output(xlsx_content, base_name, export_format):
    """Return (bytes, extension, MIME type) using the centrally configured workbook format."""
    if export_format == "xlsx":
        return xlsx_content.getvalue(), ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if export_format != "excel_5_95":
        raise ValueError(f"Unsupported Excel export format: {export_format}")

    with tempfile.TemporaryDirectory(prefix="pon_excel_") as folder:
        folder_path = Path(folder)
        source = folder_path / f"{base_name}.xlsx"
        source.write_bytes(xlsx_content.getvalue())
        command = [
            "ssconvert",
            "--export-type=Gnumeric_Excel:excel_biff7",
            str(source),
            str(folder_path / f"{base_name}.xls"),
        ]
        try:
            completed = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
        except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
            raise RuntimeError("Microsoft Excel 5.0/95 conversion is unavailable in this installation.") from exc
        destination = folder_path / f"{base_name}.xls"
        if completed.returncode or not destination.exists():
            detail = (completed.stderr or completed.stdout or "unknown conversion error").strip()
            raise RuntimeError(f"Could not create the Microsoft Excel 5.0/95 file: {detail}")
        content = destination.read_bytes()
        if not _is_excel_5_95(content):
            raise RuntimeError("The generated file was not confirmed as Microsoft Excel 5.0/95 (BIFF7).")
        return content, ".xls", "application/vnd.ms-excel"


def _is_excel_5_95(content):
    """Confirm the BIFF5/7 BOF marker stored inside the OLE workbook stream."""
    return b"\x09\x08\x08\x00\x00\x05" in content


def _workbook(path):
    return load_workbook(path, data_only=True, read_only=False, keep_vba=str(path).lower().endswith(".xlsm"))


def inspect_workbook(path):
    wb = _workbook(path)
    sheets = []
    for ws in wb.worksheets:
        nonempty = 0
        preview = []
        for row in ws.iter_rows():
            values = [cell.value for cell in row]
            if any(value not in (None, "") for value in values):
                nonempty += 1
                if len(preview) < 8:
                    preview.append({"row": row[0].row, "values": [str(v) if v is not None else "" for v in values[:12]]})
        sheets.append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "nonempty": nonempty, "preview": preview})
    return sheets


def build_source_file_preview(path, sheet_name, start_row, start_column, mapping=None, row_limit=15, column_limit=12):
    """Read a positional workbook window without importing or changing operational data."""
    wb = _workbook(path)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError("The selected worksheet is not present in the source file.")
        ws = wb[sheet_name]
        first_row = max(1, int(start_row))
        first_column = column_index_from_string(str(start_column or "A").strip().upper())
        if row_limit < 1 or column_limit < 1:
            raise ValueError("Preview limits must be positive.")

        last_column = min(max(ws.max_column, first_column), first_column + column_limit - 1)
        last_row = min(ws.max_row, first_row + row_limit - 1)
        current_mapping = dict(mapping or {})
        header_row = int(current_mapping.get("_header_row") or 0)
        meanings = (
            ("code_column", "Product / Main Code"),
            ("long_code_column", "Long / Secondary Code"),
            ("description_column", "Description"),
            ("quantity_column", "Quantity"),
            ("container_column", "Container"),
        )
        mapped_by_column = defaultdict(list)
        for field, label in meanings:
            column = str(current_mapping.get(field) or "").strip().upper()
            if column:
                mapped_by_column[column].append(label)
        if current_mapping.get("location_stream"):
            column = str(current_mapping.get("code_column") or "").strip().upper()
            if column:
                mapped_by_column[column].append("Pallet / Location stream")

        columns = []
        for column_number in range(first_column, last_column + 1):
            letter = get_column_letter(column_number)
            original_header = ""
            if 1 <= header_row <= ws.max_row:
                value = ws.cell(header_row, column_number).value
                original_header = str(value).strip() if value not in (None, "") else ""
            columns.append({
                "letter": letter,
                "original_header": original_header,
                "meaning": " / ".join(mapped_by_column.get(letter, [])) or "Not mapped",
            })

        rows = []
        if first_row <= ws.max_row:
            for row_number in range(first_row, last_row + 1):
                rows.append({
                    "number": row_number,
                    "values": [
                        "" if ws.cell(row_number, column_number).value is None else str(ws.cell(row_number, column_number).value)
                        for column_number in range(first_column, last_column + 1)
                    ],
                })
        first_letter = get_column_letter(first_column)
        return {
            "worksheet": ws.title,
            "start_row": first_row,
            "start_column": first_letter,
            "first_cell": f"{first_letter}{first_row}",
            "header_row": header_row,
            "columns": columns,
            "rows": rows,
        }
    finally:
        wb.close()


def normalize_container_identifier(value):
    """Return the canonical 4-letter/7-digit container identifier or an empty string."""
    if value in (None, ""):
        return ""
    candidate = re.sub(r"\s+", "", str(value).strip().upper())
    return candidate if CONTAINER_PATTERN.fullmatch(candidate) else ""


def _normalise_header(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"[_-]+", " ", text)
    text = re.sub(r"[^a-z0-9. ]+", "", text)
    return re.sub(r"\s+", " ", text).strip()


def _header_score(header, field):
    header = _normalise_header(header)
    aliases = HEADER_ALIASES.get(field, set())
    if not header:
        return 0
    if header in aliases:
        return 60
    for alias in aliases:
        alias_n = _normalise_header(alias)
        if alias_n and (alias_n in header or header in alias_n):
            return 42
    return 0


def _positive_whole_number(value):
    if value in (None, "") or isinstance(value, bool):
        return False
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return False
    return number > 0 and number == number.to_integral_value()


def _column_profile(ws, column_number, start_row, known_product_codes=None, max_samples=220):
    values = []
    for row_number in range(start_row, min(ws.max_row, start_row + max_samples - 1) + 1):
        value = ws.cell(row_number, column_number).value
        if value not in (None, ""):
            values.append(value)
    if not values:
        return {
            "nonempty": 0, "container_ratio": 0, "quantity_ratio": 0, "product_ratio": 0,
            "master_ratio": 0, "text_ratio": 0, "unique_ratio": 0, "avg_length": 0,
        }
    texts = [str(v).strip() for v in values]
    container_hits = sum(1 for v in values if normalize_container_identifier(v))
    quantity_hits = sum(1 for v in values if _positive_whole_number(v))
    product_hits = sum(1 for t in texts if PRODUCT_CODE_PATTERN.fullmatch(re.sub(r"\s+", "", t.upper())))
    master_hits = 0
    known = known_product_codes or set()
    if known:
        master_hits = sum(1 for t in texts if re.sub(r"\s+", "", t.upper()) in known)
    text_hits = sum(1 for t in texts if len(t) >= 12 and not _positive_whole_number(t))
    unique = len(set(texts))
    return {
        "nonempty": len(values),
        "container_ratio": container_hits / len(values),
        "quantity_ratio": quantity_hits / len(values),
        "product_ratio": product_hits / len(values),
        "master_ratio": master_hits / len(values) if known else 0,
        "text_ratio": text_hits / len(values),
        "unique_ratio": unique / len(values),
        "avg_length": sum(len(t) for t in texts) / len(texts),
    }


def _field_confidence(field, header, profile):
    score = _header_score(header, field)
    if field == "code_column":
        score += round(profile["product_ratio"] * 22)
        score += round(profile["master_ratio"] * 38)
        if profile["unique_ratio"] >= 0.35:
            score += 8
    elif field == "description_column":
        score += round(profile["text_ratio"] * 32)
        if profile["avg_length"] >= 15:
            score += 8
    elif field == "quantity_column":
        score += round(profile["quantity_ratio"] * 38)
        if profile["unique_ratio"] < 0.40:
            score += 4
    elif field == "container_column":
        # Data pattern is deliberately weighted above the header because client headers can be misleading.
        score += round(profile["container_ratio"] * 84)
        if profile["container_ratio"] >= 0.70 and profile["unique_ratio"] <= 0.15:
            score += 12
    elif field == "long_code_column":
        score += round(profile["product_ratio"] * 22)
    return max(0, min(100, score))


def _sheet_mapping_analysis(ws, known_product_codes=None):
    max_header_row = min(ws.max_row, 20)
    max_columns = min(ws.max_column, 40)
    best = None
    for header_row in range(1, max_header_row + 1):
        headers = [ws.cell(header_row, c).value for c in range(1, max_columns + 1)]
        start_row = min(ws.max_row + 1, header_row + 1)
        columns = []
        for column_number, header in enumerate(headers, 1):
            profile = _column_profile(ws, column_number, start_row, known_product_codes)
            field_scores = {
                field: _field_confidence(field, header, profile)
                for field in ("code_column", "description_column", "quantity_column", "container_column", "long_code_column")
            }
            columns.append({
                "number": column_number,
                "letter": get_column_letter(column_number),
                "header": str(header or "").strip(),
                "profile": profile,
                "scores": field_scores,
            })
        chosen = {}
        confidences = {}
        for field in ("code_column", "description_column", "quantity_column", "container_column", "long_code_column"):
            ranked = sorted(columns, key=lambda c: c["scores"][field], reverse=True)
            candidate = ranked[0] if ranked else None
            confidence = candidate["scores"][field] if candidate else 0
            # Do not force a mapping below the review threshold.
            chosen[field] = candidate["letter"] if candidate and confidence >= 75 else ""
            confidences[field] = confidence if candidate else 0
        required_quality = confidences["code_column"] + confidences["description_column"] + confidences["quantity_column"]
        if confidences["container_column"] >= 75:
            required_quality += confidences["container_column"] * 0.65
        header_matches = sum(1 for c in columns for score in c["scores"].values() if score >= 60)
        sheet_score = required_quality + header_matches * 4
        candidate = {
            "sheet": ws,
            "header_row": header_row,
            "start_row": start_row,
            "columns": columns,
            "mapping": chosen,
            "confidence": confidences,
            "score": sheet_score,
        }
        if best is None or candidate["score"] > best["score"]:
            best = candidate
    return best


def _structure_signature(ws, header_row):
    headers = [_normalise_header(ws.cell(header_row, c).value) for c in range(1, min(ws.max_column, 40) + 1)]
    # Worksheet name is intentionally excluded: the same customer layout can arrive under a renamed sheet.
    raw = f"cols:{len(headers)}|" + "|".join(headers)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:24]


def analyse_workbook_mapping(path, kind, known_product_codes=None, saved_profiles=None, sheet_name=None):
    """Analyse the workbook and return a conservative mapping plus confidence and container summary."""
    wb = _workbook(path)
    known = {re.sub(r"\s+", "", str(v).upper()) for v in (known_product_codes or set()) if v}

    # Headerless 2-CODES scanner files use one physical Code column plus one
    # LongCode column. Pallet rows (T9999) can be interleaved in the Code column,
    # therefore the first real source row must be preserved instead of being
    # mistaken for a header row.
    if kind == "RECEIVED":
        paired_candidates = []
        for ws in wb.worksheets:
            if sheet_name and ws.title != sheet_name:
                continue
            max_cols = min(ws.max_column, 12)
            max_rows = min(ws.max_row, 500)
            for code_column_number in range(1, max_cols + 1):
                for long_column_number in range(1, max_cols + 1):
                    if long_column_number == code_column_number:
                        continue
                    first_row = None
                    product_rows = 0
                    long_hits = 0
                    location_hits = 0
                    paired_rows = 0
                    code_nonempty = 0
                    long_nonempty = 0
                    for row_number in range(1, max_rows + 1):
                        code_value = ws.cell(row_number, code_column_number).value
                        long_value = ws.cell(row_number, long_column_number).value
                        if code_value not in (None, "") and first_row is None:
                            first_row = row_number
                        if code_value not in (None, ""):
                            code_nonempty += 1
                        if long_value not in (None, ""):
                            long_nonempty += 1
                        code_text = re.sub(r"\s+", "", str(code_value or "").strip().upper())
                        long_text = str(long_value or "").strip().upper()
                        if LOCATION_PATTERN.fullmatch(code_text):
                            location_hits += 1
                            continue
                        if PRODUCT_CODE_PATTERN.fullmatch(code_text):
                            product_rows += 1
                            if long_value not in (None, ""):
                                paired_rows += 1
                                if is_pon_long_code(long_text):
                                    long_hits += 1
                    if product_rows < 2 or paired_rows < 2 or not first_row:
                        continue
                    long_ratio = long_hits / paired_rows
                    paired_ratio = paired_rows / product_rows
                    if long_ratio < 0.70 or paired_ratio < 0.70:
                        continue
                    other_nonempty = 0
                    for other_col in range(1, max_cols + 1):
                        if other_col in {code_column_number, long_column_number}:
                            continue
                        other_nonempty += sum(
                            1 for row_number in range(1, max_rows + 1)
                            if ws.cell(row_number, other_col).value not in (None, "")
                        )
                    if other_nonempty > max(8, int((code_nonempty + long_nonempty) * 0.15)):
                        continue
                    paired_candidates.append({
                        "sheet": ws,
                        "code_column": get_column_letter(code_column_number),
                        "long_column": get_column_letter(long_column_number),
                        "first_row": first_row,
                        "products": product_rows,
                        "locations": location_hits,
                        "long_hits": long_hits,
                        "score": (long_ratio * 100) + (paired_ratio * 50) + min(location_hits, 25),
                    })
        if paired_candidates:
            best_pair = max(paired_candidates, key=lambda item: item["score"])
            ws = best_pair["sheet"]
            signature_raw = f"two-code-scan-stream|{best_pair['code_column']}|{best_pair['long_column']}"
            signature = hashlib.sha256(signature_raw.encode("utf-8")).hexdigest()[:24]
            suggestion = {
                "sheet_name": ws.title,
                "start_row": best_pair["first_row"],
                "start_column": best_pair["code_column"],
                "row_step": 1,
                "code_column": best_pair["code_column"],
                "long_code_column": best_pair["long_column"],
                "description_column": "",
                "quantity_column": "",
                "container_column": "",
                "coloured_rows_only": False,
                "required_fill_signature": "",
                "location_stream": True,
            }
            profile_match = None
            for profile in saved_profiles or []:
                profile_mapping = dict(getattr(profile, "mapping", {}) or {})
                if profile_mapping.get("_structure_signature") == signature:
                    for key in ("start_column", "row_step", "code_column", "long_code_column", "description_column", "quantity_column", "location_stream"):
                        if key in profile_mapping:
                            suggestion[key] = profile_mapping[key]
                    profile_match = getattr(profile, "name", "Saved profile")
                    break
            confidence = {
                "code_column": 100, "long_code_column": 100, "description_column": 0,
                "quantity_column": 0, "container_column": 0,
            }
            suggestion["_structure_signature"] = signature
            suggestion["_header_row"] = 0
            suggestion["_mapping_confidence"] = confidence
            suggestion["_analysis_version"] = "2026.10"
            return {
                "mapping": suggestion,
                "confidence": confidence,
                "field_labels": {
                    "code_column": f"{best_pair['code_column']} · Bike Code / pallet stream",
                    "long_code_column": f"{best_pair['long_column']} · LongCode",
                    "description_column": "Not required for scanned bikes",
                    "quantity_column": "Not required; each bike is one unit",
                    "container_column": "Not applicable",
                },
                "worksheet": ws.title,
                "header_row": 0,
                "data_rows": best_pair["products"],
                "bike_units": best_pair["products"],
                "spare_units": 0,
                "containers": [],
                "profile_match": profile_match,
                "structure_signature": signature,
            }

    # Preserve the warehouse single-column First Scan format: T9999 rows change location,
    # bike codes follow underneath. This format has no header row and must start at row 1.
    if kind == "RECEIVED":
        stream_candidates = []
        for ws in wb.worksheets:
            if sheet_name and ws.title != sheet_name:
                continue
            max_cols = min(ws.max_column, 12)
            max_rows = min(ws.max_row, 300)
            for column_number in range(1, max_cols + 1):
                values = []
                first_row = None
                for row_number in range(1, max_rows + 1):
                    value = ws.cell(row_number, column_number).value
                    if value not in (None, ""):
                        if first_row is None:
                            first_row = row_number
                        values.append(value)
                if not values:
                    continue
                location_hits = sum(1 for v in values if LOCATION_PATTERN.fullmatch(str(v).strip().upper()))
                product_hits = sum(
                    1 for v in values
                    if PRODUCT_CODE_PATTERN.fullmatch(re.sub(r"\s+", "", str(v).strip().upper()))
                    and not LOCATION_PATTERN.fullmatch(str(v).strip().upper())
                )
                other_nonempty = 0
                for other_col in range(1, max_cols + 1):
                    if other_col == column_number:
                        continue
                    other_nonempty += sum(
                        1 for row_number in range(1, max_rows + 1)
                        if ws.cell(row_number, other_col).value not in (None, "")
                    )
                stream_ratio = (location_hits + product_hits) / len(values)
                if location_hits >= 1 and product_hits >= 1 and stream_ratio >= 0.75 and other_nonempty <= max(5, int(len(values) * 0.20)):
                    stream_candidates.append({
                        "sheet": ws,
                        "column": get_column_letter(column_number),
                        "first_row": first_row or 1,
                        "locations": location_hits,
                        "products": product_hits,
                        "score": stream_ratio * 100 + min(location_hits, 25),
                    })
        if stream_candidates:
            best_stream = max(stream_candidates, key=lambda item: item["score"])
            ws = best_stream["sheet"]
            signature_raw = f"location-stream|{best_stream['column']}"
            signature = hashlib.sha256(signature_raw.encode("utf-8")).hexdigest()[:24]
            suggestion = {
                "sheet_name": ws.title,
                "start_row": best_stream["first_row"],
                "start_column": best_stream["column"],
                "row_step": 1,
                "code_column": best_stream["column"],
                "long_code_column": "",
                "description_column": "",
                "quantity_column": "",
                "container_column": "",
                "coloured_rows_only": False,
                "required_fill_signature": "",
                "location_stream": True,
            }
            profile_match = None
            for profile in saved_profiles or []:
                profile_mapping = dict(getattr(profile, "mapping", {}) or {})
                if profile_mapping.get("_structure_signature") == signature:
                    for key in ("start_column", "row_step", "code_column", "long_code_column", "description_column", "quantity_column", "location_stream"):
                        if key in profile_mapping:
                            suggestion[key] = profile_mapping[key]
                    profile_match = getattr(profile, "name", "Saved profile")
                    break
            confidence = {
                "code_column": 100, "long_code_column": 0, "description_column": 0,
                "quantity_column": 0, "container_column": 0,
            }
            suggestion["_structure_signature"] = signature
            suggestion["_header_row"] = 0
            suggestion["_mapping_confidence"] = confidence
            suggestion["_analysis_version"] = "2026.09"
            return {
                "mapping": suggestion,
                "confidence": confidence,
                "field_labels": {
                    "code_column": f"{best_stream['column']} · Single-column scan stream",
                    "long_code_column": "Not reliably detected",
                    "description_column": "Not required for scan stream",
                    "quantity_column": "Not required; each scan is one unit",
                    "container_column": "Not applicable",
                },
                "worksheet": ws.title,
                "header_row": 0,
                "data_rows": best_stream["products"],
                "bike_units": best_stream["products"],
                "spare_units": 0,
                "containers": [],
                "profile_match": profile_match,
                "structure_signature": signature,
            }

    candidates = []
    for ws in wb.worksheets:
        if sheet_name and ws.title != sheet_name:
            continue
        candidates.append(_sheet_mapping_analysis(ws, known))
    candidates = [c for c in candidates if c]
    if not candidates:
        raise ValueError("No readable worksheet was found.")
    analysis = max(candidates, key=lambda item: item["score"])
    ws = analysis["sheet"]
    first_source_column = next(
        (column["letter"] for column in analysis["columns"] if column["header"] or column["profile"]["nonempty"]),
        "A",
    )
    suggestion = {
        "sheet_name": ws.title,
        "start_row": analysis["start_row"],
        "start_column": first_source_column,
        "row_step": 1,
        "code_column": analysis["mapping"]["code_column"],
        "long_code_column": analysis["mapping"]["long_code_column"],
        "description_column": analysis["mapping"]["description_column"],
        "quantity_column": analysis["mapping"]["quantity_column"],
        "container_column": analysis["mapping"]["container_column"] if kind == "CLIENT" else "",
        "coloured_rows_only": False,
        "required_fill_signature": "",
        "location_stream": kind == "RECEIVED",
    }
    signature = _structure_signature(ws, analysis["header_row"])
    profile_match = None
    for profile in saved_profiles or []:
        profile_mapping = dict(getattr(profile, "mapping", {}) or {})
        if profile_mapping.get("_structure_signature") == signature:
            for key in ("start_column", "row_step", "code_column", "long_code_column", "description_column", "quantity_column", "container_column", "coloured_rows_only", "required_fill_signature", "location_stream"):
                if key in profile_mapping:
                    suggestion[key] = profile_mapping[key]
            for field in analysis["confidence"]:
                if suggestion.get(field):
                    analysis["confidence"][field] = 100
            profile_match = getattr(profile, "name", "Saved profile")
            break

    if kind == "CLIENT" and suggestion.get("code_column"):
        green_signatures = defaultdict(int)
        for row_number in range(int(suggestion["start_row"]), ws.max_row + 1):
            code_cell = ws[f"{suggestion['code_column']}{row_number}"]
            if code_cell.value in (None, ""):
                continue
            for cell in ws[row_number]:
                if _is_green_fill(cell, wb):
                    green_signatures[_cell_fill_signature(cell)] += 1
                    break
        if green_signatures:
            suggestion["coloured_rows_only"] = True
            suggestion["required_fill_signature"] = max(green_signatures, key=green_signatures.get)

    detected = defaultdict(lambda: {"rows": 0, "units": 0, "spare_units": 0})
    data_rows = 0
    bike_units = 0
    spare_units = 0
    identifier_col = suggestion.get("code_column") or suggestion.get("long_code_column")
    container_col = suggestion.get("container_column")
    quantity_col = suggestion.get("quantity_column")
    current_section = ""
    if identifier_col:
        for row_number in range(int(suggestion["start_row"]), ws.max_row + 1, int(suggestion.get("row_step") or 1)):
            section = _manifest_section_value(ws, suggestion, row_number) if kind == "CLIENT" else ""
            if section:
                current_section = section
                continue
            code = ws[f"{identifier_col}{row_number}"].value
            if code in (None, ""):
                continue
            data_rows += 1
            qty_raw = ws[f"{quantity_col}{row_number}"].value if quantity_col else 1
            units = int(Decimal(str(qty_raw))) if _positive_whole_number(qty_raw) else 1
            is_spare = kind == "CLIENT" and is_non_scannable_manifest_section(current_section)
            if is_spare:
                spare_units += units
            else:
                bike_units += units
            if kind == "CLIENT" and container_col:
                identifier = normalize_container_identifier(ws[f"{container_col}{row_number}"].value)
                if identifier:
                    detected[identifier]["rows"] += 1
                    if is_spare:
                        detected[identifier]["spare_units"] += units
                    else:
                        detected[identifier]["units"] += units

    headers_by_letter = {c["letter"]: c["header"] for c in analysis["columns"]}
    field_labels = {}
    for field in ("code_column", "long_code_column", "description_column", "quantity_column", "container_column"):
        letter = suggestion.get(field, "")
        field_labels[field] = f"{letter} · {headers_by_letter.get(letter, '')}".rstrip(" ·") if letter else "Not reliably detected"

    suggestion["_structure_signature"] = signature
    suggestion["_header_row"] = analysis["header_row"]
    suggestion["_mapping_confidence"] = dict(analysis["confidence"])
    suggestion["_analysis_version"] = "2026.09"
    return {
        "mapping": suggestion,
        "confidence": dict(analysis["confidence"]),
        "field_labels": field_labels,
        "worksheet": ws.title,
        "header_row": analysis["header_row"],
        "data_rows": data_rows,
        "bike_units": bike_units,
        "spare_units": spare_units,
        "containers": [
            {
                "identifier": identifier,
                "rows": values["rows"],
                "units": values["units"],
                "spare_units": values["spare_units"],
            }
            for identifier, values in sorted(detected.items())
        ],
        "profile_match": profile_match,
        "structure_signature": signature,
    }


def suggest_mapping(path, kind, sheet_name=None, known_product_codes=None, saved_profiles=None):
    """Compatibility wrapper used by tests and older code paths."""
    return analyse_workbook_mapping(
        path,
        kind,
        known_product_codes=known_product_codes,
        saved_profiles=saved_profiles,
        sheet_name=sheet_name,
    )["mapping"]


def _theme_palette(wb):
    if not wb.loaded_theme:
        return []
    root = ElementTree.fromstring(wb.loaded_theme)
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", namespace)
    palette = []
    if scheme is None:
        return palette
    for item in list(scheme):
        colour = next(iter(item), None)
        if colour is None:
            palette.append("")
        else:
            palette.append(colour.attrib.get("lastClr") or colour.attrib.get("val") or "")
    return palette


def _apply_tint(channel, tint):
    if tint < 0:
        return round(channel * (1 + tint))
    return round(channel + (255 - channel) * tint)


def _resolved_fill_rgb(cell, wb):
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return None
    colour = fill.fgColor
    raw = ""
    if colour.type == "rgb" and colour.rgb:
        raw = colour.rgb[-6:]
    elif colour.type == "theme" and colour.theme is not None:
        palette = _theme_palette(wb)
        if 0 <= colour.theme < len(palette):
            raw = palette[colour.theme][-6:]
    if len(raw) != 6:
        return None
    try:
        red, green, blue = (int(raw[index:index + 2], 16) for index in (0, 2, 4))
    except ValueError:
        return None
    tint = colour.tint or 0
    return tuple(_apply_tint(channel, tint) for channel in (red, green, blue))


def _is_green_fill(cell, wb):
    resolved = _resolved_fill_rgb(cell, wb)
    if not resolved:
        return False
    red, green, blue = resolved
    return green >= red * 1.08 and green >= blue * 1.08 and green - min(red, blue) >= 12


def _cell_fill_signature(cell):
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    colour = fill.fgColor
    if colour.type == "rgb":
        return f"rgb:{colour.rgb}"
    if colour.type == "theme":
        return f"theme:{colour.theme}:tint:{round(colour.tint or 0, 4)}"
    if colour.type == "indexed":
        return f"indexed:{colour.indexed}"
    return f"{colour.type}:{colour.value}"


def _row_fill_signature(ws, row_number):
    signatures = [_cell_fill_signature(ws.cell(row_number, col)) for col in range(1, ws.max_column + 1)]
    return next((value for value in signatures if value), "")


def _value(ws, column_letter, row_number):
    if not column_letter:
        return None
    index = column_index_from_string(column_letter.upper())
    return ws.cell(row_number, index).value


def _clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _quantity(value):
    if value in (None, ""):
        return 1
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"Invalid quantity: {value}") from exc
    if number <= 0 or number != number.to_integral_value():
        raise ValueError(f"Quantity must be a positive whole number: {value}")
    return int(number)


def _manifest_section_value(ws, mapping, row_number):
    """Return a repeated section label without hard-coding customer group names."""
    mapped_columns = []
    for field in ("code_column", "long_code_column", "description_column", "quantity_column"):
        column = str(mapping.get(field, "") or "").strip().upper()
        if column and column not in mapped_columns:
            mapped_columns.append(column)
    values = [_clean_text(_value(ws, column, row_number)) for column in mapped_columns]
    values = [value for value in values if value]
    if len(values) < 2:
        return ""
    normalised = [re.sub(r"\s+", " ", value).strip().upper() for value in values]
    if len(set(normalised)) != 1:
        return ""
    quantity_value = _value(ws, mapping.get("quantity_column", ""), row_number)
    try:
        _quantity(quantity_value)
    except ValueError:
        return re.sub(r"\s+", " ", values[0]).strip()
    return ""


def parse_workbook(path, mapping):
    wb = _workbook(path)
    ws = wb[mapping["sheet_name"]]
    start_row = int(mapping.get("start_row") or 1)
    row_step = int(mapping.get("row_step") or 1)
    code_column = (mapping.get("code_column") or "").upper()
    long_code_column = (mapping.get("long_code_column") or "").upper()
    identifier_column = code_column or long_code_column
    coloured_only = bool(mapping.get("coloured_rows_only"))
    required_fill = mapping.get("required_fill_signature", "")
    location_stream = bool(mapping.get("location_stream"))
    current_location = ""
    current_section = ""
    rows = []
    errors = []

    # A configured first data row can intentionally start below a visible
    # section heading. Read only the preceding labels so the first product row
    # still retains its source section; no source row is rewritten or merged.
    for prior_row in range(1, start_row):
        section = _manifest_section_value(ws, mapping, prior_row)
        if section:
            current_section = section

    for row_number in range(start_row, ws.max_row + 1, row_step):
        section = _manifest_section_value(ws, mapping, row_number)
        if section:
            current_section = section
            continue
        raw_code = _value(ws, identifier_column, row_number)
        code = _clean_text(raw_code).upper()
        if not code:
            continue

        if location_stream and LOCATION_PATTERN.fullmatch(code):
            current_location = code
            continue

        fill_signature = _row_fill_signature(ws, row_number)
        if coloured_only and not fill_signature:
            continue
        if coloured_only and required_fill and fill_signature != required_fill:
            continue

        try:
            quantity = _quantity(_value(ws, mapping.get("quantity_column", ""), row_number))
        except ValueError as exc:
            errors.append({"row": row_number, "message": str(exc)})
            continue

        container_identifier = ""
        container_column = mapping.get("container_column", "")
        if container_column:
            raw_container = _value(ws, container_column, row_number)
            container_identifier = normalize_container_identifier(raw_container)
            if raw_container not in (None, "") and not container_identifier:
                errors.append({"row": row_number, "message": f"Invalid container identifier: {raw_container}"})
                continue
            if not container_identifier:
                errors.append({"row": row_number, "message": "Container identifier is required for this mapped row."})
                continue

        raw_data = {}
        for col in range(1, min(ws.max_column, 50) + 1):
            cell = ws.cell(row_number, col)
            if cell.value not in (None, ""):
                raw_data[cell.column_letter] = _clean_text(cell.value)
        if current_section:
            raw_data["_manifest_section"] = current_section

        rows.append(
            {
                "source_sheet": ws.title,
                "source_row": row_number,
                "code": code,
                "long_code": _clean_text(_value(ws, long_code_column, row_number)).upper(),
                "description": _clean_text(_value(ws, mapping.get("description_column", ""), row_number)),
                "quantity": quantity,
                "location": current_location,
                "fill_signature": fill_signature,
                "raw_data": raw_data,
                "container_identifier": container_identifier,
            }
        )
    return rows, errors


def compare_lines(client_lines, received_lines, two_codes=None):
    client_lines = list(client_lines)
    received_lines = list(received_lines)
    if two_codes is None:
        two_codes = _uses_two_code_reconciliation(received_lines)
    expected = defaultdict(int)
    received = defaultdict(int)
    source_rows = defaultdict(list)
    sections = defaultdict(list)
    physical_codes = defaultdict(set)
    for line in client_lines:
        if is_non_scannable_manifest_line(line):
            continue
        code = _line_reconciliation_code(line, use_long_code=two_codes)
        if not code or (two_codes and not is_pon_long_code(code)):
            continue
        expected[code] += line.quantity
        source_rows[code].append(line.source_row)
        section = manifest_section(line)
        if section and section not in sections[code]:
            sections[code].append(section)
    for line in received_lines:
        code = _line_reconciliation_code(line, use_long_code=two_codes)
        if not code:
            continue
        received[code] += line.quantity
        if two_codes:
            physical_code = str(getattr(line, "code", "") or "").strip().upper()
            if physical_code:
                physical_codes[code].add(physical_code)

    result = []
    for code in sorted(set(expected) | set(received)):
        exp = expected[code]
        rec = received[code]
        if exp == rec:
            status = "MATCH"
        elif rec == 0:
            status = "MISSING"
        elif exp == 0:
            status = "UNADVISED" if two_codes else "UNEXPECTED"
        elif rec < exp:
            status = "SHORT"
        else:
            status = "OVER"
        quantity_status = status
        observed_codes = sorted(physical_codes.get(code, set()))
        warning = ""
        if len(observed_codes) > 1:
            status = "REVIEW"
            warning = "Long Code observed with multiple physical Codes: " + ", ".join(observed_codes)
        result.append({
            "code": code,
            "expected": exp,
            "received": rec,
            "difference": rec - exp,
            "status": status,
            "quantity_status": quantity_status,
            "warning": warning,
            "physical_codes": observed_codes,
            "sections": sections.get(code, []),
            "source_rows": source_rows.get(code, []),
        })
    return result
