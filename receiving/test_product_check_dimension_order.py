from io import BytesIO
from types import SimpleNamespace

from django.test import SimpleTestCase
from openpyxl import load_workbook

from .services import build_product_check_workbook

PATCH_MARKER = "PON_PRODUCT_CHECK_DIMENSION_ORDER_1008_1540"


class ProductCheckDimensionOrderTests(SimpleTestCase):
    def _row(self, definition=None, history=None, is_pbp=False):
        return {
            "code": "BIKE-ORDER-TEST",
            "long_code": "LONG-ORDER-TEST",
            "count": 1,
            "location": "T0001",
            "status": "New",
            "definition": definition,
            "pbp_to_pon_data": history or {},
            "is_pbp_to_pon": is_pbp,
            "is_new": not is_pbp,
            "tl_code": "",
            "matched_by": "",
        }

    def _sheet(self, row):
        output = build_product_check_workbook("TEST", [row])
        return load_workbook(BytesIO(output.getvalue()), data_only=True).active

    def test_headers_and_definition_values_are_length_height_width(self):
        definition = SimpleNamespace(
            length_mm=1340,
            height_mm=750,
            width_mm=250,
            weight_kg=15.63,
        )
        ws = self._sheet(self._row(definition=definition))
        self.assertEqual(
            [ws.cell(1, column).value for column in range(6, 9)],
            ["Length (cm)", "Height (cm)", "Width (cm)"],
        )
        self.assertEqual(
            [ws.cell(2, column).value for column in range(6, 9)],
            [134.0, 75.0, 25.0],
        )

    def test_pbp_history_values_use_same_length_height_width_order(self):
        history = {
            "valid": True,
            "length_mm": 1340,
            "height_mm": 750,
            "width_mm": 250,
            "weight_kg": 15.63,
        }
        ws = self._sheet(self._row(history=history, is_pbp=True))
        self.assertEqual(
            [ws.cell(2, column).value for column in range(6, 9)],
            [134.0, 75.0, 25.0],
        )
