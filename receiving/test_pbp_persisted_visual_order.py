from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from django.test import SimpleTestCase
from openpyxl import load_workbook

from .services import EXPORT_DEFAULTS, build_new_product_workbook, build_product_check_workbook
from .views import (
    _historical_import_selection,
    _persisted_import_classification,
    _replication_start_for_product_forms,
)


PATCH_MARKER = "PON_PBP_PERSISTED_VISUAL_ORDER_1009_0815"
PBP_YELLOW = "FFFFFF00"


class PbpPersistedVisualOrderTests(SimpleTestCase):
    def _definition(self, code, is_pbp=False):
        return SimpleNamespace(
            code=code,
            long_code=f"LONG-{code}",
            short_name=f"SHORT {code}"[:30],
            full_name=f"Full name {code}",
            group1="PONCVL",
            length_mm=1340,
            height_mm=750,
            width_mm=250,
            weight_kg=Decimal("15.630"),
            cubic_override=Decimal("0.25125"),
            is_pbp_to_pon=is_pbp,
        )

    def _row(self, code, location, classification, definition=True):
        is_pbp = classification == "PBP_TO_PON"
        is_new = classification == "NEW"
        history = {
            "valid": True,
            "length_mm": 1340,
            "height_mm": 750,
            "width_mm": 250,
            "weight_kg": Decimal("15.630"),
        } if is_pbp else {}
        return {
            "code": code,
            "long_code": f"LONG-{code}",
            "count": 1,
            "location": location,
            "classification": classification,
            "is_new": is_new,
            "is_pbp_to_pon": is_pbp,
            "status": "PBP → PON" if is_pbp else ("New" if is_new else "Existing"),
            "definition": self._definition(code, is_pbp) if definition else None,
            "pbp_to_pon_data": history,
            "tl_code": "",
            "matched_by": "",
        }

    def _product_check_sheet(self, rows):
        output = build_product_check_workbook("ORDER-TEST", rows)
        return load_workbook(BytesIO(output.getvalue()), data_only=True)["PON Product Check"]

    def _new_product_sheet(self, products, pbp_codes=None):
        output = build_new_product_workbook(
            products,
            dict(EXPORT_DEFAULTS),
            date(2026, 10, 8),
            pbp_to_pon_codes=pbp_codes,
        )
        return load_workbook(BytesIO(output.getvalue()), data_only=True)["New Product Import"]

    def test_pbp_rows_appear_before_truly_new(self):
        rows = [
            self._row("NEW-1", "T0001", "NEW"),
            self._row("PBP-B", "T9999", "PBP_TO_PON"),
            self._row("PBP-A", "", "PBP_TO_PON"),
        ]
        sheet = self._product_check_sheet(rows)
        self.assertEqual([sheet.cell(row, 1).value for row in range(2, 5)], ["PBP-A", "PBP-B", "NEW-1"])

    def test_truly_new_rows_keep_physical_pallet_order(self):
        rows = [
            self._row("NEW-LATE", "T2000", "NEW"),
            self._row("PBP-1", "", "PBP_TO_PON"),
            self._row("NEW-EARLY", "T1000", "NEW"),
        ]
        sheet = self._product_check_sheet(rows)
        self.assertEqual([sheet.cell(row, 1).value for row in range(3, 5)], ["NEW-EARLY", "NEW-LATE"])

    def test_pbp_rows_are_outside_replicate_tail(self):
        forms = [
            {"row": {"is_pbp_to_pon": True, "is_new": False}},
            {"row": {"is_pbp_to_pon": True, "is_new": False}},
            {"row": {"is_pbp_to_pon": False, "is_new": True}},
        ]
        self.assertEqual(_replication_start_for_product_forms(forms), 2)

    def test_product_check_colours_entire_pbp_row_yellow(self):
        sheet = self._product_check_sheet([self._row("PBP-YELLOW", "", "PBP_TO_PON")])
        self.assertTrue(all(cell.fill.fill_type == "solid" for cell in sheet[2]))
        self.assertTrue(all(cell.fill.fgColor.rgb == PBP_YELLOW for cell in sheet[2]))

    def test_product_check_does_not_colour_truly_new_row_pbp_yellow(self):
        sheet = self._product_check_sheet([self._row("NEW-NORMAL", "T0001", "NEW")])
        self.assertTrue(any(cell.fill.fgColor.rgb != PBP_YELLOW for cell in sheet[2]))
        self.assertNotEqual(sheet["E2"].fill.fgColor.rgb, PBP_YELLOW)

    def test_new_product_import_colours_entire_pbp_row_yellow(self):
        sheet = self._new_product_sheet([self._definition("PBP-IMPORT", is_pbp=True)])
        self.assertTrue(all(cell.fill.fill_type == "solid" for cell in sheet[2]))
        self.assertTrue(all(cell.fill.fgColor.rgb == PBP_YELLOW for cell in sheet[2]))

    def test_persisted_pbp_code_stays_yellow_when_current_object_is_not_pbp(self):
        product = self._definition("CURRENT-PON", is_pbp=False)
        sheet = self._new_product_sheet([product], pbp_codes={"CURRENT-PON"})
        self.assertTrue(all(cell.fill.fgColor.rgb == PBP_YELLOW for cell in sheet[2]))

    def test_original_existing_is_excluded_from_new_product_import_selection(self):
        rows = [
            self._row("EXISTING-1", "T0001", "EXISTING"),
            self._row("PBP-1", "", "PBP_TO_PON"),
            self._row("NEW-1", "T0002", "NEW"),
        ]
        _, _, import_codes = _historical_import_selection(rows)
        self.assertEqual(import_codes, ["PBP-1", "NEW-1"])

    def test_persisted_original_classification_dominates_cached_current(self):
        state = SimpleNamespace(import_classification="PBP_TO_PON")
        saved_current = {"classification": "EXISTING"}
        self.assertEqual(_persisted_import_classification(state, saved_current), "PBP_TO_PON")

    def test_length_height_width_order_remains_intact(self):
        sheet = self._product_check_sheet([self._row("NEW-DIMS", "T0001", "NEW")])
        self.assertEqual(
            [sheet.cell(1, column).value for column in range(6, 9)],
            ["Length (cm)", "Height (cm)", "Width (cm)"],
        )
        self.assertEqual([sheet.cell(2, column).value for column in range(6, 9)], [134, 75, 25])
