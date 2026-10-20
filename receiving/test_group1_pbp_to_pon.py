"""Regression tests for historical PBP Group1 mapping in New Product Import.

No production database changes. Django tests use an isolated test database.
"""
from datetime import date
from io import BytesIO
import tempfile
from types import SimpleNamespace

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from .models import (AppSetting, Container, Customer, Group1Family, ImportBatch,
                     NormalizedLine, ProductCatalog, ProductCatalogEntry, SourceFile)
from .views import _pbp_group1_mapping, _container_workspace_data


class OriginalPbpFamilyResolverTests(SimpleTestCase):
    def setUp(self):
        self.family = SimpleNamespace(
            pk=42, family_name="Cervelo", suffix="CVL", active=True,
            pbp_group1="PBPCVL", pon_group1="PONCVL",
        )
        self.container = SimpleNamespace(product_statuses=None)

    def test_matches_source_pbp_group1_without_name_phrase(self):
        row = {"code": "0L0CABFX1C54", "catalog_customer": "PBP",
               "catalog_group1": "PBPCVL", "pbp_to_pon_data": {"long_name": "CALEDONIA FORCE AXS MOCHA SIZE 54"}}
        result = _pbp_group1_mapping(self.container, row, [self.family])
        self.assertTrue(result["resolved"])
        self.assertEqual(result["group1"], "PONCVL")
        self.assertEqual(result["reason"], "original_pbp_group1")

    def test_historical_check_row_is_used_when_active_catalog_is_pon(self):
        row = {"code": "0L0CABFX1C54", "catalog_customer": "PON",
               "catalog_group1": "PONCVL",
               "pbp_to_pon_data": {"_check_row": {"catalog_customer": "PBP", "catalog_group1": "PBPCVL"}}}
        self.assertEqual(_pbp_group1_mapping(self.container, row, [self.family])["group1"], "PONCVL")

    def test_missing_pbp_group1_keeps_existing_fallback_available(self):
        self.assertIsNone(_pbp_group1_mapping(self.container,
                          {"code": "X", "catalog_customer": "PBP", "catalog_group1": ""}, [self.family]))

    def test_unmapped_group1_does_not_guess_from_name(self):
        row = {"code": "X", "catalog_customer": "PBP", "catalog_group1": "PBPXXX"}
        result = _pbp_group1_mapping(self.container, row, [self.family])
        self.assertFalse(result["resolved"])
        self.assertEqual(result["reason"], "unmapped_pbp_group1")

    def test_inactive_family_cannot_authorise_conversion(self):
        self.family.active = False
        row = {"code": "X", "catalog_customer": "PBP", "catalog_group1": "PBPCVL"}
        self.assertFalse(_pbp_group1_mapping(self.container, row, [self.family])["resolved"])

    def test_mapped_source_pbp_does_not_use_a_pon_catalog_value(self):
        row = {"code": "X", "catalog_customer": "PON", "catalog_group1": "PBPCVL"}
        self.assertIsNone(_pbp_group1_mapping(self.container, row, [self.family]))


class PbpNewProductDownloadRegressionTests(TestCase):
    SKU = "0L0CABFX1C54"

    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=self.media.name)
        self.override.enable()
        self.user = get_user_model().objects.create_user(username="group1-mapping", password="test")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(code="PON", name="Pon.Bike")
        self.container = Container.objects.create(
            identifier="ONEU6047465", customer=self.customer,
            container_date=date(2026, 10, 13), year=2026, created_by=self.user,
        )
        AppSetting.objects.create(key="TL_EXPORT_FORMAT", value="xlsx")
        source = SourceFile.objects.create(
            container=self.container, kind="RECEIVED",
            file=SimpleUploadedFile("received.xlsx", b"placeholder"),
            original_name="received.xlsx", sha256="1" * 64,
            status="CONFIRMED", uploaded_by=self.user,
        )
        batch = ImportBatch.objects.create(
            source_file=source, status="CONFIRMED", total_rows=1,
            total_units=1, created_by=self.user,
        )
        NormalizedLine.objects.create(
            batch=batch, source_sheet="Original", source_row=3,
            code=self.SKU, quantity=1, location="T4486", long_code="",
        )
        self.catalog = ProductCatalog.objects.create(
            customer=self.customer,
            file=SimpleUploadedFile("PRODUCTS_PON_PBP.xml", b"placeholder"),
            original_name="PRODUCTS_PON_PBP.xml", sha256="2" * 64,
            row_count=1, active=True, uploaded_by=self.user,
        )
        self.entry = ProductCatalogEntry.objects.create(
            catalog=self.catalog, source_row=10, code=self.SKU, pon_sku="", code2="",
            customer="PBP", short_name="CaledoniaForceAXSMocha54",
            long_name="CALEDONIA FORCE AXS MOCHA SIZE 54",
            group1="PBPCVL", group2="", raw_data={
                "cubic": "0.239", "weight": "13.050", "length": "1350",
                "height": "770", "width": "230",
            },
        )
        # Migration 0006 seeds Cervelo in every Django test database.
        # Reuse the seeded mapping instead of violating unique family_name.
        self.family = Group1Family.objects.get(family_name="Cervelo")
        self.assertEqual(self.family.suffix, "CVL")
        self.assertTrue(self.family.active)
        self.assertNotIn("CALEDONIA", self.family.match_phrases.upper())

    def tearDown(self):
        self.override.disable()
        self.media.cleanup()

    def _export(self):
        return self.client.get(reverse("receiving:export_new_products", args=[self.container.pk]))

    def _check(self):
        return self.client.get(reverse("receiving:product_check", args=[self.container.pk]))

    def test_caledonia_pbp_exports_group1_from_catalog_without_caIedonia_keyword(self):
        self.assertEqual(self._check().status_code, 200)
        self.assertTrue(_container_workspace_data(self.container)["new_product_import_download_ready"])
        response = self._export()
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), data_only=True)["New Product Import"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, self.SKU)
        self.assertEqual(sheet["D2"].value, "PON")
        self.assertEqual(sheet["J2"].value, "PONCVL")
        self.assertEqual(sheet["M2"].value, "CALEDONIA FORCE AXS MOCHA SIZE 54")
        self.assertEqual(sheet["U2"].value, 13.05)
        self.assertEqual((sheet["V2"].value,sheet["W2"].value,sheet["X2"].value),(770,230,1350))
        self.assertIsNone(sheet["B2"].value, "An empty long_sku must not be invented or filled")
        self.entry.refresh_from_db()
        self.assertEqual(self.entry.group1, "PBPCVL", "PBP original must remain untouched")

    def test_unmapped_pbp_group1_blocks_export_even_if_longname_is_a_match(self):
        self.entry.group1 = "PBPXXX"
        self.entry.save(update_fields=["group1"])
        self.family.match_phrases = "CALEDONIA"
        self.family.save(update_fields=["match_phrases"])
        self._check()
        self.assertFalse(_container_workspace_data(self.container)["new_product_import_download_ready"])
        response = self._export()
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.container.generated_exports.filter(kind="NEW_PRODUCT").count(), 0)

    def test_group1_remains_mapped_after_unrelated_pon_catalog_sync(self):
        self._check()
        ProductCatalogEntry.objects.create(
            catalog=self.catalog, source_row=11, code=self.SKU, pon_sku=self.SKU,
            customer="PON", code2=self.SKU, short_name="new pon name",
            long_name="NEW PON PRODUCT", group1="PONCVL", raw_data={
                "cubic": "0.252", "weight": "15.63", "height": "750",
                "width": "250", "length": "1340",
            },
        )
        response = self._export()
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        self.assertEqual(sheet["J2"].value,"PONCVL")
        self.assertEqual(sheet["T2"].value,0.252)
        self.assertEqual(sheet["U2"].value,15.63)

    def test_inactive_pbp_mapping_blocks_export(self):
        self.family.active = False
        self.family.save(update_fields=["active"])
        self._check()
        response = self._export()
        self.assertEqual(response.status_code,302)

    def test_historical_pbp_catalog_lookup_restores_missing_check_snapshot(self):
        # Existing containers may have valid PBP details but a previous version
        # did not persist original Group1 in the compact history structure.
        self._check()
        state = self.container.product_statuses.get(code=self.SKU)
        history = dict(state.import_history)
        history.pop("_check_row", None)
        state.import_history = history
        state.save(update_fields=["import_history"])
        self.assertTrue(_container_workspace_data(self.container)["new_product_import_download_ready"])
        response = self._export()
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        self.assertEqual(sheet["J2"].value, "PONCVL")
        state.refresh_from_db()
        self.assertNotIn("source_pbp_group1", state.import_history)

    def test_informed_long_sku_is_preserved_exactly(self):
        self.entry.pon_sku = "LONG-ORIGINAL-123"
        self.entry.save(update_fields=["pon_sku"])
        self._check()
        response = self._export()
        self.assertEqual(response.status_code,200)
        sheet=load_workbook(BytesIO(response.content),data_only=True).active
        self.assertEqual(sheet["B2"].value,"LONG-ORIGINAL-123")
