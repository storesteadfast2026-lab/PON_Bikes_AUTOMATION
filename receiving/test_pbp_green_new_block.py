
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

from django.conf import settings
from django.test import SimpleTestCase
from openpyxl import load_workbook

from .services import EXPORT_DEFAULTS, build_new_product_workbook, convert_excel_output

PATCH_MARKER = 'PON_PBP_GREEN_NEW_BLOCK_1008_1500'
ORDER_MARKER = 'PON_PRODUCT_CHECK_SCREEN_PRINT_ORDER_1008_1515'
EXPECTED_YELLOW = 'FFFFFF00'

class PbpGreenNewBlockTests(SimpleTestCase):
    def _product(self, code="PBP-GREEN-1"):
        return SimpleNamespace(
            code=code,
            long_code="58-12345-123-1-123-123456",
            short_name="PBP MIGRATION BIKE",
            full_name="PBP Migration Bike",
            group1="PONCVL",
            length_mm=1340,
            height_mm=750,
            width_mm=250,
            weight_kg=Decimal("15.630"),
            cubic_override=Decimal("0.25125"),
            customer_override="PON",
            status_override="L",
            quantity_override=1,
            pallet_override=0,
            lift_override=1,
            outer_override=1,
            is_pbp_to_pon=True,
        )

    def test_product_check_screen_new_rows_follow_printed_order(self):
        views = Path(settings.BASE_DIR) / "receiving" / "views.py"
        text = views.read_text(encoding="utf-8")
        self.assertIn(PATCH_MARKER, text)
        self.assertIn(ORDER_MARKER, text)
        self.assertIn("_group_product_rows_by_code(rows)", text)
        self.assertIn(
            'non_new_rows = [row for row in unique_new_rows if row.get("is_pbp_to_pon")]',
            text,
        )
        self.assertIn(
            'if row.get("is_new") and not row.get("is_pbp_to_pon")',
            text,
        )
        self.assertIn("unique_new_rows = non_new_rows + truly_new_rows", text)
        self.assertIn(
            'replication_start_index = _replication_start_for_product_forms(product_forms)',
            text,
        )
        self.assertNotIn(
            'unique_new_rows.sort(key=lambda row: 1 if row.get("definition") is None else 0)',
            text,
        )

    def test_replicate_uses_only_tail_after_replication_start_index(self):
        template = Path(settings.BASE_DIR) / "receiving" / "templates" / "receiving" / "product_check.html"
        text = template.read_text(encoding="utf-8")
        self.assertIn(PATCH_MARKER, text)
        self.assertIn(
            "form.querySelectorAll('.new-product-row')).slice({{ replication_start_index|default:0 }})",
            text,
        )

    def test_pbp_new_product_import_is_yellow(self):
        workbook = build_new_product_workbook([self._product()], dict(EXPORT_DEFAULTS), date(2026, 10, 8))
        ws = load_workbook(workbook, data_only=True)["New Product Import"]
        self.assertTrue(all(cell.fill.fill_type == "solid" for cell in ws[2]))
        self.assertTrue(all(cell.fill.fgColor.type == "rgb" for cell in ws[2]))
        self.assertTrue(all(cell.fill.fgColor.rgb == EXPECTED_YELLOW for cell in ws[2]))

    def test_excel_5_95_conversion_still_succeeds(self):
        workbook = build_new_product_workbook([self._product("PBP-GREEN-XLS")], dict(EXPORT_DEFAULTS), date(2026, 10, 8))
        content, extension, _ = convert_excel_output(workbook, "pbp_green_conversion", "excel_5_95")
        self.assertEqual(extension.lower(), ".xls")
        self.assertGreater(len(content), 1000)
        self.assertTrue(content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")))
