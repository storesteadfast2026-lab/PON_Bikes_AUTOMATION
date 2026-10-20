from collections import defaultdict
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

from receiving.models import Container, ContainerProductStatus, ProductDefinition
from receiving.views import _persisted_product_check_data

# SEARCH_VALUE and OUTPUT are supplied by the Windows launcher.
# The workbook layout below is intentionally identical to the original exporter.
for required_parameter in ("SEARCH_VALUE", "OUTPUT"):
    if required_parameter not in globals():
        raise RuntimeError(f"Missing export argument: {required_parameter}")

search_value = str(SEARCH_VALUE or "").strip().upper()
if not re.fullmatch(r"[A-Z0-9_-]{1,40}", search_value):
    raise RuntimeError("Search must be a Container ID or Job Order (letters, digits, '_' or '-').")

# Resolve the actual container inside Django, never by interpolating arbitrary input
# into ORM lookup paths or changing the container's historical classification.
if re.fullmatch(r"[A-Z]{4}[0-9]{7}", search_value):
    candidates = list(Container.objects.filter(identifier__iexact=search_value)[:2])
    if not candidates:
        candidates = list(Container.objects.filter(job_order__iexact=search_value)[:2])
else:
    candidates = list(Container.objects.filter(job_order__iexact=search_value)[:2])
if not candidates:
    raise RuntimeError(f"No container found for Container ID / Job Order: {search_value}.")
if len(candidates) != 1:
    raise RuntimeError(f"More than one container matches {search_value}; export blocked.")
container = candidates[0]
CONTAINER_ID = str(container.identifier or "").strip().upper()
if not re.fullmatch(r"[A-Z0-9_-]{1,30}", CONTAINER_ID):
    raise RuntimeError(f"Unsafe container identifier in database: {CONTAINER_ID!r}.")
job_order = str(getattr(container, "job_order", "") or "").strip()

# ONEU is an acceptance fixture, not a global product allowlist.
EXPECTED_ONEU6047465 = {
    "0L0CABFX1C48",
    "0M0STBVX2C61",
    "0M0STBUI1C61",
    "0M0STBUI2C61",
}

statuses = list(ContainerProductStatus.objects.filter(container=container).order_by("code"))
if not statuses:
    raise RuntimeError(f"No saved Product Check classifications for {CONTAINER_ID}. Run Product Check first.")

status_by_code = {}
for state in statuses:
    code = str(state.code or "").strip().upper()
    if code in status_by_code:
        raise RuntimeError(f"Duplicate Product Check classification for code {code}.")
    status_by_code[code] = str(state.import_classification or "").strip().upper()

# Never infer 'NEW' from product names, ProductDefinition existence,
# today's product master or the old was_new flag.
new_codes = sorted(code for code, classification in status_by_code.items() if classification == "NEW")
pbp_codes = sorted(code for code, classification in status_by_code.items() if classification == "PBP_TO_PON")
other_codes = sorted(code for code, classification in status_by_code.items() if classification not in {"NEW", "PBP_TO_PON"})
if not new_codes:
    raise RuntimeError(f"No historical NEW products found in container {CONTAINER_ID}.")

if CONTAINER_ID == "ONEU6047465" and set(new_codes) != EXPECTED_ONEU6047465:
    missing = sorted(EXPECTED_ONEU6047465 - set(new_codes))
    extra = sorted(set(new_codes) - EXPECTED_ONEU6047465)
    raise RuntimeError(
        "ONEU6047465 historical NEW classification differs from the four accepted SKUs. "
        f"Missing={missing}; Unexpected NEW={extra}. Nothing exported. "
        "Review saved Product Check classifications; do not reclassify automatically."
    )

catalog, received_batch, rows, summary = _persisted_product_check_data(container)
if received_batch is None:
    raise RuntimeError(f"No confirmed RECEIVED batch found for container {CONTAINER_ID}.")

received = defaultdict(lambda: {"count": 0, "locations": [], "long_codes": []})
for row in rows:
    code = str(row.get("code") or "").strip().upper()
    # Keep PBP_TO_PON and EXISTING completely outside the export.
    if code not in new_codes:
        continue
    received[code]["count"] += int(row.get("count") or 0)
    loc = str(row.get("location") or "").strip()
    if loc and loc not in received[code]["locations"]:
        received[code]["locations"].append(loc)
    long_code = str(row.get("long_code") or "").strip()
    if long_code and long_code not in received[code]["long_codes"]:
        received[code]["long_codes"].append(long_code)

missing_received = [code for code in new_codes if received[code]["count"] <= 0]
if missing_received:
    raise RuntimeError(f"Historical NEW SKUs missing from confirmed received lines: {missing_received}.")

definitions = {
    str(p.code).strip().upper(): p
    for p in ProductDefinition.objects.filter(code__in=new_codes).order_by("code")
}

def positive(value):
    try:
        return value is not None and Decimal(str(value)) > 0
    except Exception:
        return False

invalid = []
export_rows = []
for code in new_codes:
    definition = definitions.get(code)
    if definition is None:
        invalid.append(f"{code}: no saved ProductDefinition")
        continue
    dims = {
        "Length": definition.length_mm,
        "Height": definition.height_mm,
        "Width": definition.width_mm,
        "Weight": definition.weight_kg,
    }
    missing_fields = [name for name, value in dims.items() if not positive(value)]
    if missing_fields:
        invalid.append(f"{code}: missing/invalid {', '.join(missing_fields)}")
        continue
    length_mm = Decimal(str(definition.length_mm))
    height_mm = Decimal(str(definition.height_mm))
    width_mm = Decimal(str(definition.width_mm))
    weight_kg = Decimal(str(definition.weight_kg))
    info = received[code]
    export_rows.append([
        code,
        " / ".join(info["long_codes"]),
        info["count"],
        " / ".join(info["locations"]),
        "New",
        float(length_mm / Decimal("10")),
        float(height_mm / Decimal("10")),
        float(width_mm / Decimal("10")),
        float(weight_kg),
        float((length_mm * height_mm * width_mm) / Decimal("1000000000")),
        float(weight_kg),
        float(height_mm),
        float(width_mm),
        float(length_mm),
    ])

print("=" * 65)
print(f"CONTAINER: {CONTAINER_ID} | JOB ORDER: {job_order or 'Not recorded'}")
print(f"Historical NEW: {len(new_codes)} | PBP_TO_PON excluded: {len(pbp_codes)} | Other excluded: {len(other_codes)}")
print("NEW SKUs exported:")
for code in new_codes:
    print(f"  NEW  {code} | Bikes: {received[code]['count']}")
print("PBP_TO_PON SKUs excluded:")
for code in pbp_codes:
    print(f"  PBP  {code}")
print("=" * 65)

if invalid:
    raise RuntimeError("Cannot export incomplete saved measurements:\n" + "\n".join(invalid))

# Single-step export: no confirmation, no intermediate preview, same original sheet.
wb = Workbook()
ws = wb.active
ws.title = "New Products"
ws.append([
    CONTAINER_ID, "Long_Code", "Count", "LOC", "New",
    "L", "H", "W", "Kg", "M3", "Weight", "Height", "Width", "Length",
])
for item in export_rows:
    ws.append(item)
header_fill = PatternFill("solid", fgColor="D9EAF7")
new_fill = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=14):
    for cell in row:
        cell.border = border
        cell.fill = new_fill
        cell.alignment = Alignment(vertical="center")
for n in range(2, ws.max_row + 1):
    for col in ("F", "G", "H"):
        ws[f"{col}{n}"].number_format = "0.0"
    for col in ("I", "J", "K"):
        ws[f"{col}{n}"].number_format = "0.000"
    for col in ("L", "M", "N"):
        ws[f"{col}{n}"].number_format = "0"
for col, width in {
    "A": 20, "B": 22, "C": 10, "D": 18, "E": 10,
    "F": 10, "G": 10, "H": 10, "I": 10, "J": 11,
    "K": 11, "L": 12, "M": 12, "N": 12,
}.items():
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
# Keep identical to the original TGBU exporter: ONE sheet named New Products,
# same 14 columns, fills, borders, filters, number formats and widths.
wb.save(OUTPUT)
assert Path(OUTPUT).is_file()
print(f"[PASS] Excel created: {OUTPUT}; Products: {len(export_rows)}")
# Strictly one machine-readable marker on success for the Windows launcher.
print(f"PON_EXPORT_CONTAINER_ID={CONTAINER_ID}")
